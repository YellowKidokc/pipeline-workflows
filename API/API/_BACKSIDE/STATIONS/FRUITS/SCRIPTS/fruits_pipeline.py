#!/usr/bin/env python3
"""DeepSeek Fruits pilot: HTML -> evidence JSON, Excel, and preview HTML."""
from __future__ import annotations
import argparse, datetime as dt, hashlib, html, json, os, pathlib, re, shutil, time, urllib.request, uuid
from collections import Counter
from html.parser import HTMLParser
from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve()
FRUITS = ["love","joy","peace","patience","kindness","goodness","faithfulness","gentleness","self-control"]
RESULTS = {"FRUIT","ANTI_FRUIT","COUNTERFEIT","MIXED","UNKNOWN"}
APPS = {"EMBODIED","ADVOCATED","ATTRIBUTED","APPLIED_TO_OTHER","SELF_APPLIED","GAP_ACKNOWLEDGED","REPAIR_ATTEMPTED","SELF_EXEMPTED","PERFORMATIVE_CONTRADICTION","INSUFFICIENT_CONTEXT","NOT_APPLICABLE"}

def now(): return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00","Z")
def digest(p): return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def dump(p,v):
    p=pathlib.Path(p); p.parent.mkdir(parents=True,exist_ok=True); q=p.with_suffix(p.suffix+".tmp"); q.write_text(json.dumps(v,ensure_ascii=False,indent=2)+"\n",encoding="utf-8"); q.replace(p)

class Parser(HTMLParser):
    blocks={"p","li","blockquote","figcaption","td","th","h1","h2","h3","h4"}
    def __init__(self): super().__init__(convert_charrefs=True); self.skip=0; self.tag=""; self.buf=[]; self.out=[]
    def handle_starttag(self,t,a):
        if t in {"script","style","nav"}: self.skip+=1
        if not self.skip and t in self.blocks: self.flush(); self.tag=t
    def handle_endtag(self,t):
        if t in {"script","style","nav"} and self.skip: self.skip-=1
        elif not self.skip and t in self.blocks: self.flush()
    def handle_data(self,d):
        if not self.skip: self.buf.append(d)
    def flush(self):
        s=re.sub(r"\s+"," "," ".join(self.buf)).strip()
        if s:self.out.append((self.tag or "p",s))
        self.buf=[]

def sentences(path):
    p=Parser(); p.feed(path.read_text(encoding="utf-8",errors="replace")); p.flush(); title=path.stem; sec="Opening"; para=0; rows=[]
    for tag,text in p.out:
        if tag=="h1": title=text; sec=text; continue
        if tag in {"h2","h3","h4"}: sec=text; continue
        para+=1
        for s in re.split(r'(?<=[.!?])\s+(?=[\"\'“‘(]*[A-Z0-9])',text):
            if s.strip(): rows.append({"line":len(rows)+1,"section":sec,"paragraph":para,"sentence":s.strip()})
    for i,x in enumerate(rows): x["previous"]=rows[i-1]["sentence"] if i else ""; x["next"]=rows[i+1]["sentence"] if i+1<len(rows) else ""
    return title,rows

def call(cfg,system,user,max_tokens):
    body={"model":cfg["model"],"temperature":0.1,"max_tokens":max_tokens,"response_format":{"type":"json_object"},"messages":[{"role":"system","content":system},{"role":"user","content":user}]}
    req=urllib.request.Request(cfg["base_url"].rstrip("/")+"/chat/completions",data=json.dumps(body).encode(),method="POST",headers={"Authorization":"Bearer "+os.environ[cfg["api_key_env"]],"Content-Type":"application/json"})
    for n in range(3):
        try:
            with urllib.request.urlopen(req,timeout=cfg.get("timeout_seconds",300)) as r: raw=json.loads(r.read().decode())
            ch=raw["choices"][0]; content=ch["message"].get("content") or ""
            if ch.get("finish_reason")=="length" or not content.strip(): raise RuntimeError("empty or truncated response")
            return json.loads(content),{"model":raw.get("model",cfg["model"]),"usage":raw.get("usage"),"finish_reason":ch.get("finish_reason")}
        except Exception:
            if n==2: raise
            time.sleep(2**n)

def validate(v,rows):
    keys={"entities","sentence_assessments","passages","counterfeits","reader_impacts","repairs","relationships","summary"}
    if keys-set(v): raise ValueError(f"missing keys {sorted(keys-set(v))}")
    lookup={x["line"]:x["sentence"] for x in rows}; joined=re.sub(r"\s+"," "," ".join(lookup.values())).strip()
    for x in v["sentence_assessments"]:
        if x.get("line") not in lookup: raise ValueError(f"bad line {x.get('line')}")
        if x.get("fruit") not in FRUITS or x.get("result") not in RESULTS or x.get("standard_application") not in APPS: raise ValueError(f"bad classification line {x.get('line')}")
    for x in v["passages"]:
        q=re.sub(r"\s+"," ",str(x.get("exact_quotation","")).strip())
        if q and q not in joined: raise ValueError(f"non-verbatim quote {q[:80]}")

def clear(ws,r0,cols,upto=400):
    for r in range(r0,max(ws.max_row,upto)+1):
        for c in range(1,cols+1): ws.cell(r,c).value=None

def excel(v):
    if isinstance(v,list): return ", ".join(map(str,v))
    if isinstance(v,dict): return json.dumps(v,ensure_ascii=False)
    return v

def workbook(template,out,title,src,rows,v,receipt):
    shutil.copy2(template,out); w=load_workbook(out); s=w["Evaluation Setup"]
    for cell,val in {"B5":title,"B6":"Be Glad You're a Loser","B7":"Article narrator","B8":"Series readers","B9":now()[:10]}.items(): s[cell]=val
    e=w["Entity Register"]; clear(e,5,10)
    for r,x in enumerate(v["entities"],5):
        vals=[f"E{r-4}",x.get("name"),x.get("kind"),", ".join(x.get("aliases") or []),", ".join(x.get("roles") or []),x.get("mentions"),x.get("substantive_passages"),x.get("treatment"),x.get("evidence_sufficiency"),x.get("confidence")]
        for c,z in enumerate(vals,1):e.cell(r,c).value=excel(z)
    p=w["Passage Ratings"]; clear(p,5,14)
    pk=["passage_id","exact_quotation","actor","target","fruit","result","mechanism","power_relationship","benefits","bears_cost","importance","confidence","explanation","possible_repair"]
    for r,x in enumerate(v["passages"],5):
        for c,k in enumerate(pk,1):p.cell(r,c).value=excel(x.get(k))
    l=w["Sentence Layer"]; clear(l,5,19,max(400,len(rows)+4)); by={x["line"]:x for x in v["sentence_assessments"]}
    for r,x in enumerate(rows,5):
        a=by.get(x["line"],{}); vals=[x["line"],x["section"],x["paragraph"],"",x["previous"],x["sentence"],x["next"],a.get("speaker"),a.get("actor"),a.get("target"),a.get("subject"),a.get("audience"),a.get("standard_expressed"),a.get("applied_to"),a.get("standard_application"),a.get("fruit"),a.get("result"),a.get("rationale"),a.get("confidence")]
        for c,z in enumerate(vals,1):l.cell(r,c).value=z
    last=len(rows)+4; f=w["Fruit Profile"]
    for r in range(5,14):
        f.cell(r,2).value=f'=COUNTIF(\'Sentence Layer\'!$P$5:$P${last},A{r})'; f.cell(r,3).value=f'=COUNTIFS(\'Sentence Layer\'!$P$5:$P${last},A{r},\'Sentence Layer\'!$Q$5:$Q${last},"FRUIT")'; f.cell(r,4).value=f'=COUNTIFS(\'Sentence Layer\'!$P$5:$P${last},A{r},\'Sentence Layer\'!$Q$5:$Q${last},"ANTI_FRUIT")+COUNTIFS(\'Sentence Layer\'!$P$5:$P${last},A{r},\'Sentence Layer\'!$Q$5:$Q${last},"COUNTERFEIT")'; f.cell(r,6).value=f'=COUNTIFS(\'Sentence Layer\'!$P$5:$P${last},A{r},\'Sentence Layer\'!$Q$5:$Q${last},"MIXED")'; f.cell(r,7).value=f'=IFERROR(AVERAGEIF(\'Sentence Layer\'!$P$5:$P${last},A{r},\'Sentence Layer\'!$S$5:$S${last}),"—")'; f.cell(r,8).value=", ".join(x for x,_ in Counter(a.get("target") for a in by.values() if a.get("fruit")==f.cell(r,1).value and a.get("target")).most_common(3))
    for sn,items,ks in [("Counterfeit Fruit",v["counterfeits"],["passage_id","claimed_fruit","missing_companion","produced_result","harmed_target","notes"]),("Reader Impact",v["reader_impacts"],["reader_type","likely_experience","fruit_received","harm_risked","passage_ids"]),("Repair Map",v["repairs"],["passage_id","passage","missing_fruit","repair","priority"])]:
        ws=w[sn]; clear(ws,5,len(ks),100)
        for r,x in enumerate(items,5):
            for c,k in enumerate(ks,1): ws.cell(r,c).value=", ".join(map(str,x.get(k))) if isinstance(x.get(k),list) else x.get(k)
    rm=w["Relationship Matrix"]; clear(rm,5,30,100)
    actors=[]; targets=[]
    for x in v["relationships"]:
        if x.get("actor") and x["actor"] not in actors: actors.append(x["actor"])
        if x.get("target") and x["target"] not in targets: targets.append(x["target"])
    for c,target in enumerate(targets,2): rm.cell(4,c).value=target
    for r,actor in enumerate(actors,5):
        rm.cell(r,1).value=actor
        for x in v["relationships"]:
            if x.get("actor")==actor and x.get("target") in targets:
                rm.cell(r,targets.index(x["target"])+2).value=excel(x.get("treatment") or x.get("summary") or x.get("result"))
    su=w["Paper Summary"]
    for r,k in {5:"relational_fruit_thesis",8:"greatest_strength",10:"greatest_danger",12:"most_important_asymmetry",14:"most_revealing_passage",16:"recommended_revision"}.items():su.cell(r,1).value=v["summary"].get(k)
    w["Data Sources"].append([f"Sentence Layer!A5:S{last}",str(src),title,"Local + DeepSeek","Fruits relational evaluation",receipt["evaluation"]["model"],receipt["completed_at"],"David",len(rows),"Extracted and classified",f"source_sha256={receipt['source_sha256']}; run_id={receipt['run_id']}"])
    w.calculation.fullCalcOnLoad=True; w.calculation.forceFullCalc=True; w.calculation.calcMode="auto"; w.save(out)

def page(title,src,v,prose,receipt):
    ct={f:Counter() for f in FRUITS}
    for x in v["sentence_assessments"]:ct[x["fruit"]][x["result"]]+=1
    def ps(xs):return "".join(f"<p>{html.escape(x)}</p>" for x in xs)
    quotes="".join(f"<article><blockquote>{html.escape(str(x.get('exact_quotation','')))}</blockquote><p><b>{html.escape(str(x.get('fruit','')).title())} · {html.escape(str(x.get('result','')))}</b> — {html.escape(str(x.get('explanation','')))}</p></article>" for x in v["passages"][:10])
    labels=json.dumps([x.title() for x in FRUITS]); pos=json.dumps([ct[x]["FRUIT"] for x in FRUITS]); neg=json.dumps([ct[x]["ANTI_FRUIT"]+ct[x]["COUNTERFEIT"] for x in FRUITS])
    return f'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(prose["title"])}</title><script src="https://cdn.jsdelivr.net/npm/chart.js"></script><style>:root{{--ink:#17212b;--paper:#f7f3ea;--gold:#c7963f;--green:#315b4a;--red:#8e3d37}}*{{box-sizing:border-box}}body{{margin:0;background:var(--paper);color:var(--ink);font:18px/1.65 Georgia,serif}}header,footer{{background:#17212b;color:#ddd;padding:18px 5vw;font:14px system-ui}}header b{{color:#f1c66d}}main{{max-width:1100px;margin:auto;padding:6vw}}h1{{font-size:clamp(42px,7vw,88px);line-height:.96}}h2{{font-size:38px;margin-top:2em}}.dek{{font-size:24px;color:#445}}.thesis{{border-left:6px solid var(--gold);padding:22px;background:white;font-size:24px;margin:40px 0}}.grid{{display:grid;grid-template-columns:1.2fr .8fr;gap:32px}}.card{{background:white;padding:24px;border:1px solid #ddd7ca;border-radius:10px}}article{{padding:22px 0;border-bottom:1px solid #d8d0c2}}blockquote{{font-size:22px;border-left:4px solid var(--green);padding-left:18px}}@media(max-width:800px){{.grid{{grid-template-columns:1fr}}}}</style></head><body><header><b>GENERATED FRUITS ANALYSIS</b> · DeepSeek {html.escape(receipt["evaluation"]["model"])} · run {receipt["run_id"]} · source SHA-256 {receipt["source_sha256"]} · {receipt["completed_at"]}</header><main><p>FRUITS OF THE SPIRIT · RELATIONAL READING</p><h1>{html.escape(prose["title"])}</h1><p class="dek">{html.escape(prose["dek"])}</p><div class="thesis">{html.escape(v["summary"]["relational_fruit_thesis"])}</div><section class="grid"><div><h2>{html.escape(prose["hook_heading"])}</h2>{ps(prose["hook_paragraphs"])}</div><div class="card"><canvas id="chart"></canvas><small>Classified evidence sentences, not a morality score.</small></div></section><h2>{html.escape(prose["mirror_heading"])}</h2>{ps(prose["mirror_paragraphs"])}<h2>{html.escape(prose["story_heading"])}</h2>{ps(prose["story_paragraphs"])}<h2>Evidence passages</h2>{quotes}<div class="thesis">{html.escape(prose["closing"])}</div></main><footer>{html.escape(src.name)} · Observable text only; no claim about hidden motives or spiritual status.</footer><script>new Chart(document.getElementById('chart'),{{type:'bar',data:{{labels:{labels},datasets:[{{label:'Fruit',data:{pos},backgroundColor:'#315b4a'}},{{label:'Anti / counterfeit',data:{neg},backgroundColor:'#8e3d37'}}]}},options:{{indexAxis:'y',scales:{{x:{{beginAtZero:true,ticks:{{precision:0}}}}}}}}}});</script></body></html>'''

def markdown(src,v,prose,receipt):
    evidence="\n\n".join(f"> {x.get('exact_quotation','')}\n\n**{str(x.get('fruit','')).title()} · {x.get('result','')}** — {x.get('explanation','')}" for x in v["passages"][:10])
    sections=[]
    for key in ("hook","mirror","story"):
        sections.append(f"## {prose[key+'_heading']}\n\n"+"\n\n".join(prose[key+"_paragraphs"]))
    return f'''<!-- GENERATED FRUITS ANALYSIS | model={receipt["evaluation"]["model"]} | run_id={receipt["run_id"]} | source_sha256={receipt["source_sha256"]} | completed_at={receipt["completed_at"]} -->
# {prose["title"]}

{prose["dek"]}

## Relational fruit thesis

{v["summary"]["relational_fruit_thesis"]}

{chr(10).join(sections)}

## Evidence passages

{evidence}

## Closing

{prose["closing"]}

_Observable text only; no claim about hidden motives or spiritual status. Source: {src.name}._
'''

def run(src,out,cfg_path):
    cfg=json.loads(cfg_path.read_text()); title,rows=sentences(src); rid=str(uuid.uuid4()); out=out/src.stem; out.mkdir(parents=True,exist_ok=True)
    ev_path=out/f"{src.stem}.fruits.json"; prose_path=out/f"{src.stem}.page.json"
    if ev_path.exists() and prose_path.exists():
        ev=json.loads(ev_path.read_text(encoding="utf-8")); prose=json.loads(prose_path.read_text(encoding="utf-8")); em={"model":cfg["model"],"usage":None,"finish_reason":"resumed"}; wm={"model":cfg["model"],"usage":None,"finish_reason":"resumed"}
    else:
        ev_prompt=(cfg_path.parent.parent/"PROMPTS"/"evaluate.txt").read_text(); packet={"title":title,"source":src.name,"sentences":rows}
        ev,em=call(cfg,ev_prompt,"Return json for this packet:\n"+json.dumps(packet,ensure_ascii=False,separators=(",",":")),cfg["max_tokens_evaluate"]); validate(ev,rows)
        wr_prompt=(cfg_path.parent.parent/"PROMPTS"/"write_page.txt").read_text(); prose,wm=call(cfg,wr_prompt,"Return json prose from validated evidence:\n"+json.dumps(ev,ensure_ascii=False,separators=(",",":")),cfg["max_tokens_write"])
    receipt={"status":"SUCCEEDED","run_id":rid,"paper_id":src.stem,"source":str(src),"source_sha256":digest(src),"completed_at":now(),"evaluation":em,"writing":wm,"sentence_count":len(rows),"classified_sentence_count":len(ev["sentence_assessments"]),"passage_count":len(ev["passages"])}
    template=(cfg_path.parent/cfg["template"]).resolve(); receipt["template_sha256"]=digest(template)
    dump(out/f"{src.stem}.sentences.json",rows); dump(out/f"{src.stem}.fruits.json",ev); dump(out/f"{src.stem}.page.json",prose)
    workbook(template,out/f"{src.stem}.fruits.xlsx",title,src,rows,ev,receipt)
    (out/f"{src.stem}.fruits.html").write_text(page(title,src,ev,prose,receipt),encoding="utf-8")
    (out/f"{src.stem}.fruits.md").write_text(markdown(src,ev,prose,receipt),encoding="utf-8")
    dump(out/f"{src.stem}.run.json",receipt); return out

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("source",type=pathlib.Path); ap.add_argument("--output",type=pathlib.Path,required=True); ap.add_argument("--config",type=pathlib.Path,required=True); a=ap.parse_args()
    try: print(json.dumps({"status":"SUCCEEDED","output":str(run(a.source.resolve(),a.output.resolve(),a.config.resolve()))},indent=2)); return 0
    except Exception as e: print(json.dumps({"status":"FAILED","error":str(e)},indent=2)); return 1
if __name__=="__main__":raise SystemExit(main())
