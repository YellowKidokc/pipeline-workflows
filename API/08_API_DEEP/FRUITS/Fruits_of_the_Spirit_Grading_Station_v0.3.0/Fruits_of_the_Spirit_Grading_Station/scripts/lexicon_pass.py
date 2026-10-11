"""lexicon_pass.py -- Stage 1 of the Fruits of the Spirit grading pipeline.

Sentence-by-sentence, word-by-word lexicon scoring. No API calls.

Inputs:
  --paper   Path to a Markdown/text paper (UTF-8).
  Lexicon workbooks (read-only):
    * Fruits Template (4).xlsx  (Fruit / Anti-Fruit / Grounding / Contradiction /
      Jargon / Propaganda / Archetypes sheets, Weights sheet)
    * LEXICON_V2.xlsx           (SEMANTIC_BUCKETS, HEDGE_TERMS, ABSOLUTE_TERMS,
      FRUITS / FRUITS_LEX / FRUIT_BLEND for fruit-axis attribution,
      ANTI_FRUITS / ANTI_FRUIT_MAP for anti-fruit attribution)
      The AAVE-Ebonics sheet is skipped entirely: it contains broken
      #NAME? formulas and is recorded as skipped-broken in the manifest.

Lexicons are extracted once to scripts/lexicon_cache.json and rebuilt when the
cache is missing or older than either workbook.

Output:
  grading_out/<paper>.lexicon.json -- per-sentence records with char offsets,
  word counts, per-category hits/matched terms/weighted scores, fruit-axis and
  anti-fruit-axis attribution, claim-strength counts, hedge/absolute counts.
"""

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

STATION_DIR = Path(__file__).resolve().parent.parent
OUT_DIR = STATION_DIR / "grading_out"
CACHE_PATH = Path(__file__).resolve().parent / "lexicon_cache.json"
WEIGHTS_PATH = Path(__file__).resolve().parent / "weights.json"

TEMPLATE_XLSX = Path(
    r"D:\GitHub\pipeline-workflows\API_ALL\05_API_DEEP_STATIONS\FRUITS\Fruits Template (4).xlsx"
)
LEXICON_V2_XLSX = Path(
    r"D:\GitHub\pipeline-workflows\API_ALL\05_API_DEEP_STATIONS\FRUITS\LEXICON_V2.xlsx"
)

FRUIT_AXES = [
    "love", "joy", "peace", "patience", "kindness",
    "goodness", "faithfulness", "gentleness", "self_control",
]

# Category weights come from scripts/weights.json (fallback) and the
# workbook's Weights sheet (authoritative when present).

CLAIM_SUBBUCKET_MAP = {
    "weak_inference": "weak",
    "moderate_claim": "moderate",
    "strong_claim": "strong",
    "forbidden_unless_formal": "forbidden",
}


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean(v):
    return str(v).strip() if v is not None else ""


def is_broken(v):
    return "#NAME?" in clean(v)


def load_weights():
    """Load scripts/weights.json (the tuning file). Fail loudly if missing."""
    if not WEIGHTS_PATH.exists():
        sys.exit(f"ERROR: tuning file not found: {WEIGHTS_PATH}\n"
                 "weights.json holds the category weights and tunables; restore it.")
    with open(WEIGHTS_PATH, encoding="utf-8") as f:
        return json.load(f)


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_axis(name):
    return clean(name).lower().replace("-", "_").replace(" ", "_")


# ---------------------------------------------------------------------------
# Workbook extraction
# ---------------------------------------------------------------------------

def _read_sheet_rows(wb, name):
    if name not in wb.sheetnames:
        return None
    return [tuple(r) for r in wb[name].iter_rows(values_only=True)]


def _find_header(rows, first_cell_names):
    """Return (header_row_idx, header_cells_lower) for the row whose first
    cell matches one of first_cell_names."""
    for i, row in enumerate(rows):
        if row and clean(row[0]).lower() in first_cell_names:
            return i, [clean(c).lower() for c in row]
    return None, None


def _extract_template_terms(rows, active_required=True):
    """Extract terms from a Fruit/Anti-Fruit/Grounding/Contradiction/Jargon/
    Propaganda sheet: term in col A, Active flag must be 'Y'."""
    hi, header = _find_header(rows, {"term", "term/phrase"})
    if hi is None:
        return []
    try:
        active_col = header.index("active")
    except ValueError:
        active_col = None
    sub_col = None
    for name in ("subcategory", "type", "tier", "pattern type"):
        if name in header:
            sub_col = header.index(name)
            break
    terms = []
    for row in rows[hi + 1:]:
        if not row:
            continue
        term = clean(row[0])
        if not term or is_broken(term):
            continue
        if active_required and active_col is not None:
            if clean(row[active_col]).upper() != "Y":
                continue
        sub = clean(row[sub_col]) if sub_col is not None and sub_col < len(row) else ""
        terms.append({"term": term, "subcategory": sub})
    return terms


def extract_template(path, default_weights):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_map = {
            "fruit": "Fruit",
            "anti_fruit": "Anti-Fruit",
            "grounding": "Grounding",
            "contradiction": "Contradiction",
            "jargon": "Jargon",
            "propaganda": "Propaganda",
        }
        categories = {}
        for key, sheet in sheet_map.items():
            rows = _read_sheet_rows(wb, sheet)
            categories[key] = _extract_template_terms(rows) if rows else []

        # Archetypes: key markers (comma-separated) become informational terms.
        archetypes = []
        rows = _read_sheet_rows(wb, "Archetypes")
        if rows:
            hi, header = _find_header(rows, {"archetype/role"})
            if hi is not None:
                marker_col = next(
                    (i for i, h in enumerate(header) if h.startswith("key markers")), None
                )
                for row in rows[hi + 1:]:
                    if not row:
                        continue
                    name = clean(row[0])
                    if not name or is_broken(name):
                        continue
                    if marker_col is not None and marker_col < len(row):
                        for marker in clean(row[marker_col]).split(","):
                            marker = marker.strip()
                            if marker and not is_broken(marker):
                                archetypes.append({"term": marker, "subcategory": name})
        categories["archetype"] = archetypes

        # Weights sheet (authoritative); weights.json category_weights is the fallback.
        weights = dict(default_weights)
        rows = _read_sheet_rows(wb, "Weights")
        if rows:
            hi, header = _find_header(rows, {"category"})
            if hi is not None:
                wcol = header.index("weight") if "weight" in header else 1
                name_map = {"academic_jargon": "jargon"}
                for row in rows[hi + 1:]:
                    if not row:
                        continue
                    cat = clean(row[0]).lower()
                    cat = name_map.get(cat, cat)
                    if cat in weights and not is_broken(row[wcol]):
                        try:
                            weights[cat] = float(row[wcol])
                        except (TypeError, ValueError):
                            pass
        return categories, weights
    finally:
        wb.close()


def extract_v2(path):
    """Extract claim-strength buckets, hedges, absolutes, and fruit/anti-fruit
    axis maps from LEXICON_V2.xlsx. The AAVE-Ebonics sheet is skipped
    (broken #NAME? formulas)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        claim_strength = {k: [] for k in ("weak", "moderate", "strong", "forbidden")}
        hedges, absolutes = [], []
        fruit_axes = {}        # term -> set of fruit axes
        anti_fruit_axes = {}   # term -> set of axis labels

        rows = _read_sheet_rows(wb, "SEMANTIC_BUCKETS")
        if rows:
            header = [clean(c).lower() for c in rows[0]]
            idx = {name: header.index(name) for name in
                   ("term", "bucket", "subbucket") if name in header}
            for row in rows[1:]:
                if not row or any(is_broken(c) for c in row):
                    continue
                term = clean(row[idx["term"]])
                bucket = clean(row[idx["bucket"]])
                sub = clean(row[idx["subbucket"]])
                if not term:
                    continue
                if bucket == "CLAIM_STRENGTH" and sub in CLAIM_SUBBUCKET_MAP:
                    claim_strength[CLAIM_SUBBUCKET_MAP[sub]].append(term)

        for sheet, target in (("HEDGE_TERMS", hedges), ("ABSOLUTE_TERMS", absolutes)):
            rows = _read_sheet_rows(wb, sheet)
            if rows:
                header = [clean(c).lower() for c in rows[0]]
                tcol = header.index("term") if "term" in header else 3
                for row in rows[1:]:
                    if not row or tcol >= len(row):
                        continue
                    term = clean(row[tcol])
                    if term and not is_broken(term):
                        target.append(term)

        def add_axis(mapping, term, axis):
            term, axis = clean(term), norm_axis(axis)
            if term and axis and not is_broken(term) and not is_broken(axis):
                mapping.setdefault(term.lower(), set()).add(axis)

        # key = fruit axis, value = term
        for sheet in ("FRUITS", "FRUITS_LEX"):
            rows = _read_sheet_rows(wb, sheet)
            if rows:
                header = [clean(c).lower() for c in rows[0]]
                if "key" in header and "value" in header:
                    kc, vc = header.index("key"), header.index("value")
                    for row in rows[1:]:
                        if row and len(row) > max(kc, vc):
                            add_axis(fruit_axes, row[vc], row[kc])

        # FRUIT_BLEND: subkey positive -> fruit signal, negative -> anti signal.
        rows = _read_sheet_rows(wb, "FRUIT_BLEND")
        if rows:
            header = [clean(c).lower() for c in rows[0]]
            if {"key", "subkey", "value"} <= set(header):
                kc, sc, vc = header.index("key"), header.index("subkey"), header.index("value")
                for row in rows[1:]:
                    if not row or len(row) <= max(kc, sc, vc):
                        continue
                    sub = clean(row[sc]).lower()
                    if sub == "positive":
                        add_axis(fruit_axes, row[vc], row[kc])
                    elif sub == "negative":
                        add_axis(anti_fruit_axes, row[vc], row[kc])

        # ANTI_FRUITS: key = axis, value = anti-term.
        rows = _read_sheet_rows(wb, "ANTI_FRUITS")
        if rows:
            header = [clean(c).lower() for c in rows[0]]
            if "key" in header and "value" in header:
                kc, vc = header.index("key"), header.index("value")
                for row in rows[1:]:
                    if row and len(row) > max(kc, vc):
                        add_axis(anti_fruit_axes, row[vc], row[kc])

        # ANTI_FRUIT_MAP: key = anti-fruit label, value = marker term.
        rows = _read_sheet_rows(wb, "ANTI_FRUIT_MAP")
        if rows:
            header = [clean(c).lower() for c in rows[0]]
            if "key" in header and "value" in header:
                kc, vc = header.index("key"), header.index("value")
                for row in rows[1:]:
                    if row and len(row) > max(kc, vc):
                        add_axis(anti_fruit_axes, row[vc], row[kc])

        return {
            "claim_strength": {k: sorted(set(v)) for k, v in claim_strength.items()},
            "hedges": sorted(set(hedges)),
            "absolutes": sorted(set(absolutes)),
            "fruit_axes": {t: sorted(a) for t, a in sorted(fruit_axes.items())},
            "anti_fruit_axes": {t: sorted(a) for t, a in sorted(anti_fruit_axes.items())},
        }
    finally:
        wb.close()


def build_cache(template_path, v2_path, default_weights, force=False):
    if not force and CACHE_PATH.exists():
        cache_mtime = CACHE_PATH.stat().st_mtime
        if all(cache_mtime >= p.stat().st_mtime for p in (template_path, v2_path)):
            with open(CACHE_PATH, encoding="utf-8") as f:
                return json.load(f), False
    categories, weights = extract_template(template_path, default_weights)
    v2 = extract_v2(v2_path)
    cache = {
        "built_at": now_iso(),
        "sources": {
            str(template_path): {"mtime": template_path.stat().st_mtime},
            str(v2_path): {"mtime": v2_path.stat().st_mtime},
        },
        "skipped_sources": [
            {
                "sheet": "AAVE-Ebonics (LEXICON_V2.xlsx)",
                "reason": "skipped-broken: contains #NAME? broken formulas",
            }
        ],
        "weights": weights,
        "categories": categories,
        **v2,
    }
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)
    return cache, True


# ---------------------------------------------------------------------------
# Sentence splitting and matching
# ---------------------------------------------------------------------------

SENT_END_RE = re.compile(r"[.!?]+[\"”’'\)\]]*(?=\s|$)")
PARA_BREAK_RE = re.compile(r"\n\s*\n")
# Transcript-style lines: "[00:12] ..." — each timestamped line is a boundary.
TIMESTAMP_LINE_RE = re.compile(r"\n(?=\[\d{1,2}:\d{2}\])")
TS_PREFIX_RE = re.compile(r"\[\d{1,2}:\d{2}\]\s*")
WORD_RE = re.compile(r"[A-Za-z0-9]+(?:['’\-][A-Za-z0-9]+)*")


def split_sentences(text):
    """Split text into sentences, preserving char offsets. A sentence ends at
    terminal punctuation followed by whitespace/EOS, at a blank-line break, or
    at a transcript timestamp line boundary. Returns list of
    (start, end, sentence_text)."""
    boundaries = []  # end offsets (exclusive)
    for m in SENT_END_RE.finditer(text):
        boundaries.append(m.end())
    for m in PARA_BREAK_RE.finditer(text):
        boundaries.append(m.start())
    for m in TIMESTAMP_LINE_RE.finditer(text):
        boundaries.append(m.start())
    boundaries = sorted(set(boundaries))

    sentences = []
    prev = 0
    for b in boundaries:
        chunk = text[prev:b]
        if chunk.strip():
            start = prev + (len(chunk) - len(chunk.lstrip()))
            end = b - (len(chunk) - len(chunk.rstrip()))
            # strip a leading [mm:ss] timestamp marker from the sentence text
            ts = TS_PREFIX_RE.match(text[start:end])
            if ts:
                start += ts.end()
            if start < end:
                sentences.append((start, end, text[start:end]))
        prev = b
    tail = text[prev:]
    if tail.strip():
        start = prev + (len(tail) - len(tail.lstrip()))
        end = len(text) - (len(tail) - len(tail.rstrip()))
        ts = TS_PREFIX_RE.match(text[start:end])
        if ts:
            start += ts.end()
        if start < end:
            sentences.append((start, end, text[start:end]))
    return sentences


def paragraph_index(text, para_breaks, offset):
    idx = 0
    for b in para_breaks:
        if b < offset:
            idx += 1
        else:
            break
    return idx


def compile_matcher(terms):
    pats = sorted({t.lower() for t in terms if t}, key=len, reverse=True)
    if not pats:
        return None
    return re.compile(
        r"(?<!\w)(" + "|".join(re.escape(p) for p in pats) + r")(?!\w)",
        re.IGNORECASE,
    )


def match_terms(regex, text):
    if regex is None:
        return []
    return [m.group(0).lower() for m in regex.finditer(text)]


def main():
    ap = argparse.ArgumentParser(description="Stage 1: lexicon pass (no API).")
    ap.add_argument("--paper", required=True, help="Path to the paper (md/txt).")
    ap.add_argument("--template", default=str(TEMPLATE_XLSX))
    ap.add_argument("--lexicon-v2", default=str(LEXICON_V2_XLSX))
    ap.add_argument("--force-rebuild", action="store_true",
                    help="Rebuild lexicon_cache.json even if fresh.")
    args = ap.parse_args()

    paper_path = Path(args.paper)
    template_path = Path(args.template)
    v2_path = Path(args.lexicon_v2)

    tuning = load_weights()
    weights_sha = sha256_file(WEIGHTS_PATH)
    cache, rebuilt = build_cache(template_path, v2_path,
                                 tuning["category_weights"], force=args.force_rebuild)
    print(f"[lexicon] cache {'rebuilt' if rebuilt else 'loaded'}: {CACHE_PATH}")

    text = paper_path.read_text(encoding="utf-8")

    weights = cache["weights"]
    cat_terms = {c: [t["term"] for t in tl] for c, tl in cache["categories"].items()}
    cat_regex = {c: compile_matcher(t) for c, t in cat_terms.items()}
    cs_regex = {k: compile_matcher(v) for k, v in cache["claim_strength"].items()}
    hedge_re = compile_matcher(cache["hedges"])
    abs_re = compile_matcher(cache["absolutes"])
    fruit_axes_map = cache["fruit_axes"]
    anti_axes_map = cache["anti_fruit_axes"]

    sent_tuples = split_sentences(text)
    para_starts = [m.start() for m in PARA_BREAK_RE.finditer(text)]

    sentences = []
    for i, (start, end, stext) in enumerate(sent_tuples):
        words = WORD_RE.findall(stext)
        cats = {}
        for cat, regex in cat_regex.items():
            hits = match_terms(regex, stext)
            uniq = sorted(set(hits))
            cats[cat] = {
                "hits": len(hits),
                "unique_terms": len(uniq),
                "matched_terms": uniq,
                "weighted_score": round(len(hits) * weights.get(cat, 0.0), 4),
            }
        fruit_axes = {}
        for term in cats["fruit"]["matched_terms"]:
            for axis in fruit_axes_map.get(term, []):
                fruit_axes[axis] = fruit_axes.get(axis, 0) + 1
        anti_axes = {}
        for term in cats["anti_fruit"]["matched_terms"]:
            for axis in anti_axes_map.get(term, []):
                anti_axes[axis] = anti_axes.get(axis, 0) + 1
        sentences.append({
            "idx": i,
            "start": start,
            "end": end,
            "paragraph": paragraph_index(text, para_starts, start),
            "text": stext,
            "word_count": len(words),
            "categories": cats,
            "fruit_axes": fruit_axes,
            "anti_fruit_axes": anti_axes,
            "claim_strength": {k: len(match_terms(r, stext)) for k, r in cs_regex.items()},
            "hedges": len(match_terms(hedge_re, stext)),
            "absolutes": len(match_terms(abs_re, stext)),
        })

    totals = {
        cat: {
            "hits": sum(s["categories"][cat]["hits"] for s in sentences),
            "unique_terms": len({t for s in sentences for t in s["categories"][cat]["matched_terms"]}),
            "weighted_score": round(sum(s["categories"][cat]["weighted_score"] for s in sentences), 4),
        }
        for cat in cat_terms
    }
    totals["claim_strength"] = {
        k: sum(s["claim_strength"][k] for s in sentences)
        for k in ("weak", "moderate", "strong", "forbidden")
    }
    totals["hedges"] = sum(s["hedges"] for s in sentences)
    totals["absolutes"] = sum(s["absolutes"] for s in sentences)
    totals["fruit_axes"] = {}
    for s in sentences:
        for a, n in s["fruit_axes"].items():
            totals["fruit_axes"][a] = totals["fruit_axes"].get(a, 0) + n

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / f"{paper_path.stem}.lexicon.json"
    payload = {
        "paper": str(paper_path),
        "paper_name": paper_path.stem,
        "generated_at": now_iso(),
        "stage": "lexicon_pass_v1",
        "char_count": len(text),
        "sentence_count": len(sentences),
        "paragraph_count": (sentences[-1]["paragraph"] + 1) if sentences else 0,
        "weights": weights,
        "weights_file": str(WEIGHTS_PATH),
        "weights_sha256": weights_sha,
        "lexicon_cache": {"path": str(CACHE_PATH), "sha256": sha256_file(CACHE_PATH)},
        "skipped_sources": cache["skipped_sources"],
        "totals": totals,
        "sentences": sentences,
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)

    print(f"[lexicon] paper: {paper_path.name}")
    print(f"[lexicon] sentences: {len(sentences)}  words: {sum(s['word_count'] for s in sentences)}")
    for cat, t in totals.items():
        if isinstance(t, dict) and "hits" in t:
            print(f"[lexicon]   {cat}: hits={t['hits']} unique={t['unique_terms']} weighted={t['weighted_score']}")
        else:
            print(f"[lexicon]   {cat}: {t}")
    print(f"[lexicon] wrote {out_path}")


if __name__ == "__main__":
    main()
