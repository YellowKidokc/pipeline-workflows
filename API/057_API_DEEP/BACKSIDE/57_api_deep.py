"""57_API_DEEP (was API_HOME station 48): Fruits, Axioms, Atoms, Lean, Stories, Master Equation and Coherence as ONE chain per paper.

David, 2026-09-26: "There's no way that I don't run one and not run them all." One run per paper writes
one JSON per station, a combined API_DEEP.json, a per-sentence Fruits workbook, and API_DEEP.md with the
stations in page order: Fruits, Axioms, Atoms, Lean, Stories, Master Equation, Coherence.

The prompts are the original station prompts, copied verbatim into prompts/ (and the tested prompt builders
into lib/). Edit a prompt file to change the call; its hash goes into every receipt. `--copy` builds the same
prompts without calling any API and puts one on the clipboard for pasting into a chat.

Call plan (all calls JSON mode, max_tokens 8192, whole paper in every call, sentences numbered locally):
  phase 1 (parallel)  atoms | fruits sentence ranges | master equation Part A x2 (temperature 0.7)
  phase 2 (parallel)  fruits verdict (+ sentence summary) | axiom nodes | lean4 | stories | coherence | ME Part B
                      (phase 2 prompts receive the atom list, so atom ids line up across stations)
"""
from __future__ import annotations
import argparse, hashlib, json, math, re, subprocess, sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
FRONT = HERE.parent                                      # 057_API_DEEP: INBOX/ and OUTBOX/ sit next to BACKSIDE/
sys.path.insert(0, str(FRONT.parent / "_system")); sys.path.insert(0, str(HERE / "lib"))
from engine import llm                                   # noqa: E402
from engine.focus import append, compose                 # noqa: E402
from engine.paths import API_HOME                        # noqa: E402
from engine import pick                                  # noqa: E402  (_PICK.md: only ticked notes go to the API)
from engine.publish import publish_on_note               # noqa: E402  (answers go onto the source note)
import fruits_grade as FG                                # noqa: E402  (copied from 05_API_DEEP_STATIONS/FRUITS)
from atom_prompt import build_atom_prompt                # noqa: E402  (copied from 05_API_DEEP_STATIONS/ATOMS)
from axiom_prompt import build_axiom_prompt              # noqa: E402  (copied from 05_API_DEEP_STATIONS/AXIOM_NODES)
import love_truth as LT                                  # noqa: E402  (Truth Engine v2.0 lexicons + character profiles, local)
import report_html                                       # noqa: E402  (API_DEEP.html, matrix house style)
import runs_index                                        # noqa: E402  (OUTBOX/INDEX.html)

LABEL = "57_API_DEEP"
OUTBOX = FRONT / "OUTBOX"
INBOX = FRONT / "INBOX"
PROMPTS = HERE / "prompts"
FRUITS = ["love", "joy", "peace", "patience", "kindness", "goodness", "faithfulness", "gentleness", "self_control"]
PAGE_ORDER = ["fruits", "axiom_nodes", "atoms", "lean4", "stories", "master_equation", "coherence"]
RANGE_SIZE = 70
MAX_TOKENS = 8192
JSON_ONLY = "\n\nReturn ONLY one JSON object. No markdown fences, no commentary."


def now() -> str: return datetime.now(timezone.utc).isoformat()
def sha(text: str) -> str: return hashlib.sha256(text.encode("utf-8")).hexdigest()
def read(name: str) -> str: return (PROMPTS / name).read_text(encoding="utf-8-sig")


# ---------------------------------------------------------------- source prep (local, deterministic)

def split_front_matter(text: str) -> tuple[str, str]:
    m = re.match(r"\ufeff?\s*(?:```ya?ml\s*)?---\s*\n(.*?)\n---\s*\n(?:```\s*\n)?", text, re.S)
    return (m.group(1), text[m.end():]) if m else ("", text)

SENT_END = re.compile(r"(?:(?<=[.!?])|(?<=[.!?][\"'”’)\]]))\s+(?=[\"'“‘(\[*_]?[A-Z0-9])")

def number_source(text: str) -> tuple[list[dict], str]:
    """Paragraphs P01.. and sentences S001..; headings, list items and table rows are sentences of their own."""
    sentences, blocks, in_code = [], [], False
    paras = re.split(r"\n\s*\n", text)
    p_no = 0
    for para in paras:
        lines = [x.rstrip() for x in para.strip().splitlines() if x.strip()]
        if not lines: continue
        units: list[str] = []
        for line in lines:
            if line.lstrip().startswith("```"): in_code = not in_code; continue
            if in_code: units.append(line.strip()); continue
            if re.match(r"\s*(#|[-*+] |\d+[.)] |\||>)", line) or len(lines) > 1 and not units and line.endswith(":"):
                units.append(line.strip())
            elif units and not re.match(r"\s*(#|[-*+] |\d+[.)] |\|)", units[-1]) and not units[-1].endswith(("|",)):
                units[-1] += " " + line.strip()
            else:
                units.append(line.strip())
        parts = [s.strip() for u in units for s in (SENT_END.split(u) if not u.startswith(("|", "#")) else [u]) if s.strip()]
        parts = [s for s in parts if re.search(r"\w", s)]
        if not parts: continue
        p_no += 1; pid = f"P{p_no:02d}"; out = [f"[{pid}]"]
        for s in parts:
            sid = f"S{len(sentences)+1:03d}"; sentences.append({"id": sid, "para": pid, "text": s}); out.append(f"{sid} {s}")
        blocks.append("\n".join(out))
    return sentences, "\n\n".join(blocks)


# ---------------------------------------------------------------- JSON calls

def parse_json(text: str):
    try: return FG.extract_json_object(text), None
    except Exception as exc: return None, f"{type(exc).__name__}: {exc}"

def call_json(name: str, system: str, user: str, a, temperature=0.1, validate=None) -> dict:
    """One station call. Retries once when the reply is truncated, not JSON, or fails validation,
    telling the model exactly what was wrong. Returns {parsed, receipt, error}."""
    messages = ([{"role": "system", "content": system}] if system else []) + [{"role": "user", "content": user}]
    receipts, parsed, error = [], None, None
    for attempt in (1, 2):
        r = llm.call(messages, provider=a.provider, model=a.model, json_mode=True, max_tokens=MAX_TOKENS, temperature=temperature, station=LABEL)
        receipts.append(llm.receipt(r, part=name, attempt=attempt))
        if r.error and not r.error.startswith("truncated"):
            error = r.error; break
        if r.error:  # truncated: ask again, once, for a compact answer
            error = r.error
            last = dict(messages[-1]); last["content"] += ("\n\nIMPORTANT: a previous answer hit the output limit and was cut off. "
                                                           "Be far more compact: short strings, only the most important items, same schema.")
            messages = messages[:-1] + [last]
            continue
        parsed, error = parse_json(r.text)
        if parsed is not None and validate:
            problems = validate(parsed)
            if problems: error = "validation: " + "; ".join(problems[:12])
            else: error = None
        if error is None: break
        messages = messages + [{"role": "assistant", "content": r.text}, {"role": "user", "content": f"That reply has problems: {error}. Return the corrected complete JSON object."}]
    return {"name": name, "parsed": parsed, "error": error, "receipts": receipts,
            "prompt_sha256": sha(system + "\n" + user), "tokens": sum(x["tokens"] for x in receipts)}


# ---------------------------------------------------------------- prompt builders

def source_block(meta: str, numbered: str, title: str) -> str:
    head = f"PAPER TITLE: {title}\n" + (f"PAPER METADATA (front matter):\n{meta}\n" if meta else "")
    if re.search(r"^type:\s*youtube", meta, re.M):
        head += ("SOURCE KIND: an auto-captioned YouTube transcript of spoken conversation (often an interview, so more than one speaker). "
                 "Treat it as the \"paper\": judge what the speakers say and do. Caption errors and missing speaker labels are not the speakers' faults.\n")
    return head + "\nSOURCE (paragraphs [P01].., sentences S001..; cite these ids wherever a schema asks for a span or location):\n" + numbered

def atoms_brief(atoms: dict | None) -> str:
    if not atoms or not atoms.get("atoms"): return "EXTRACTED ATOMS: none available for this run."
    rows = [f"{a['atom_id']} [{a.get('nodeType','?')}] {a.get('name','')}: {a.get('statementPlain') or a.get('statementTechnical') or ''}"[:300] for a in atoms["atoms"]]
    return "EXTRACTED ATOMS (use these atom ids wherever the schema asks for atom_uuid / atom_uuids / claim_id):\n" + "\n".join(rows)

def build_calls_phase1(ctx, a) -> list[tuple]:
    src, focus = ctx["src"], ctx["focus"]
    calls = [("atoms", "", append(build_atom_prompt(src) + "\nFor every atom also add \"source_sentences\": [sentence ids].", focus), 0.1, None)]
    sents = ctx["sentences"]
    for i in range(0, len(sents), RANGE_SIZE):
        chunk = sents[i:i+RANGE_SIZE]; rng = f"{chunk[0]['id']}-{chunk[-1]['id']}"
        ids = {s["id"] for s in chunk}
        calls.append((f"fruits_sentences_{i//RANGE_SIZE+1}", read("FRUITS_SENTENCES.md"),
                      f"TASK: score sentences {rng} ({len(chunk)} sentences).\n\n{src}", 0.1,
                      lambda p, ids=ids: sentence_problems(p, ids)))
    turns = ctx.get("turns") or []
    for i in range(0, len(turns), TURN_RANGE):
        chunk = turns[i:i+TURN_RANGE]; ids = {t["id"] for t in chunk}
        tmap = "\n".join(f"{t['id']} = {t['para']}" for t in chunk)
        calls.append((f"fruits_turns_{i//TURN_RANGE+1}", read("FRUITS_TURNS.md"),
                      f"TASK: score turns {chunk[0]['id']}-{chunk[-1]['id']} ({len(chunk)} turns). Turn map:\n{tmap}\n\n{src}", 0.1,
                      lambda p, ids=ids: turn_problems(p, ids)))
    me = read("MASTER_EQUATION_V2.md")
    task = ("Run PART A (the analog) for this paper. Return the Part A JSON object exactly as specified, with all 10 slots "
            "(G, M, E, S_eff, T, K, R, Q, F, C); quotes are exact source text followed by the sentence id in brackets. "
            "Also add \"equations_present\": true|false.")
    for k in (1, 2):
        calls.append((f"master_equation_a{k}", me, append(task + "\n\n" + src, focus), 0.7, None))
    return calls

def build_calls_phase2(ctx, a, atoms, summary) -> list[tuple]:
    src, focus, brief, uuid = ctx["src"], ctx["focus"], atoms_brief(atoms), ctx["paper_uuid"]
    rubric = json.loads(read("fruits_rubric_v0.3.0.json")); schema = json.loads(read("fruits_report_v0.3.0.schema.json"))
    fruits_user = FG.build_user_prompt(ctx["path"], src, rubric, schema, "gated_profile") + "\n\nSENTENCE-LEVEL SUMMARY (from the per-sentence pass):\n" + json.dumps(summary, ensure_ascii=False)
    station = lambda file, extra="": (read(file), append(f"paper_uuid: {uuid}\nrun_uuid: {ctx['run_uuid']}\n{extra}\n{brief}\n\n{src}{JSON_ONLY}", focus))
    calls = [
        ("fruits", read("FRUITS_SYSTEM_v0.3.0.md"), append(fruits_user, focus), 0.1, FG.validate_report),
        ("axiom_nodes", "", append(build_axiom_prompt(src) + "\n" + brief + "\nAlso return \"unmapped_claims\": [atom ids that map to no node].", focus), 0.1, None),
        ("lean4", *station("LEAN4.md", "The Lean corpus was NOT searched in this run: leave existing_matches empty, mark verification "
                            "controls NOT_RUN and verification_status NOT_ATTEMPTED or CANDIDATE. Give at most 5 formal candidates, "
                            "the strongest first; leave reproduction_receipt and corpus_links fields empty."), 0.1, None),
        ("stories", *station("STORIES.md"), 0.1, None),
        ("coherence", *station("COHERENCE_SCORE.md", "Score each dimension 0-10 and the overall 0-10; cite sentence ids in reasons."), 0.1, None),
    ]
    if ctx["has_math"]:
        calls.append(("master_equation_b", *station("MASTER_EQUATION_V1_PART_B.md", "This is PART B of the master equation station: the paper's own equations."), 0.1, None))
    return calls

TURN_MAX = 15      # sentences; a long monologue is cut into pieces so each scored unit stays readable
TURN_RANGE = 30    # turns per scoring call

def build_turns(sentences) -> list[dict]:
    """Speaker turns for spoken transcripts: auto-captions mark a speaker change with '>>'."""
    turns, cur = [], []
    for s in sentences:
        if cur and (">>" in s["text"] or len(cur) >= TURN_MAX):
            turns.append(cur); cur = []
        cur.append(s)
    if cur: turns.append(cur)
    return [{"id": f"T{i:03d}", "sids": [x["id"] for x in t], "para": f"{t[0]['id']}-{t[-1]['id']}",
             "text": " ".join(x["text"] for x in t)} for i, t in enumerate(turns, 1)]

def turn_problems(p, ids) -> list[str]:
    got = [t.get("id") for t in p.get("turns", []) if isinstance(t, dict)]
    bad = [t.get("id") for t in p.get("turns", []) if not (isinstance(t.get("v"), list) and len(t["v"]) == 9)]
    missing = sorted(ids - set(got))
    return ([f"missing turn ids {missing[:15]}"] if missing else []) + ([f"turns without 9 values: {bad[:10]}"] if bad else [])

def turns_aggregate(turns, parsed_parts) -> list[dict]:
    by_id = {}
    for p in parsed_parts:
        for t in (p or {}).get("turns", []):
            v = t.get("v")
            if isinstance(v, list) and len(v) == 9:
                by_id[t.get("id")] = {"v": [max(-2, min(2, int(x))) if isinstance(x, (int, float)) else 0 for x in v],
                                      "why": t.get("why") or {}, "speaker": t.get("speaker", "")}
    return [{"id": t["id"], "para": t["para"], "text": t["text"], **by_id.get(t["id"], {"v": None, "why": {}, "speaker": ""})} for t in turns]

def sentence_problems(p, ids) -> list[str]:
    got = [s.get("id") for s in p.get("sentences", []) if isinstance(s, dict)]
    bad = [s.get("id") for s in p.get("sentences", []) if not (isinstance(s.get("v"), list) and len(s["v"]) == 9)]
    missing = sorted(ids - set(got))
    return ([f"missing sentence ids {missing[:15]}"] if missing else []) + ([f"sentences without 9 values: {bad[:10]}"] if bad else [])


# ---------------------------------------------------------------- local aggregation

def fruits_aggregate(sentences, parsed_parts) -> tuple[list[dict], dict]:
    by_id = {}
    for p in parsed_parts:
        for s in (p or {}).get("sentences", []):
            v = s.get("v")
            if isinstance(v, list) and len(v) == 9:
                by_id[s.get("id")] = {"v": [max(-2, min(2, int(x))) if isinstance(x, (int, float)) else 0 for x in v], "why": s.get("why") or {}}
    rows = [{**s, **by_id.get(s["id"], {"v": None, "why": {}})} for s in sentences]
    scored = [r for r in rows if r["v"]]
    if not scored: return rows, {"scored": 0, "total": len(rows)}
    col = lambda i: [r["v"][i] for r in scored]
    love = col(0); roll = [sum(love[max(0, i-2):i+3]) / len(love[max(0, i-2):i+3]) for i in range(len(love))]
    turns = [scored[i]["id"] for i in range(1, len(roll)) if (roll[i-1] > 0) != (roll[i] > 0) and abs(roll[i] - roll[i-1]) > 0.1]
    def corr(x, y):
        mx, my = sum(x)/len(x), sum(y)/len(y); sx = math.sqrt(sum((a-mx)**2 for a in x)); sy = math.sqrt(sum((b-my)**2 for b in y))
        return round(sum((a-mx)*(b-my) for a, b in zip(x, y)) / (sx*sy), 3) if sx and sy else None
    total = lambda r: sum(r["v"])
    q = lambda r: {"id": r["id"], "para": r["para"], "sum": total(r), "love": r["v"][0], "text": r["text"][:240], "why": r["why"]}
    paras = {}
    for r in scored: paras.setdefault(r["para"], []).append(r)
    summary = {
        "scored": len(scored), "total": len(rows), "missing": [r["id"] for r in rows if not r["v"]],
        "fruit_means": {f: round(sum(col(i))/len(scored), 3) for i, f in enumerate(FRUITS)},
        "fruit_extremes": {f: {"plus2": sum(1 for x in col(i) if x == 2), "minus2": sum(1 for x in col(i) if x == -2),
                               "share_negative": round(sum(1 for x in col(i) if x < 0)/len(scored), 3)} for i, f in enumerate(FRUITS)},
        "track_love": {f: corr(love, col(i)) for i, f in enumerate(FRUITS) if i},
        "top_toward": [q(r) for r in sorted(scored, key=total, reverse=True)[:5] if total(r) > 0],
        "top_away": [q(r) for r in sorted(scored, key=total)[:5] if total(r) < 0],
        "love_turns": turns[:20],
        "counterfeit_candidates": [q(r) for r in scored if "counterfeit" in r["why"]][:15],
        "hidden_fruit": [q(r) for r in scored if "hidden" in r["why"]][:15],
        "paragraph_love": {p: round(sum(r["v"][0] for r in rs)/len(rs), 2) for p, rs in paras.items()},
        "paragraph_total": {p: round(sum(total(r) for r in rs)/len(rs), 2) for p, rs in paras.items()},
    }
    return rows, summary

def spark(values) -> str:
    bars = "▁▂▃▄▅▆▇█"
    return "".join(bars[min(7, max(0, int(round((v + 2) / 4 * 7))))] for v in values)

def me_compare(a1, a2) -> dict:
    if not a1 or not a2: return {"runs": [x for x in (a1, a2) if x], "note": "only one Part A run succeeded; nothing to compare"}
    slots1 = {s.get("slot"): s for s in a1.get("slots", [])}; slots2 = {s.get("slot"): s for s in a2.get("slots", [])}
    rows = []
    for slot in ["G", "M", "E", "S_eff", "T", "K", "R", "Q", "F", "C"]:
        x, y = slots1.get(slot, {}), slots2.get(slot, {})
        rows.append({"slot": slot, "fit_1": x.get("fit"), "fit_2": y.get("fit"), "status": "agreed" if x.get("fit") == y.get("fit") else "contested",
                     "plays_role": x.get("plays_role") or y.get("plays_role"), "quote": x.get("quote") or y.get("quote")})
    s1, s2 = a1.get("analog_strength"), a2.get("analog_strength")
    nums = [v for v in (s1, s2) if isinstance(v, (int, float))]
    pt1, pt2 = (a1.get("product_test") or {}).get("answer"), (a2.get("product_test") or {}).get("answer")
    return {"core_relation": a1.get("core_relation"), "core_relation_run2": a2.get("core_relation"),
            "product_test": {"run1": pt1, "run2": pt2, "status": "agreed" if pt1 == pt2 else "contested",
                             "quote": (a1.get("product_test") or {}).get("quote"), "reason": (a1.get("product_test") or {}).get("reason")},
            "slots": rows, "contested": [r["slot"] for r in rows if r["status"] == "contested"],
            "analog_strength": min(nums) if nums else None, "analog_strength_spread": f"{min(nums)}-{max(nums)}" if nums else None,
            "confidence": [a1.get("confidence"), a2.get("confidence")], "one_line": a1.get("one_line"),
            "breaks": list(dict.fromkeys([*(a1.get("breaks") or []), *(a2.get("breaks") or [])]))[:10] if all(isinstance(b, str) for b in [*(a1.get("breaks") or []), *(a2.get("breaks") or [])]) else [a1.get("breaks"), a2.get("breaks")],
            "extra_factors": [a1.get("extra_factors"), a2.get("extra_factors")],
            "equations_present": bool(a1.get("equations_present") or a2.get("equations_present"))}


# ---------------------------------------------------------------- page (Markdown)

def finding(x) -> str:
    if not isinstance(x, dict): return cell(x)
    ids = x.get("atoms") or x.get("atom_uuids") or ([x["atom_uuid"]] if x.get("atom_uuid") else [])
    head = " × ".join(map(str, ids)) if ids else ""
    body = x.get("description") or x.get("note") or ""
    term = f"**{x['term']}**: " if x.get("term") else ""
    tag = f" ({x['severity']})" if x.get("severity") else ""
    return cell(f"{term}{head + ' — ' if head else ''}{body}{tag}") if (body or term) else cell(json.dumps(x, ensure_ascii=False))

def cell(x) -> str:
    return str("" if x is None else x).replace("|", "\\|").replace("\n", " ")[:300]

def render_page(ctx, R) -> str:
    L = [f"# API DEEP — {ctx['title']}", "",
         f"Source `{ctx['path'].name}` · sha256 `{ctx['source_sha'][:16]}…` · {ctx['n_sentences']} sentences in {ctx['n_paras']} paragraphs · "
         f"{ctx['provider']}/{ctx['model']} · {ctx['finished']}", "",
         "Every result below is an AI proposal pending David's review.", ""]
    # 1 FRUITS
    s = R.get("fruits_sentences_summary") or {}
    L += ["## 1 · Fruits of Love and Truth", ""]
    lt = R.get("love_truth") or {}; pr = lt.get("profile") or {}; te = lt.get("truth_engine") or {}
    if pr.get("status") == "ASSESSED":
        L += [f"**{pr['quadrant']}** · love axis {pr['love_axis']} · truth axis {pr['truth_axis']} "
              f"({', '.join(f'{k} {round(v, 2)}' for k, v in pr['truth_parts'].items())})", "",
              "Level: " + (", ".join(pr["levels"]) or "none") + " · Shape: " +
              (", ".join(f"**{x['name']}** (gap {x['gap']})" for x in pr["shapes_matched"]) or f"no clear shape (closest {pr['shapes_ranked'][0]['name']}, gap {pr['shapes_ranked'][0]['gap']})"), "",
              f"Scored from {pr['sentences']} {pr.get('unit', 'sentences')} ({round(100 * pr['active_share'])}% engage a fruit). Net points per 100 {pr.get('unit', 'sentences')}:", "",
              "| " + " | ".join(f.replace('_', ' ') for f in pr["net_per_100"]) + " |", "|" + "---:|" * 9,
              "| " + " | ".join(f"{v:+}" for v in pr["net_per_100"].values()) + " |", "",
              "Paper verdict (rubric 0-4) read as a second opinion: " + str((lt.get("verdict_profile") or {}).get("quadrant", "—")), ""]
    if te:
        L += [f"Truth Engine v2.0 (lexicon, {te['words']} words): TRUTH {te['truth_score_raw']} → {te['truth_score_0_1']} · per 100 words: " +
              ", ".join(f"{c.replace('_', ' ')} {v}" for c, v in te["per_100_words"].items()), "",
              "Workbook characterizations: " + (", ".join(f"**{c['name']}**" for c in lt.get("characterizations", [])) or "none triggered"),
              "Top words: " + " · ".join(f"{c.replace('_', ' ')}: {', '.join(list(w)[:5])}" for c, w in te["top_words"].items() if w), ""]
    L += ["Profiles are V0 drafts (prompts/CHARACTER_PROFILES.md); a word hit is a trace, never character evidence.", "",
          "### Sentence by sentence", ""]
    if s.get("scored"):
        L += [f"Sentences scored: {s['scored']}/{s['total']}" + (f" (missing {', '.join(s['missing'][:10])})" if s["missing"] else ""), "",
              "**Love by paragraph** " + f"`{spark(list(s['paragraph_love'].values()))}`" + "  (▁ = −2 … █ = +2)", "",
              "| Fruit | mean (−2..+2) | +2 | −2 | share negative | tracks love (r) |", "|---|---:|---:|---:|---:|---:|"]
        for f in FRUITS:
            e = s["fruit_extremes"][f]; L.append(f"| {f.replace('_',' ')} | {s['fruit_means'][f]} | {e['plus2']} | {e['minus2']} | {e['share_negative']} | {s['track_love'].get(f) if s['track_love'].get(f) is not None else '—'} |")
        for head, key in (("Strongest moments toward the fruits", "top_toward"), ("Strongest moments away", "top_away"),
                          ("Counterfeit candidates (fruit words, anti-fruit work)", "counterfeit_candidates"), ("Hidden fruit (fruit work, no fruit words)", "hidden_fruit")):
            if s.get(key):
                L += ["", f"**{head}**", ""] + [f"- {x['id']} ({x['para']}, sum {x['sum']:+d}): “{x['text']}” — {'; '.join(f'{k}: {v}' for k, v in x['why'].items())}" for x in s[key]]
        if s.get("love_turns"): L += ["", "Love curve turns (5-sentence rolling mean crosses zero): " + ", ".join(s["love_turns"])]
    fr = R.get("fruits")
    if fr:
        md = FG.render_markdown(fr, ctx["path"], "").split("## Original source")[0]
        L += ["", "### Paper verdict (rubric v0.3.0)", ""] + ["##" + x if x.startswith("#") else x for x in md.splitlines()[1:]]
    # 2 AXIOMS
    ax = R.get("axiom_nodes") or {}
    L += ["", "## 2 · Axiom nodes", "", f"Primary mode: **{ax.get('primary_mode','—')}** · {cell(ax.get('summary'))}", "",
          "| Node | Name | Mode | Alignment | Confidence | Quote |", "|---|---|---|---|---|---|"]
    L += [f"| {cell(n.get('node_id'))} | {cell(n.get('name'))} | {cell(n.get('mode'))} | {cell(n.get('alignment'))} | {cell(n.get('confidence'))} | {cell(n.get('evidence_quote'))} |" for n in ax.get("axiom_nodes", [])]
    if ax.get("unmapped_claims"): L += ["", "Unmapped atoms: " + ", ".join(map(str, ax["unmapped_claims"]))]
    # 3 ATOMS
    at = R.get("atoms") or {}
    L += ["", "## 3 · Atoms", "", f"Domain: **{at.get('domainType','—')}** · {len(at.get('atoms', []))} atoms", "",
          "| Id | Type | Stage | Name | Plain statement | Falsified if | Sentences |", "|---|---|---|---|---|---|---|"]
    L += [f"| {x.get('atom_id')} | {cell(x.get('nodeType'))} | {cell(x.get('stage'))} | {cell(x.get('name'))} | {cell(x.get('statementPlain'))} | {cell(x.get('falsificationCondition'))} | {cell(', '.join(x.get('source_sentences') or []))} |" for x in at.get("atoms", [])]
    # 4 LEAN
    le = R.get("lean4") or {}
    L += ["", "## 4 · Lean 4 formalization candidates", "", "The Lean corpus was not searched in this run; these are targets, not results.", "",
          "| Claim | Disposition | Object | Proposed statement | Does not establish |", "|---|---|---|---|---|"]
    for c in le.get("formal_candidates", []):
        L.append(f"| {cell((c.get('identity') or {}).get('claim_id'))} | {cell((c.get('source_selection') or {}).get('selection_disposition'))} | "
                 f"{cell((c.get('formal_object') or {}).get('object_type'))} | {cell((c.get('formal_object') or {}).get('exact_proposed_statement'))} | "
                 f"{cell((c.get('result_and_boundary') or {}).get('what_remains_open'))} |")
    # 5 STORIES
    st = R.get("stories") or {}
    L += ["", "## 5 · Stories", ""]
    arc = st.get("narrative_arc") or {}
    if arc: L += [f"Arc: opening **{cell(arc.get('opening_device'))}** → tension **{cell(arc.get('central_tension'))}** → payoff **{cell(arc.get('resolution_or_payoff'))}** · returns to opening: {arc.get('return_to_opening')}", ""]
    for x in st.get("stories", []):
        b = x.get("boundaries") or {}
        L += [f"- **{cell(x.get('title_or_label'))}** ({x.get('story_type')}, {x.get('intended_function')}, span {cell(x.get('source_span'))}) — shows: {cell(b.get('what_it_shows'))} · does not show: {cell(b.get('what_it_does_not_show'))} · load-bearing: {not x.get('is_separable_from_argument', True)}"]
    if not st.get("stories"): L.append("- No stories found.")
    # 6 MASTER EQUATION
    me = R.get("master_equation") or {}
    L += ["", "## 6 · Master equation (analog, two independent runs)", ""]
    if me.get("slots"):
        cr = me.get("core_relation") or {}; pt = me.get("product_test") or {}
        L += [f"Core relation: {cell(cr.get('sentence') if isinstance(cr, dict) else cr)}", "",
              f"Product test: **{pt.get('run1')}** / **{pt.get('run2')}** ({pt.get('status')}) — {cell(pt.get('reason'))}", "",
              f"Analog strength **{me.get('analog_strength')}** (spread {me.get('analog_strength_spread')}) · contested slots: {', '.join(me.get('contested') or []) or 'none'}", "",
              f"> {cell(me.get('one_line'))}", "", "| Slot | Fit run 1 | Fit run 2 | Status | Plays the role | Quote |", "|---|---|---|---|---|---|"]
        L += [f"| {r['slot']} | {r['fit_1']} | {r['fit_2']} | {r['status']} | {cell(r['plays_role'])} | {cell(r['quote'])} |" for r in me["slots"]]
    mb = R.get("master_equation_b")
    if mb: L += ["", f"Part B (the paper's own equations): relationship **{(mb.get('master_equation_analysis') or {}).get('relationship')}**, {len(mb.get('equations', []))} equations"]
    # 7 COHERENCE
    co = R.get("coherence") or {}; ov = co.get("overall") or {}
    L += ["", "## 7 · Coherence", "", f"Overall **{ov.get('score')}/{ov.get('score_ceiling', 10)}** — {cell(ov.get('summary'))}", "",
          "| Dimension | Score | Reasons |", "|---|---:|---|"]
    L += [f"| {cell(d.get('name'))} | {d.get('score')} | {cell('; '.join(map(str, d.get('reasons') or [])))} |" for d in co.get("dimensions", [])]
    for head, key in (("Contradictions", "contradictions"), ("Tensions", "tensions"), ("Missing definitions", "missing_definitions")):
        if co.get(key): L += ["", f"**{head}**", ""] + [f"- {finding(x)}" for x in co[key]]
    # AUDIT
    L += ["", "## Audit", "", "| Call | Tokens | Attempts | Status |", "|---|---:|---:|---|"]
    L += [f"| {c['name']} | {c['tokens']} | {len(c['receipts'])} | {'reused' if c.get('reused') else ''} {'ERROR: ' + cell(c['error']) if c['error'] else 'ok'} |" for c in ctx["calls"]]
    L += ["", f"Total tokens: {sum(c['tokens'] for c in ctx['calls'] if not c.get('reused'))} (new) · prompts hashed in API_DEEP.run.json", ""]
    return "\n".join(L)


# ---------------------------------------------------------------- runner

def resolve_items(items, limit) -> list[tuple[Path, Path | None]]:
    """Returns (source file, paper folder or None)."""
    out = []
    for raw in items:
        p = Path(raw).expanduser()
        if p.is_dir() and (p / "paper.json").is_file():
            meta = json.loads((p / "paper.json").read_text(encoding="utf-8")); out.append((p / "00_SOURCE" / meta["source_file"], p))
        else: out += [(x, None) for x in pick.resolve([raw])]       # notes, @lists, folders (they obey their _PICK.md)
    return out[:limit] if limit else out

def run_dir(src: Path, paper: Path | None, out_root: str | None) -> Path:
    base = paper / "02_RUNS" / LABEL if paper else Path(out_root or OUTBOX) / src.stem
    return base / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")

def prior_cache(base: Path) -> dict:
    """Successful calls from earlier runs, keyed by prompt hash, so a rerun only pays for what failed."""
    cache = {}
    for f in sorted(base.glob("*/calls/*.json")):
        try:
            c = json.loads(f.read_text(encoding="utf-8"))
            if c.get("parsed") is not None and not c.get("error"): cache[c["prompt_sha256"]] = c
        except (OSError, ValueError): pass
    return cache

def run_calls(specs, a, cache, out: Path) -> dict:
    def one(spec):
        name, system, user, temp, validate = spec
        key = sha(system + "\n" + user)
        if key in cache and not a.redo: return {**cache[key], "name": name, "reused": True}
        print(f"  -> {name}", flush=True)
        return call_json(name, system, user, a, temp, validate)
    with ThreadPoolExecutor(max_workers=max(1, min(a.workers, len(specs)))) as pool:
        results = list(pool.map(one, specs))
    (out / "calls").mkdir(parents=True, exist_ok=True)
    for r in results:
        (out / "calls" / f"{r['name']}.json").write_text(json.dumps(r, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"  {'reused' if r.get('reused') else 'done  '} {r['name']}: {'ERROR ' + r['error'] if r['error'] else 'ok'} ({r['tokens']} tokens)", flush=True)
    return {r["name"]: r for r in results}

def prepare(src: Path, paper: Path | None, a) -> dict:
    text = src.read_text(encoding="utf-8", errors="replace")
    meta, body = split_front_matter(text)
    body = re.sub(r"<!-- (analysis|scorecard):start -->.*?<!-- \1:end -->\n?", "", body, flags=re.S)   # our own blocks, not the speaker's words (AGENTS.md: read the source only)
    body = re.sub(r"^\[Watch on YouTube\].*$", "", body, flags=re.M)
    title = (re.search(r'^title:\s*"?(.*?)"?\s*$', meta, re.M) or re.search(r"^#\s+(.+)$", body, re.M))
    title = title.group(1).strip() if title else src.stem
    sentences, numbered = number_source(body)
    focus, focus_hash = compose(HERE, paper, a.focus)
    src_block = source_block(meta, numbered, title)
    transcript = bool(re.search(r"^type:\s*youtube", meta, re.M))
    return {"path": src, "transcript": transcript, "turns": build_turns(sentences) if transcript else [], "title": title, "sentences": sentences, "src": src_block, "focus": focus, "focus_hash": focus_hash,
            "source_sha": hashlib.sha256(text.encode("utf-8")).hexdigest(), "paper_uuid": (re.search(r"paper_uuid:\s*\"?([\w-]+)", meta) or [None, src.stem])[1],
            "run_uuid": "deep-" + hashlib.sha256(text.encode("utf-8")).hexdigest()[:12],  # stable, so reruns can reuse cached calls
"n_sentences": len(sentences), "n_paras": len({s['para'] for s in sentences}),
            "has_math": bool(re.search(r"\$[^$\n]+\$|\\\(|\\\[|\\frac|[χΨΦ∫∑]|\b[A-Za-z]\w?\s*=\s*[A-Za-z0-9(]", body)),
            "provider": a.provider, "model": a.model}

def copy_mode(ctx, a, out: Path) -> None:
    """No API calls: write every prompt, then put one on the clipboard for pasting into a chat."""
    pdir = out / "prompts"; pdir.mkdir(parents=True, exist_ok=True)
    specs = build_calls_phase1(ctx, a) + build_calls_phase2(ctx, a, None, {"note": "per-sentence pass not run in copy mode"})
    for name, system, user, *_ in specs:
        (pdir / f"{name}.md").write_text((f"## SYSTEM\n\n{system}\n\n## USER\n\n" if system else "") + user, encoding="utf-8")
    stations = [(n, s, u) for n, s, u, *_ in specs if not n.startswith("fruits_sentences") and n != "master_equation_a2"]
    allin = ["Run the seven API DEEP stations below on the one paper at the end, in this order: " + ", ".join(PAGE_ORDER) + ".",
             "Return seven JSON objects, each under a heading with the station name. Follow each station's own rules and schema exactly.", ""]
    for n, s, u in stations:
        instr = (s + "\n\n" if s else "") + u.split("\nSOURCE (paragraphs")[0].split("SOURCE:\n")[0].split("--- BEGIN SOURCE ---")[0]
        allin += [f"=== STATION: {n} ===", instr.strip(), ""]
    allin += ["=== THE PAPER (shared by every station) ===", ctx["src"]]
    (pdir / "ALL_IN_ONE.md").write_text("\n".join(allin), encoding="utf-8")
    pick = pdir / ("ALL_IN_ONE.md" if a.copy == "all" else f"{a.copy}.md")
    if pick.exists():
        subprocess.run(["powershell", "-NoProfile", "-Command", f"Get-Content -Raw -Encoding UTF8 -LiteralPath '{pick}' | Set-Clipboard"], check=False)
        print(f"  copied to clipboard: {pick.name} ({pick.stat().st_size//1024} KB)")
    print(f"  prompts written: {pdir}")

def process(src: Path, paper: Path | None, a) -> bool:
    print(f"\n[{LABEL}] {src.name}", flush=True)
    ctx = prepare(src, paper, a); out = run_dir(src, paper, a.out); cache = {} if a.redo else prior_cache(out.parent)
    out.mkdir(parents=True, exist_ok=True)
    (out / "00_source_numbered.md").write_text(ctx["src"], encoding="utf-8")
    print(f"  {ctx['n_sentences']} sentences, {ctx['n_paras']} paragraphs, math={ctx['has_math']} -> {out}")
    if a.copy: copy_mode(ctx, a, out); return True
    llm.configure(a.workers)
    c1 = run_calls(build_calls_phase1(ctx, a), a, cache, out)
    atoms = (c1["atoms"]["parsed"] or {}) if c1.get("atoms") else {}
    for i, x in enumerate(atoms.get("atoms", []) if isinstance(atoms.get("atoms"), list) else [], 1): x["atom_id"] = f"A{i:02d}"
    rows, summary = fruits_aggregate(ctx["sentences"], [c["parsed"] for n, c in c1.items() if n.startswith("fruits_sentences")])
    turn_rows = turns_aggregate(ctx["turns"], [c["parsed"] for n, c in c1.items() if n.startswith("fruits_turns")]) if ctx["turns"] else []
    me_a = me_compare(c1.get("master_equation_a1", {}).get("parsed"), c1.get("master_equation_a2", {}).get("parsed"))
    ctx["has_math"] = ctx["has_math"] or me_a.get("equations_present", False)
    c2 = run_calls(build_calls_phase2(ctx, a, atoms, summary), a, cache, out)
    fr = c2["fruits"]["parsed"]
    if fr and not FG.validate_report(fr):
        fr = FG.normalize_and_recompute(fr, src, sha(read("fruits_rubric_v0.3.0.json")), c2["fruits"]["prompt_sha256"], {"provider": a.provider, "model": a.model})
    R = {"fruits": fr, "fruits_sentences_summary": summary, "axiom_nodes": c2["axiom_nodes"]["parsed"], "atoms": atoms,
         "lean4": c2["lean4"]["parsed"], "stories": c2["stories"]["parsed"], "master_equation": me_a,
         "master_equation_b": c2.get("master_equation_b", {}).get("parsed"), "coherence": c2["coherence"]["parsed"]}
    lex_rows, lex_summary = LT.lexicon_pass(ctx["sentences"])
    R["love_truth"] = {"truth_engine": lex_summary, "characterizations": LT.characterizations(lex_summary, R["coherence"], fr),
                       "profile": (LT.sentence_profile(turn_rows, fr, R["coherence"], lex_summary, unit="turns") if turn_rows
                                   else LT.sentence_profile(rows, fr, R["coherence"], lex_summary)),
                       "sentence_profile": LT.sentence_profile(rows, fr, R["coherence"], lex_summary) if turn_rows else None,
                       "verdict_profile": LT.profiles(fr, R["coherence"], lex_summary)}
    for row, lx in zip(rows, lex_rows): row["lexicon"] = lx["hits"]
    ctx["calls"] = list(c1.values()) + list(c2.values()); ctx["finished"] = now()
    for name in PAGE_ORDER + ["love_truth"]: (out / f"{name}.json").write_text(json.dumps(R[name], indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "fruits_sentences.json").write_text(json.dumps({"summary": summary, "sentences": rows}, indent=2, ensure_ascii=False), encoding="utf-8")
    if turn_rows: (out / "fruits_turns.json").write_text(json.dumps(turn_rows, indent=2, ensure_ascii=False), encoding="utf-8")
    ctx["turn_rows"] = turn_rows
    write_xlsx(out / "fruits_sentences.xlsx", rows, R, ctx)
    errors = {c["name"]: c["error"] for c in ctx["calls"] if c["error"]}
    receipt = {"station": LABEL, "source": str(src), "source_sha256": ctx["source_sha"], "provider": a.provider, "model": a.model,
               "focus_text": ctx["focus"], "focus_hash": ctx["focus_hash"], "finished_at": ctx["finished"], "errors": errors,
               "tokens_new": sum(c["tokens"] for c in ctx["calls"] if not c.get("reused")),
               "prompt_files": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(PROMPTS.iterdir()) if p.is_file()},
               "calls": [{k: c.get(k) for k in ("name", "prompt_sha256", "tokens", "error", "reused")} for c in ctx["calls"]]}
    (out / "API_DEEP.run.json").write_text(json.dumps(receipt, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "API_DEEP.json").write_text(json.dumps({"title": ctx["title"], "source": str(src), "order": PAGE_ORDER, **R}, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "API_DEEP.md").write_text(render_page(ctx, R), encoding="utf-8")
    (out / "API_DEEP.html").write_text(report_html.build(ctx, R, rows), encoding="utf-8")
    if paper:
        meta = json.loads((paper / "paper.json").read_text(encoding="utf-8"))
        meta.setdefault("stations_run", []).append({"station": LABEL, "run": str(out.relative_to(paper))})
        (paper / "paper.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    print(f"  page: {out / 'API_DEEP.html'}\n  tokens: {receipt['tokens_new']}  errors: {len(errors)}" + "".join(f"\n    FAIL {k}: {v}" for k, v in errors.items()))
    return not errors

def write_xlsx(path: Path, rows, R: dict, ctx: dict) -> None:
    """Sentences (colour-coded, one row each) · Profile (Love and Truth) · Paragraphs. Fills follow the report's palette."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError: path.with_suffix(".xlsx.txt").write_text("pip install openpyxl for the workbook\n"); return
    fill = {v: PatternFill("solid", fgColor=c) for v, c in {-2: "A8292A", -1: "F0A39D", 0: "E4E3DE", 1: "86B6EF", 2: "104281"}.items()}
    bold = Font(bold=True); white = Font(color="FFFFFF", bold=True)
    wb = Workbook(); ws = wb.active; ws.title = "Sentences"
    head = ["id", "paragraph", *FRUITS, "sum", "why", *[f"TE {c}" for c in LT.CATS], "TE trigger words", "sentence"]
    ws.append(head)
    for c in ws[1]: c.font = bold
    for r in rows:
        v = r["v"] or [None]*9; lx = r.get("lexicon") or {}
        ws.append([r["id"], r["para"], *v, sum(v) if r["v"] else None, "; ".join(f"{k}: {x}" for k, x in r["why"].items()) if r["why"] else "",
                   *[len(lx.get(c, [])) for c in LT.CATS], "; ".join(f"{c}: {', '.join(lx[c])}" for c in LT.CATS if lx.get(c)), r["text"]])
        row = ws.max_row
        for j in range(9):
            val = v[j]
            if isinstance(val, int):
                cell = ws.cell(row, 3 + j); cell.fill = fill[val]
                if abs(val) == 2: cell.font = white
        ws.cell(row, len(head)).alignment = Alignment(wrap_text=False)
    for j, w in enumerate([7, 6, *[7]*9, 5, 40, *[6]*6, 40, 120], 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    lt = R.get("love_truth") or {}; pr = lt.get("profile") or {}; te = lt.get("truth_engine") or {}; vp = lt.get("verdict_profile") or {}
    ps = wb.create_sheet("Profile")
    lines = [("FRUITS OF LOVE AND TRUTH", ctx["title"]), ("source", ctx["path"].name), ("model", f"{ctx['provider']}/{ctx['model']}"), ("", ""),
             ("quadrant", pr.get("quadrant")), ("character shape", ", ".join(f"{x['name']} (gap {x['gap']})" for x in pr.get("shapes_matched", [])) or "no clear shape"),
             ("level", ", ".join(pr.get("levels", [])) or "none"), ("love axis (0-1)", pr.get("love_axis")), ("truth axis (0-1)", pr.get("truth_axis")),
             *[(f"  truth part: {k}", round(v, 3)) for k, v in (pr.get("truth_parts") or {}).items()],
             ("sentences scored", pr.get("sentences")), ("share engaging a fruit", pr.get("active_share")), ("", ""),
             ("FRUIT", "net per 100 sentences | sentences engaging | deviation from paper average | verdict 0-4")]
    for k, v in lines: ps.append([k, v])
    for f in FRUITS:
        ps.append([f, (pr.get("net_per_100") or {}).get(f), (pr.get("active_sentences") or {}).get(f), (pr.get("profile_deviation") or {}).get(f), (vp.get("scores") or {}).get(f)])
    ps.append([]); ps.append(["SHAPES (all, ranked)", "gap", "match", "high", "low"])
    for sh in pr.get("shapes_ranked", []): ps.append([sh["name"], sh["gap"], "yes" if sh["match"] else "", ", ".join(sh["high"]), ", ".join(sh["low"])])
    ps.append([]); ps.append(["TRUTH ENGINE v2.0", "per 100 words", "hits", "sentence share"])
    for c in LT.CATS: ps.append([c, (te.get("per_100_words") or {}).get(c), (te.get("hits") or {}).get(c), (te.get("sentence_share") or {}).get(c)])
    ps.append(["TRUTH score", te.get("truth_score_raw"), "0-1", te.get("truth_score_0_1")])
    ps.append(["workbook characterizations", ", ".join(c["name"] for c in lt.get("characterizations", [])) or "none"])
    ps.append(["paper verdict read (second opinion)", vp.get("quadrant")])
    for row in ps.iter_rows():
        if row[0].value and str(row[0].value).isupper(): row[0].font = bold
    ps.column_dimensions["A"].width = 36; ps.column_dimensions["B"].width = 44
    pg = wb.create_sheet("Paragraphs"); pg.append(["paragraph", "sentences", *FRUITS, "mean of nine"])
    for c in pg[1]: c.font = bold
    paras: dict[str, list] = {}
    for r in rows:
        if r["v"]: paras.setdefault(r["para"], []).append(r["v"])
    for p_id, vs in paras.items():
        means = [round(sum(v[i] for v in vs) / len(vs), 2) for i in range(9)]
        pg.append([p_id, len(vs), *means, round(sum(means) / 9, 2)])
    pg.freeze_panes = "B2"
    if ctx.get("turn_rows"):
        ts = wb.create_sheet("Turns", 1); ts.append(["turn", "sentences", "speaker", *FRUITS, "sum", "why", "text"])
        for c in ts[1]: c.font = bold
        for t in ctx["turn_rows"]:
            v = t["v"] or [None]*9
            ts.append([t["id"], t["para"], t.get("speaker", ""), *v, sum(v) if t["v"] else None,
                       "; ".join(f"{k}: {x}" for k, x in t["why"].items()) if t["why"] else "", t["text"][:2000]])
            for j in range(9):
                if isinstance(v[j], int):
                    cell = ts.cell(ts.max_row, 4 + j); cell.fill = fill[v[j]]
                    if abs(v[j]) == 2: cell.font = white
        ts.freeze_panes = "D2"; ts.column_dimensions["I"].width = 8; ts.column_dimensions["N"].width = 50
    wb.save(path)

def main() -> int:
    p = argparse.ArgumentParser(prog=LABEL, description=__doc__.splitlines()[0])
    p.add_argument("items", nargs="*", help="paper folders (paper.json), .md files, or folders of .md files")
    p.add_argument("--limit", type=int); p.add_argument("--workers", type=int, default=30)
    p.add_argument("--provider", default="deepseek"); p.add_argument("--model", default="deepseek-chat")
    p.add_argument("--focus", default=""); p.add_argument("--redo", action="store_true", help="ignore cached successful calls")
    p.add_argument("--out", help="output root for plain files (default 057_API_DEEP/OUTBOX)")
    p.add_argument("--copy", nargs="?", const="all", help="no API: write prompts and copy one to the clipboard (default ALL_IN_ONE; or a call name such as fruits)")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--publish", action="store_true", help="afterwards write the analysis onto each source note (publish_analysis)")
    p.add_argument("--parallel", type=int, default=4, help="papers processed at the same time (default 4)")
    a = p.parse_args()
    inbox = INBOX
    items = resolve_items(a.items or [str(inbox)], a.limit)
    if not items: print(f"No input. Drop .md papers into {inbox} or pass files.", file=sys.stderr); return 0 if a.dry_run else 2
    if a.dry_run:
        for src, paper in items: print("PLAN", src, "->", run_dir(src, paper, a.out).parent)
        return 0
    with ThreadPoolExecutor(max_workers=max(1, min(a.parallel, len(items)))) as pool:   # papers at once; calls share the limiter
        ok = list(pool.map(lambda sp: process(*sp, a), items))
    root = Path(a.out or OUTBOX)
    if not a.copy and root.is_dir(): print(f"  index: {runs_index.build(root)}")
    print(f"\n{LABEL}: {sum(ok)}/{len(ok)} papers without errors")
    if a.publish and not a.copy: publish_on_note([src for (src, paper), good in zip(items, ok) if good and not paper])
    return 0 if all(ok) else 1

if __name__ == "__main__":
    raise SystemExit(main())
