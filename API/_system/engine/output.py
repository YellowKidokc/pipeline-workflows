"""Output contract: every station run on an item writes, in a dated folder that is never overwritten,
  <NN_STATION>.json      canonical data
  <NN_STATION>.xlsx      its own rows (e.g. per sentence), one sheet per table
  <NN_STATION>.html      the public-facing section
  <NN_STATION>.run.json  receipt: source hash, model, prompt version, focus text + hash, tokens, time, errors
"""
from __future__ import annotations

import hashlib
import html as html_lib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S.%fZ")


def dated_run_dir(item_folder: Path, label: str) -> Path:
    out = item_folder / "02_RUNS" / label / stamp()
    out.mkdir(parents=True, exist_ok=False)
    return out


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000]
    return value


def write_xlsx(path: Path, sheets: dict[str, list[dict]]) -> bool:
    """One sheet per table. Returns False (and writes nothing) when openpyxl is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items() or {"empty": []}.items():
        ws = wb.create_sheet(str(name)[:31] or "sheet")
        if not rows:
            ws.append(["(no rows)"])
            continue
        headers: list[str] = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_cell(row.get(h)) for h in headers])
        ws.freeze_panes = "A2"
    if not wb.worksheets:
        wb.create_sheet("empty").append(["(no rows)"])
    wb.save(path)
    return True


def write_bundle(out: Path, label: str, data: Any, html: str, receipt: dict[str, Any],
                 sheets: dict[str, list[dict]] | None = None) -> list[str]:
    """Write the four files. Returns warnings (e.g. xlsx skipped)."""
    warnings = []
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{label}.json").write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / f"{label}.html").write_text(html, encoding="utf-8")
    if sheets is None:
        sheets = {label: data if isinstance(data, list) else [data] if isinstance(data, dict) else []}
    if not write_xlsx(out / f"{label}.xlsx", sheets):
        warnings.append("openpyxl is not installed, so no .xlsx was written (pip install openpyxl)")
    receipt = {**receipt, "warnings": receipt.get("warnings", []) + warnings}
    (out / f"{label}.run.json").write_text(json.dumps(receipt, indent=2, ensure_ascii=False), encoding="utf-8")
    return warnings


def page(title: str, body: str) -> str:
    """Minimal standalone HTML section, readable in light and dark."""
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{html_lib.escape(title)}</title>
<style>:root{{--bg:#fbfaf7;--fg:#1d1d1b;--mute:#6b6a66;--line:#e2dfd8;--accent:#2f5d8a}}
@media (prefers-color-scheme:dark){{:root{{--bg:#151514;--fg:#ecebe6;--mute:#a09e97;--line:#34332f;--accent:#8fb5de}}}}
body{{background:var(--bg);color:var(--fg);font:15px/1.55 system-ui,sans-serif;max-width:980px;margin:0 auto;padding:24px 16px}}
h1,h2,h3{{line-height:1.25}} table{{border-collapse:collapse;width:100%;font-size:13px;margin:8px 0 18px}}
th,td{{border-bottom:1px solid var(--line);padding:5px 6px;text-align:left;vertical-align:top}}
th{{color:var(--mute);font-weight:600}} .mute{{color:var(--mute)}} code,pre{{font-size:12.5px}}
pre{{white-space:pre-wrap;background:rgba(127,127,127,.08);padding:10px;border-radius:6px}}
.tag{{display:inline-block;border:1px solid var(--line);border-radius:10px;padding:0 7px;margin:1px;font-size:12px}}
</style></head><body>{body}</body></html>"""


def esc(value: Any) -> str:
    return html_lib.escape(value if isinstance(value, str) else json.dumps(value, ensure_ascii=False))


def table(rows: list[dict], columns: list[str] | None = None) -> str:
    if not rows:
        return '<p class="mute">(none)</p>'
    columns = columns or list(rows[0])
    head = "".join(f"<th>{esc(c)}</th>" for c in columns)
    body = "".join("<tr>" + "".join(f"<td>{esc(r.get(c, ''))}</td>" for c in columns) + "</tr>" for r in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def markdown_to_html(md: str) -> str:
    """Small, dependency-free markdown renderer for reports (headings, lists, bold, paragraphs)."""
    import re
    out, in_list = [], False
    for line in md.splitlines():
        text = html_lib.escape(line)
        text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
        text = re.sub(r"`(.+?)`", r"<code>\1</code>", text)
        heading = re.match(r"^(#{1,6})\s+(.*)$", text)
        bullet = re.match(r"^\s*[-*]\s+(.*)$", text)
        if bullet:
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{bullet.group(1)}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        if heading:
            level = len(heading.group(1))
            out.append(f"<h{level}>{heading.group(2)}</h{level}>")
        elif text.strip():
            out.append(f"<p>{text}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)
