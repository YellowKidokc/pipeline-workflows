"""Fill David's Excel templates from an item's assembled JSON.

A map file in config/excel_templates/<name>.json says which value goes in which cell:

  {"template": "templates/excel/fruits_page.xlsx",          (relative to API_HOME, or absolute)
   "output": "03_REPORT/fruits_page.xlsx",                  (relative to the item folder)
   "cells": {
     "Summary!B2": "item.title",
     "Summary!B3": "metric:Flesch–Kincaid grade",
     "Summary!C5": "station:40_ANALYTICAL_ARMS.curves.love.mean"},
   "rows": {                                               (optional: tables written downwards)
     "Sentences!A2": {"from": "station:40_ANALYTICAL_ARMS.sentence_vectors", "columns": ["$key", "v.0", "v.1"]}}}

Sources: item.<field> (paper.json / video.json), metric:<name> (statistics.json), station:<NN_LABEL>.<path>
(the station's latest successful JSON). A missing value leaves the cell empty and is listed in the output.

  python tools/excel_fill.py <item folder> <map name>      (station 46 runs every map automatically)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine.items import Item, load  # noqa: E402
from engine.paths import API_HOME, expand, inside  # noqa: E402
from engine.station import latest_data  # noqa: E402


def dig(obj, path: str):
    for part in [p for p in path.split(".") if p]:
        if isinstance(obj, dict):
            obj = obj.get(part)
        elif isinstance(obj, list) and part.lstrip("-").isdigit():
            idx = int(part)
            obj = obj[idx] if -len(obj) <= idx < len(obj) else None
        else:
            return None
    return obj


def resolve(item: Item, ref: str, stats: dict):
    if ref.startswith("item."):
        return dig(item.meta, ref[5:])
    if ref.startswith("metric:"):
        return next((m["value"] for m in stats.get("metrics", []) if m["name"] == ref[7:]), None)
    if ref.startswith("station:"):
        label, _, path = ref[8:].partition(".")
        return dig(latest_data(item, label), path)
    return None


def fill(item: Item, name: str) -> tuple[Path, list[str]]:
    from openpyxl import load_workbook
    spec = json.loads(inside("config", "excel_templates", f"{name}.json").read_text(encoding="utf-8"))
    template = expand(spec["template"]) if not Path(spec["template"]).is_absolute() else Path(spec["template"])
    template = template if template.exists() else API_HOME / spec["template"]
    wb = load_workbook(template)
    stats_path = item.folder / "03_REPORT" / "statistics.json"
    stats = json.loads(stats_path.read_text(encoding="utf-8")) if stats_path.exists() else {}
    missing = []
    for cell, ref in spec.get("cells", {}).items():
        sheet, _, addr = cell.partition("!")
        value = resolve(item, ref, stats)
        if value is None:
            missing.append(f"{cell} <- {ref}")
            continue
        wb[sheet][addr] = value if isinstance(value, (int, float, str)) else json.dumps(value, ensure_ascii=False)
    for anchor, table in spec.get("rows", {}).items():
        sheet, _, addr = anchor.partition("!")
        data = resolve(item, table["from"], stats)
        rows = list(data.items()) if isinstance(data, dict) else list(enumerate(data or []))
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        col_letter, row0 = coordinate_from_string(addr)
        col0 = column_index_from_string(col_letter)
        for r, (key, value) in enumerate(rows):
            for c, column in enumerate(table["columns"]):
                v = key if column == "$key" else dig(value, column)
                wb[sheet].cell(row=row0 + r, column=col0 + c, value=v if isinstance(v, (int, float, str)) or v is None else json.dumps(v))
    out = item.folder / spec.get("output", f"03_REPORT/{name}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out, missing


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        raise SystemExit(2)
    item = load(Path(sys.argv[1]))
    if not item:
        raise SystemExit(f"not an item folder: {sys.argv[1]}")
    path, missing = fill(item, sys.argv[2])
    print(path)
    for m in missing:
        print("  empty:", m)
