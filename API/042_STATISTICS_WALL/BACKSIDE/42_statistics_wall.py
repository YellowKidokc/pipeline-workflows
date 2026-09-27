"""42_STATISTICS_WALL: every metric in Python (same numbers every run), plus context on every number.

Phase 1 (per item, in parallel, no API):
  * the built-in metric set (engine/metrics.py): text, readability, richness, structure, cohesion,
    information theory, stance, tone, citations, math, concept graph, Obsidian;
  * the analytical-arm numbers from the item's latest 40_ANALYTICAL_ARMS run (fruit means, spikes,
    counterfeit / hidden fruit, master-equation slots, axiom nodes, coherence) marked method = api;
  * external suites from config/metric_suites.json (David's Paper Intelligence suite etc.) when enabled.
Phase 2 (whole corpus): corpus percentile, series percentile and change since this item's previous run
for every metric; corpus position metrics (novelty, centroid similarity, new terms).
Output: the station bundle, plus 03_REPORT/statistics.json (read by 46) and statistics.xlsx.

  --audit-suites   list each configured suite folder, its scripts, and which answer --help (STATISTICS_WALL step 1)
"""
from __future__ import annotations

import json
import math
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(next(p for p in Path(__file__).resolve().parents if (p / "_system").is_dir()) / "_system"))
from engine.items import all_items  # noqa: E402
from engine.metrics import STOP, compute, percentile  # noqa: E402
from engine.output import esc, page, write_xlsx  # noqa: E402
from engine.paths import API_HOME, configured, external, inside  # noqa: E402
from engine.station import Context, ItemResult, Station, latest_data, latest_run  # noqa: E402
from engine.text import segment, strip_front_matter  # noqa: E402

LABEL = "42_STATISTICS_WALL"
CATALOG = json.loads(inside("config", "metric_catalog.json").read_text(encoding="utf-8"))
META = {(f["k"], m["name"]): {**m, "family_name": f["n"]} for f in CATALOG["families"] for m in f["metrics"]}
FIT = {"absent": 0, "stretched": 1, "analogous": 2, "direct": 3}
FRUIT_NAMES = {"love": "Love", "joy": "Joy", "peace": "Peace", "patience": "Patience", "kindness": "Kindness", "goodness": "Goodness",
               "faithfulness": "Faithfulness", "gentleness": "Gentleness", "self_control": "Self-control"}


def arm_metrics(item) -> list[dict]:
    d = latest_data(item, "40_ANALYTICAL_ARMS")
    if not isinstance(d, dict):
        return []
    out = []

    def add(fam, name, value, method="api"):
        if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
            out.append({"family": fam, "name": name, "value": round(float(value), 4), "method": method, "from": "40_ANALYTICAL_ARMS"})
    curves = d.get("curves", {})
    for key, label in FRUIT_NAMES.items():
        add("fruit", f"{label}, mean", (curves.get(key) or {}).get("mean"))
    spikes = (d.get("summary") or {}).get("spikes", [])
    add("fruit", "Spikes toward love", sum(1 for s in spikes if s.get("fruit") == "love" and s.get("kind") == "spike" and s.get("value", 0) > 0))
    add("fruit", "Spikes away from love", sum(1 for s in spikes if s.get("fruit") == "love" and s.get("kind") == "spike" and s.get("value", 0) < 0))
    add("fruit", "Counterfeit sentences", len({g["id"] for g in (d.get("summary") or {}).get("counterfeit", [])}))
    add("fruit", "Hidden-fruit sentences", len({g["id"] for g in (d.get("summary") or {}).get("hidden_fruit", [])}))
    me = d.get("master_equation_agreement") or {}
    slots = me.get("slots", [])
    if slots:
        add("me", "Slot coverage", sum(1 for s in slots if FIT.get(s.get("run1") or "absent", 0) > 0))
        for s in slots:
            add("me", f"{s['slot']} fit", FIT.get(s.get("run1") or "absent", 0))
        add("me", "Run agreement", sum(1 for s in slots if s.get("status") == "agreed") / len(slots))
        add("rel", "Runs compared", 2)
    pt = (me.get("product_test") or {}).get("run1")
    if pt:
        add("me", "Product test (multiplicative)", 1 if pt == "multiplicative" else 0)
    add("me", "Analog strength", me.get("analog_strength"))
    arms = d.get("arms") or {}
    ax = arms.get("axiom_nodes") or {}
    engaged = ax.get("engaged", []) if isinstance(ax, dict) else []
    if isinstance(ax, dict) and ax:
        add("ax", "Nodes engaged", len(engaged))
        for mode in ("AX_CORE", "AX_DERIVED", "AX_SCAFFOLD", "FW_EXTENDED", "HY_EVIDENCE"):
            add("ax", f"{mode} touched", sum(1 for n in engaged if n.get("mode") == mode))
        add("ax", "Directly asserted", sum(1 for n in engaged if n.get("alignment") == "directly_asserted"))
        add("ax", "Supported", sum(1 for n in engaged if n.get("alignment") == "supported"))
        add("ax", "Contested", sum(1 for n in engaged if n.get("alignment") == "contested"))
        add("ax", "Unmapped claims", len(ax.get("unmapped_claims", [])))
    coh = arms.get("coherence") or {}
    if isinstance(coh, dict) and coh:
        add("cohs", "Coherence score", coh.get("score"))
        names = {"internal_consistency": "Internal consistency", "definitional_stability": "Definitional stability",
                 "inferential_connectedness": "Inferential connectedness", "scope_discipline": "Register boundary respect"}
        for key, label in names.items():
            add("cohs", label, ((coh.get("dimensions") or {}).get(key) or {}).get("score"))
        cons = coh.get("contradictions", [])
        add("cohs", "Structural contradictions", sum(1 for c in cons if c.get("type") == "structural"))
        add("cohs", "Local contradictions", sum(1 for c in cons if c.get("type") == "local"))
        add("cohs", "Tensions", len(coh.get("tensions", [])))
        add("cohs", "Missing definitions", len(coh.get("missing_definitions", [])))
    return out


def suite_metrics(ctx) -> list[dict]:
    conf = json.loads(inside("config", "metric_suites.json").read_text(encoding="utf-8"))
    out = []
    for suite in conf["suites"]:
        if not suite.get("enabled"):
            continue
        key = suite.get("path_key")
        if key and not configured(key):
            ctx.step(f"suite {suite['name']}: skipped ({key} not configured)")
            continue
        script_dir = str(external(key)) if key else str(API_HOME)
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / f"source{ctx.item.source.suffix}"
            src.write_text(ctx.item.source.read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
            outdir = Path(tmp) / "out"
            outdir.mkdir()
            cmd = [c.replace("{python}", sys.executable).replace("{script_dir}", script_dir).replace("{API_HOME}", str(API_HOME))
                   .replace("{source}", str(src)).replace("{out}", str(outdir)) for c in suite["command"]]
            ctx.step(f"suite {suite['name']}: running")
            try:
                done = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
            except (OSError, subprocess.TimeoutExpired) as exc:
                ctx.warnings.append(f"suite {suite['name']} failed to run: {exc}")
                continue
            if done.returncode != 0:
                ctx.warnings.append(f"suite {suite['name']} exit {done.returncode}: {(done.stderr or done.stdout)[-300:]}")
                continue
            found = 0
            for jf in outdir.rglob("*.json"):
                try:
                    data = json.loads(jf.read_text(encoding="utf-8"))
                except ValueError:
                    continue
                for path, value in _flatten(data):
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
                        out.append({"family": suite["name"], "name": f"{jf.stem}:{path}"[:120], "value": value,
                                    "method": f"python:{suite['name']}"})
                        found += 1
            ctx.step(f"suite {suite['name']}: {found} numeric value(s)")
    return out


def _flatten(obj, prefix=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from _flatten(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(obj, list) and len(obj) <= 50 and all(not isinstance(x, (dict, list)) for x in obj):
        return
    elif not isinstance(obj, list):
        yield prefix, obj


def bag(text: str) -> Counter:
    import re
    return Counter(w for w in re.findall(r"[a-z][a-z'-]+", text.lower()) if w not in STOP and len(w) > 2)


def cosine(a: Counter, b: Counter) -> float:
    dot = sum(v * b.get(k, 0) for k, v in a.items())
    na, nb = math.sqrt(sum(v * v for v in a.values())), math.sqrt(sum(v * v for v in b.values()))
    return dot / (na * nb) if na and nb else 0.0


def audit_suites() -> int:
    conf = json.loads(inside("config", "metric_suites.json").read_text(encoding="utf-8"))
    for suite in conf["suites"]:
        key = suite.get("path_key")
        folder = (external(key) / suite.get("subfolder", "")) if key and configured(key) else \
            (API_HOME / "vendor" / suite["vendored"]) if suite.get("vendored") else None
        print(f"\n== {suite['name']} ({'enabled' if suite.get('enabled') else 'off'}) {folder or '(path not configured)'}")
        if not folder or not folder.exists():
            continue
        for py in sorted(folder.rglob("*.py"))[:300]:
            try:
                done = subprocess.run([sys.executable, str(py), "--help"], capture_output=True, text=True, timeout=30, cwd=py.parent)
                flags = sorted(set(__import__("re").findall(r"--[a-z][\w-]*", done.stdout)))
                print(f"  {'help' if done.returncode == 0 else 'fail'}  {py.relative_to(folder)}  {' '.join(flags)[:150]}")
            except (OSError, subprocess.TimeoutExpired):
                print(f"  fail  {py.relative_to(folder)}")
    return 0


def main() -> int:
    if "--audit-suites" in sys.argv:
        return audit_suites()
    station = Station(LABEL, kind="both", prompt_files=["../../engine/metrics.py", "../../config/metric_suites.json"])
    items = station.items()
    if station.args.dry_run or not items:
        return station.run(lambda ctx: None, items)
    phase1: dict[str, tuple[list[dict], dict, Context]] = {}

    def one(item):
        ctx = Context(station, item)
        ctx.step("phase 1: computing local metrics")
        raw = item.source.read_text(encoding="utf-8", errors="replace")
        text = strip_front_matter(item.text())
        paragraphs, sentences = segment(text)
        metrics, extras = compute(text, paragraphs, sentences, raw)
        ctx.step(f"phase 1: {len(metrics)} python metric(s)")
        arms = arm_metrics(item)
        ctx.step(f"phase 1: {len(arms)} metric(s) from the latest 40_ANALYTICAL_ARMS run" if arms else "phase 1: no 40_ANALYTICAL_ARMS run yet")
        metrics += arms + suite_metrics(ctx)
        vec = latest_data(item, "40_ANALYTICAL_ARMS")
        if isinstance(vec, dict):
            extras["fruit_vectors"] = [v["v"] for v in vec.get("sentence_vectors", {}).values()]
            extras["me_slots"] = {s["slot"]: s.get("run1") for s in (vec.get("master_equation_agreement") or {}).get("slots", [])}
            extras["axiom_nodes"] = [{"mode": n.get("mode"), "alignment": n.get("alignment"), "node": n.get("node")}
                                     for n in ((vec.get("arms") or {}).get("axiom_nodes") or {}).get("engaged", [])]
            extras["coherence_dimensions"] = {k: [v.get("score")] for k, v in (((vec.get("arms") or {}).get("coherence") or {}).get("dimensions") or {}).items() if isinstance(v, dict)}
            extras["love_words"] = []
            run = latest_run(item, "40_ANALYTICAL_ARMS")
            if run:
                try:
                    from openpyxl import load_workbook
                    ws = load_workbook(run / "40_ANALYTICAL_ARMS.xlsx", read_only=True)["sentences"]
                    rows = list(ws.iter_rows(values_only=True))
                    col = rows[0].index("lex_love")
                    extras["love_words"] = [r[col] or 0 for r in rows[1:]]
                except Exception:
                    pass
        phase1[f"{item.kind}:{item.id}"] = (metrics, extras, ctx)

    station.say(f"42_STATISTICS_WALL phase 1: {len(items)} item(s)")
    Context(station, None).parallel(one, items, width=max(1, min(station.args.workers, 16)))

    # Phase 2: corpus context.
    station.say("42_STATISTICS_WALL phase 2: corpus and series percentiles, change since last run")
    corpus: dict[str, dict] = {}
    series_of: dict[str, str] = {}
    bags: dict[str, Counter] = {}
    for other in all_items("both"):
        key = f"{other.kind}:{other.id}"
        series_of[key] = other.meta.get("series") or ""
        if key in phase1:
            corpus[key] = {(m["family"], m["name"]): m["value"] for m in phase1[key][0]}
        else:
            prev = latest_data(other, LABEL)
            if isinstance(prev, dict):
                corpus[key] = {(m["family"], m["name"]): m["value"] for m in prev.get("metrics", [])}
        try:
            bags[key] = bag(other.text())
        except OSError:
            pass

    def process(ctx):
        item = ctx.item
        key = f"{item.kind}:{item.id}"
        metrics, extras, ctx1 = phase1[key]
        ctx.steps, ctx.warnings = ctx1.steps + ctx.steps, ctx1.warnings
        peers = [k for k in corpus if k.split(":", 1)[0] == item.kind]
        series_peers = [k for k in peers if series_of.get(k) and series_of.get(k) == (item.meta.get("series") or None)]
        me_bag = bags.get(key, Counter())
        others = [k for k in peers if k != key and k in bags]
        if others:
            sims = [cosine(me_bag, bags[k]) for k in others]
            centroid = sum((bags[k] for k in others), Counter())
            metrics.append({"family": "sem", "name": "Novelty vs your corpus", "value": round(1 - max(sims), 4), "method": "python (bag-of-words cosine)"})
            metrics.append({"family": "sem", "name": "Similarity to corpus centroid", "value": round(cosine(me_bag, centroid), 4), "method": "python (bag-of-words cosine)"})
            vocab = set(centroid)
            metrics.append({"family": "sem", "name": "New terms introduced", "value": sum(1 for w in me_bag if w not in vocab), "method": "python"})
            series_others = [k for k in series_peers if k != key and k in bags]
            if series_others:
                sc = sum((bags[k] for k in series_others), Counter())
                metrics.append({"family": "sem", "name": "Similarity to series centroid", "value": round(cosine(me_bag, sc), 4), "method": "python (bag-of-words cosine)"})
        previous = latest_data(item, LABEL)
        prev_vals = {(m["family"], m["name"]): m["value"] for m in (previous or {}).get("metrics", [])} if isinstance(previous, dict) else {}
        for m in metrics:
            k = (m["family"], m["name"])
            info = META.get(k, {})
            m.setdefault("unit", info.get("unit", ""))
            m["dir"], m["dec"] = info.get("dir", "band"), info.get("dec", 3)
            m["family_name"] = info.get("family_name", m["family"])
            pop = [corpus[p][k] for p in peers if k in corpus.get(p, {})]
            spop = [corpus[p][k] for p in series_peers if k in corpus.get(p, {})]
            m["corpus_percentile"] = percentile(m["value"], pop)
            m["series_percentile"] = percentile(m["value"], spop)
            m["corpus_n"], m["series_n"] = len(pop), len(spop)
            m["academic_percentile"] = None
            m["previous"] = prev_vals.get(k)
            m["change_since_previous"] = round(m["value"] - prev_vals[k], 4) if k in prev_vals else None
            m["in_catalog"] = k in META
        ctx.step(f"phase 2: {len(metrics)} metric(s), corpus n={len(peers)}, series n={len(series_peers)}")
        headline = [{"label": h, "family": f, "name": n} for h, f, n in HEADLINE]
        data = {"schema_version": 2, "item": {"id": item.id, "kind": item.kind, "title": item.title, "series": item.meta.get("series")},
                "metrics": metrics, "headline": headline, "headline_status": "PROVISIONAL: David picks the 12 after seeing real numbers",
                "extras": extras, "coverage": {"catalog_metrics": len(META), "computed_in_catalog": sum(1 for m in metrics if m["in_catalog"]),
                                               "not_computed": sorted(f"{f}:{n}" for f, n in META if (f, n) not in {(m['family'], m['name']) for m in metrics})},
                "corpus": {"n": len(peers), "series_n": len(series_peers)}}
        report = item.folder / "03_REPORT"
        report.mkdir(parents=True, exist_ok=True)
        (report / "statistics.json").write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        rows = [{k: m.get(k) for k in ("family_name", "name", "value", "unit", "corpus_percentile", "series_percentile", "change_since_previous",
                                       "method", "dir", "corpus_n", "series_n")} for m in metrics]
        write_xlsx(report / "statistics.xlsx", {"metrics": rows, "not_computed": [{"metric": x} for x in data["coverage"]["not_computed"]]})
        ctx.step(f"wrote 03_REPORT/statistics.json and statistics.xlsx ({data['coverage']['computed_in_catalog']}/{len(META)} catalog metrics computed)")
        html = page(f"Statistics wall: {item.title}", "<h1>Statistics wall</h1><p class='mute'>Every metric, grouped by family. "
                    "The full matrix view is in 03_REPORT/report.html (station 46).</p>" + wall_html(metrics))
        return ItemResult(data, html, {"metrics": rows}, None,
                          {"words": next((m["value"] for m in metrics if m["name"] == "Word count"), None)})

    station.per_item = False
    return station.run(process, items)


HEADLINE = [("Reading grade (Flesch–Kincaid)", "read", "Flesch–Kincaid grade"), ("Vocabulary richness (MTLD)", "lex", "MTLD"),
            ("Love, mean per sentence", "fruit", "Love, mean"), ("Coherence score", "cohs", "Coherence score"),
            ("Master-equation slots filled", "me", "Slot coverage"), ("Information per sentence", "info", "Information per sentence"),
            ("Novelty vs your corpus", "sem", "Novelty vs your corpus"), ("Hedge:booster ratio", "stance", "Hedge:booster ratio"),
            ("Counterfeit sentences", "fruit", "Counterfeit sentences"), ("Kill conditions stated", "math", "Kill conditions stated"),
            ("Scripture references", "cite", "Scripture references"), ("Word count", "text", "Word count")]


def wall_html(metrics: list[dict]) -> str:
    groups: dict[str, list] = {}
    for m in metrics:
        groups.setdefault(m.get("family_name", m["family"]), []).append(m)
    parts = []
    for fam, rows in groups.items():
        cells = "".join(f"<div title='{esc(m['method'])}'><b>{esc(m['name'])}</b> {m['value']:g}{esc(m.get('unit') or '')}"
                        f"<span class='mute'> · c{'' if m['corpus_percentile'] is None else round(m['corpus_percentile'])}</span></div>" for m in rows)
        parts.append(f"<h3>{esc(fam)}</h3><div style='columns:3 220px;font-size:13px'>{cells}</div>")
    return "".join(parts)


if __name__ == "__main__":
    raise SystemExit(main())
