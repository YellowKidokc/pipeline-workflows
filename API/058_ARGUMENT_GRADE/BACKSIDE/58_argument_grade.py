"""58_ARGUMENT_GRADE (was API_HOME station 49): every argument in a source gets a Strength score and an Originality score, each 0-8,
computed from yes/partly/no checks with quotes (the model never gives a number). Machine ceiling 8;
David's human review adds 1 and a Lean receipt adds 1 (max 10).

Calls per source: 1 extraction (the argument list) + 2 independent gradings of that same list (temperature 0.7).
Scores are the mean of the two gradings; items where they differ are shown as disagreement.
Every argument is tagged to David's case map (K01-K18, F1-F4), so the ledger can list the arguments the case needs
that are still weak: the develop-next list.

Outputs: OUTBOX/<source>/<stamp>/arguments.json · ARGUMENTS.html
         OUTBOX/ARGUMENT_LEDGER.xlsx + LEDGER.html (all sources; human/Lean columns are kept across rebuilds)
"""
from __future__ import annotations
import argparse, hashlib, html, importlib.util, json, sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
FRONT = HERE.parent                                             # 058_ARGUMENT_GRADE: INBOX/ and OUTBOX/ next to BACKSIDE/
sys.path.insert(0, str(FRONT.parent / "_system"))
_spec = importlib.util.spec_from_file_location("deep48", FRONT.parent / "057_API_DEEP" / "BACKSIDE" / "57_api_deep.py")
deep48 = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(deep48)   # reuse its tested call + source prep
from engine import llm                                          # noqa: E402
from engine.publish import publish_on_note                      # noqa: E402

LABEL = "58_ARGUMENT_GRADE"
OUTBOX = FRONT / "OUTBOX"
PROMPTS = HERE / "prompts"
ST = ["st1", "st2", "st3", "st4", "st5", "st6", "st7", "st8"]
OR = ["or1", "or2", "or3", "or4"]
ST_NAMES = {"st1": "conclusion", "st2": "premises", "st3": "support", "st4": "inference", "st5": "objection",
            "st6": "scope", "st7": "falsifiable", "st8": "convergence"}
OR_NAMES = {"or1": "not restated", "or2": "new support", "or3": "new bridge", "or4": "new structure"}
DEVELOP_BELOW = 5.0     # strength below this, on an argument the case map needs, lands on the develop-next list
BATCH = 6               # arguments per grading call


def read(name): return (PROMPTS / name).read_text(encoding="utf-8")

EXTRACT = """List every distinct argument the source makes or defends (typically 2-10; merge restatements).
An argument has a conclusion and at least one reason. Include arguments a guest or quoted opponent makes if the
source engages them. Return JSON only:
{"arguments": [{"id": "G1", "name": "short name", "speaker": "", "conclusion": "one sentence",
  "premises": ["P1 ...", "P2 ..."], "where": "sentence ids or section"}]}"""


def grade_problems(p, ids):
    got = {a.get("id") for a in p.get("arguments", []) if isinstance(a, dict)}
    bad = [a.get("id") for a in p.get("arguments", []) if not all(k in (a.get("strength") or {}) for k in ST)
           or not all(k in (a.get("originality") or {}) for k in OR)]
    return ([f"missing argument ids {sorted(ids - got)}"] if ids - got else []) + ([f"incomplete checklists: {bad}"] if bad else [])


def item(d, k):
    x = (d or {}).get(k) or {}
    v = x.get("v") if isinstance(x, dict) else x
    v = int(v) if isinstance(v, (int, float)) else 0
    q = x.get("quote", "") if isinstance(x, dict) else ""
    return (min(2, max(0, v)) if q else 0), q, (x.get("why", "") if isinstance(x, dict) else "")   # no quote, no credit


def combine(args, runs) -> list[dict]:
    out = []
    by = [{a.get("id"): a for a in (r or {}).get("arguments", [])} for r in runs]
    for a in args:
        g = [b.get(a["id"]) for b in by if b.get(a["id"])]
        if not g: continue
        st = {k: [item(x.get("strength"), k) for x in g] for k in ST}
        orr = {k: [item(x.get("originality"), k) for x in g] for k in OR}
        s_runs = [sum(st[k][i][0] for k in ST) / 2 for i in range(len(g))]          # 0-16 -> 0-8
        o_runs = [sum(orr[k][i][0] for k in OR) for i in range(len(g))]            # 0-8
        mean = lambda xs: round(sum(xs) / len(xs), 2)
        disagree = [k for k in ST + OR if len({(st.get(k) or orr.get(k))[i][0] for i in range(len(g))}) > 1]
        first = g[0]
        cm = sorted({c for x in g for c in (x.get("case_map") or []) if isinstance(c, str)})
        out.append({**a, "strength": mean(s_runs), "strength_runs": s_runs, "originality": mean(o_runs), "originality_runs": o_runs,
                    "items": {k: {"v": [r[0] for r in (st.get(k) or orr.get(k))], "quote": next((r[1] for r in (st.get(k) or orr.get(k)) if r[1]), ""),
                                  "why": next((r[2] for r in (orr.get(k) or []) if r[2]), "")} for k in ST + OR},
                    "disagree": disagree, "agreement": round(1 - len(disagree) / 12, 2),
                    "prior_art": first.get("prior_art", ""), "weakest_link": first.get("weakest_link", ""),
                    "develop_next": first.get("develop_next", ""), "case_map": cm, "domain": first.get("domain", "")})
    return out


def process(src: Path, a) -> list[dict]:
    print(f"\n[{LABEL}] {src.name}", flush=True)
    ctx = deep48.prepare(src, None, a)
    base = Path(a.out or OUTBOX) / src.stem
    out = base / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"); out.mkdir(parents=True, exist_ok=True)
    cache = {} if a.redo else deep48.prior_cache(base)
    llm.configure(a.workers)
    ex = deep48.run_calls([("extract", "", EXTRACT + "\n\n" + ctx["src"], 0.1, None)], a, cache, out)["extract"]
    args = [x for x in ((ex["parsed"] or {}).get("arguments") or []) if isinstance(x, dict) and x.get("id")]
    if not args: print("  no arguments found"); return []
    # 12 quoted checks per argument: grade in batches so each reply stays under the 8k output cap
    specs = []
    for bi in range(0, len(args), BATCH):
        batch = args[bi:bi + BATCH]
        listing = json.dumps([{k: x.get(k) for k in ("id", "name", "speaker", "conclusion", "premises")} for x in batch], ensure_ascii=False, indent=1)
        task = (read("ARGUMENT_GRADE.md") + "\n\nCASE MAP:\n" + read("CASE_MAP.md") +
                f"\n\nGRADE EXACTLY THESE ARGUMENTS (keep their ids; do not add or drop any):\n{listing}")
        ids = {x["id"] for x in batch}
        specs += [(f"grade_{k}_{bi // BATCH + 1}", task, ctx["src"] + f"\n\n(independent grading run {k})", 0.7,
                   lambda p, ids=ids: grade_problems(p, ids)) for k in ("a", "b")]
    gr = deep48.run_calls(specs, a, cache, out)
    merged = [{"arguments": [x for n, c in gr.items() if n.startswith(f"grade_{k}_") for x in ((c["parsed"] or {}).get("arguments") or [])]}
              for k in ("a", "b")]
    graded = combine(args, merged)
    key = ctx["source_sha"][:8]
    for g in graded: g.update({"key": f"{key}-{g['id']}", "source": src.stem, "title": ctx["title"]})
    calls = [ex, *gr.values()]
    (out / "arguments.json").write_text(json.dumps({"title": ctx["title"], "source": str(src), "arguments": graded,
                                                   "tokens_new": sum(c["tokens"] for c in calls if not c.get("reused")),
                                                   "errors": {c["name"]: c["error"] for c in calls if c["error"]}}, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "ARGUMENTS.html").write_text(page(ctx["title"], graded, single=True), encoding="utf-8")
    print(f"  {len(graded)} arguments · tokens {sum(c['tokens'] for c in calls if not c.get('reused'))} · {out / 'ARGUMENTS.html'}")
    for g in graded: print(f"    {g['id']} S{g['strength']:.1f} O{g['originality']:.1f} agree {g['agreement']:.2f} {g['case_map']} {g['name'][:60]}")
    return graded


# ---------------------------------------------------------------- ledger (all sources) and pages

def latest_all(root: Path) -> list[dict]:
    rows = []
    for d in sorted(p for p in root.iterdir() if p.is_dir()):
        runs = sorted(d.glob("*/arguments.json"))
        if runs: rows += json.loads(runs[-1].read_text(encoding="utf-8")).get("arguments", [])
    return rows


def rulings(ledger: Path) -> dict:
    """Keep David's human-review and Lean columns when the ledger is rebuilt."""
    try:
        from openpyxl import load_workbook
        ws = load_workbook(ledger)["Arguments"]; head = [c.value for c in ws[1]]
        i_key, i_h, i_l, i_n = (head.index(x) for x in ("key", "human_review (0/1)", "lean_receipt", "david_notes"))
        return {r[i_key]: (r[i_h], r[i_l], r[i_n]) for r in ws.iter_rows(min_row=2, values_only=True) if r[i_key]}
    except Exception:
        return {}


def total(g, rul) -> float:
    h, lean, _ = rul.get(g["key"], (None, None, None))
    return round(min(8, g["strength"]) + (1 if str(h).strip() in ("1", "yes", "y", "True") else 0) + (1 if lean else 0), 2)


def write_ledger(root: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = latest_all(root); path = root / "ARGUMENT_LEDGER.xlsx"; rul = rulings(path)
    wb = Workbook(); ws = wb.active; ws.title = "Arguments"
    head = ["key", "source", "argument", "speaker", "domain", "case_map", "strength (0-8)", "originality (0-8)", "agreement",
            "human_review (0/1)", "lean_receipt", "total (0-10)", "weakest_link", "develop_next", "prior_art", "conclusion", "david_notes"]
    ws.append(head)
    for c in ws[1]: c.font = Font(bold=True)
    for g in sorted(rows, key=lambda g: (-len(g["case_map"]), g["strength"])):
        h, lean, notes = rul.get(g["key"], (None, None, None))
        ws.append([g["key"], g["source"], g["name"], g.get("speaker", ""), g["domain"], " ".join(g["case_map"]), g["strength"], g["originality"],
                   g["agreement"], h, lean, total(g, rul), g["weakest_link"], g["develop_next"], g["prior_art"], g["conclusion"], notes])
        for col, v in ((7, g["strength"]), (8, g["originality"])):
            ws.cell(ws.max_row, col).fill = PatternFill("solid", fgColor="104281" if v >= 6 else "86B6EF" if v >= 4 else "F0A39D" if v >= 2 else "A8292A")
    ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions
    dv = wb.create_sheet("Develop next")
    dv.append(["case_map", "argument", "source", "strength", "originality", "weakest_link", "develop_next"])
    for c in dv[1]: c.font = Font(bold=True)
    for g in sorted([g for g in rows if g["case_map"] and g["strength"] < DEVELOP_BELOW], key=lambda g: g["strength"]):
        dv.append([" ".join(g["case_map"]), g["name"], g["source"], g["strength"], g["originality"], g["weakest_link"], g["develop_next"]])
    cov = wb.create_sheet("Case coverage"); cov.append(["case id", "arguments found", "best strength", "best originality", "best argument"])
    for c in cov[1]: c.font = Font(bold=True)
    for cid in [f"K{i:02d}" for i in range(1, 19)] + ["F1", "F2", "F3", "F4"]:
        hits = [g for g in rows if cid in g["case_map"]]; best = max(hits, key=lambda g: g["strength"]) if hits else None
        cov.append([cid, len(hits), best["strength"] if best else None, max((g["originality"] for g in hits), default=None), best["name"] if best else "not yet argued"])
    wb.save(path)
    (root / "LEDGER.html").write_text(page("Argument ledger: every source", rows, single=False, rul=rul), encoding="utf-8")
    return path


def page(title, rows, single, rul=None) -> str:
    rul = rul or {}
    data = json.dumps([{k: g.get(k) for k in ("key", "id", "name", "source", "speaker", "domain", "case_map", "strength", "originality",
                                              "strength_runs", "originality_runs", "agreement", "disagree", "weakest_link", "develop_next",
                                              "prior_art", "conclusion", "items")} | {"total": total(g, rul)} for g in rows], ensure_ascii=False).replace("</", "<\\/")
    return PAGE.replace("__TITLE__", html.escape(title)).replace("__DATA__", data).replace("__DEV__", str(DEVELOP_BELOW)) \
               .replace("__STN__", json.dumps(ST_NAMES)).replace("__ORN__", json.dumps(OR_NAMES))


PAGE = r"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title>
<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400&display=swap" rel="stylesheet">
<style>
:root{--plane:#f5f5f2;--surface:#fcfcfb;--ink:#0b0b0b;--muted:#6f6d67;--grid:#e1e0d9;--axis:#c3c2b7;--mid:#e4e3de;--g1:#86b6ef;--g2:#2a78d6;--g3:#104281;--b1:#f0a39d;--b2:#e34948;--b3:#a8292a}
@media (prefers-color-scheme:dark){:root{--plane:#0f0f0e;--surface:#1a1a19;--ink:#fff;--muted:#9a988f;--grid:#2c2c2a;--axis:#383835;--mid:#383835;--g1:#1c5cab;--g2:#3987e5;--g3:#9ec5f4;--b1:#7d2a28;--b2:#e66767;--b3:#f4aaa6}}
*{box-sizing:border-box}body{margin:0;background:var(--plane);color:var(--ink);font:14px/1.5 "IBM Plex Sans",system-ui,sans-serif}.wrap{max-width:1180px;margin:0 auto;padding:24px 16px 60px}
h1{font-size:24px;margin:0 0 4px}h2{font-size:17px;margin:0 0 6px}.sub{color:var(--muted);font-size:12.5px;margin:0 0 12px}
section{background:var(--surface);border:1px solid var(--grid);border-radius:10px;padding:18px;margin:16px 0}
svg{display:block;width:100%;height:auto;overflow:visible}svg text{fill:var(--ink);font-family:inherit}.mut{fill:var(--muted)}
table{width:100%;border-collapse:collapse;font-size:12.5px}th{text-align:left;color:var(--muted);font-weight:500;border-bottom:1px solid var(--axis);padding:6px 8px;white-space:nowrap}
td{border-bottom:1px solid var(--grid);padding:6px 8px;vertical-align:top}.scroll{overflow-x:auto}.id{font:11px "IBM Plex Mono",monospace;color:var(--muted)}
.pill{display:inline-block;border-radius:99px;padding:0 7px;font-size:11px;border:1px solid var(--grid);margin:1px}
.dots span{display:inline-block;width:14px;height:14px;border-radius:3px;margin-right:2px;vertical-align:middle}
#tip{position:fixed;pointer-events:none;background:var(--surface);border:1px solid var(--grid);box-shadow:0 6px 20px rgba(0,0,0,.18);border-radius:8px;padding:8px 10px;font-size:12px;max-width:380px;display:none;z-index:9}
</style></head><body><div class="wrap" id="app"></div><div id="tip"></div><script>
var D=__DATA__, DEV=__DEV__, STN=__STN__, ORN=__ORN__;
function css(v){return getComputedStyle(document.documentElement).getPropertyValue(v).trim();}
function esc(s){return String(s==null?"":s).replace(/[&<>"]/g,function(c){return {"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c];});}
function el(t,a,p){var e=document.createElementNS("http://www.w3.org/2000/svg",t);for(var k in a)e.setAttribute(k,a[k]);p&&p.appendChild(e);return e;}
function tx(s,x,y,t,a){var e=el("text",Object.assign({x:x,y:y,"font-size":11},a||{}),s);e.textContent=t;return e;}
var tip=document.getElementById("tip");function hover(n,h){n.addEventListener("mousemove",function(e){tip.innerHTML=h;tip.style.display="block";var x=e.clientX+14;if(x+390>innerWidth)x=e.clientX-390;tip.style.left=Math.max(4,x)+"px";tip.style.top=(e.clientY+14)+"px";});n.addEventListener("mouseleave",function(){tip.style.display="none";});}
function col(v){return v>=6?css("--g3"):v>=4?css("--g1"):v>=2?css("--b1"):css("--b3");}
var app=document.getElementById("app");
app.innerHTML='<h1>__TITLE__</h1><p class="sub">'+D.length+' arguments. Strength and originality are each 0–8, computed from yes / partly / no checks with quotes, averaged over two independent gradings. Machine ceiling 8; human review +1; Lean receipt +1.</p>';
var s1=document.createElement("section");s1.innerHTML='<h2>Strength × originality</h2><p class="sub">Each circle is one argument. Size = how many case-map entries it serves. Dashed ring = the two gradings disagree on 3 or more checks. Lower left = restated and weak; upper right = new and strong. Red outline = the case needs it and it is below '+DEV+'.</p>';app.appendChild(s1);
function scatter(){var old=s1.querySelector("svg");old&&old.remove();var s=el("svg",{viewBox:"0 0 640 420"});s1.appendChild(s);var X=50,Y=10,W=560,H=360;
 [[0,0,"restated & weak"],[1,0,"new but weak — develop"],[0,1,"strong but familiar"],[1,1,"new & strong"]].forEach(function(q){el("rect",{x:X+q[0]*W/2,y:Y+(1-q[1])*H/2,width:W/2,height:H/2,fill:q[0]&&q[1]?css("--g1"):!q[0]&&!q[1]?css("--b1"):css("--mid"),"fill-opacity":.25,stroke:css("--surface")},s);tx(s,X+q[0]*W/2+8,Y+(1-q[1])*H/2+16,q[2],{"class":"mut","font-size":10.5});});
 for(var v=0;v<=8;v+=2){tx(s,X+v/8*W,Y+H+14,v,{"text-anchor":"middle","class":"mut","font-size":9});tx(s,X-6,Y+H-v/8*H+3,v,{"text-anchor":"end","class":"mut","font-size":9});}
 tx(s,X+W/2,Y+H+32,"originality →",{"text-anchor":"middle","class":"mut"});tx(s,14,Y+H/2,"strength →",{"text-anchor":"middle","class":"mut",transform:"rotate(-90 14 "+(Y+H/2)+")"});
 D.forEach(function(g,i){var jx=((i*37)%11-5)*1.2,jy=((i*53)%11-5)*1.2;var need=g.case_map.length&&g.strength<DEV;
  var c=el("circle",{cx:X+g.originality/8*W+jx,cy:Y+H-g.strength/8*H+jy,r:6+Math.min(6,g.case_map.length*2),fill:col(g.strength),"fill-opacity":.85,stroke:need?css("--b2"):css("--surface"),"stroke-width":need?2.5:1.5,"stroke-dasharray":g.disagree.length>=3?"3 2":"none"},s);
  hover(c,'<b>'+esc(g.name)+'</b><div class="id">'+esc(g.source)+' · '+g.id+'</div>strength '+g.strength+' ('+g.strength_runs.join(' / ')+') · originality '+g.originality+' ('+g.originality_runs.join(' / ')+')<br>'+(g.case_map.length?'serves '+g.case_map.join(' '):'not on the case map')+'<br><span style="color:var(--muted)">weakest: '+esc(g.weakest_link)+'</span>');});}
scatter();matchMedia("(prefers-color-scheme: dark)").addEventListener("change",scatter);
var s2=document.createElement("section");s2.innerHTML='<h2>Every argument</h2><p class="sub">Squares: the eight strength checks, then the four originality checks (dark blue = yes, light = partly, red = no; half-and-half = the two runs disagree). Hover a square for the quote.</p>';app.appendChild(s2);
function need(x){return x.case_map.length&&x.strength<DEV?1:0;}
var rows=D.slice().sort(function(a,b){return need(b)-need(a)||a.strength-b.strength;});
var t='<div class="scroll"><table><thead><tr><th>Argument</th><th>Strength</th><th>Orig.</th><th>Total</th><th>Checks</th><th>Case</th><th>Weakest link / develop next</th><th>Closest prior art</th></tr></thead><tbody>';
rows.forEach(function(g,ri){t+='<tr><td><b>'+esc(g.name)+'</b><div class="id">'+esc(g.source)+' · '+g.id+(g.speaker?' · '+esc(g.speaker):'')+'</div>'+esc(g.conclusion)+'</td><td>'+g.strength+'</td><td>'+g.originality+'</td><td>'+g.total+'</td><td class="dots" style="white-space:nowrap">'+
 Object.keys(STN).concat(Object.keys(ORN)).map(function(k,j){var v=g.items[k].v,a=v[0],b=v.length>1?v[1]:v[0];function c(x){return x===2?css("--g3"):x===1?css("--g1"):css("--b2");}
  return (j===8?'<span style="width:6px;background:none"></span>':'')+'<span data-r="'+ri+'" data-k="'+k+'" style="background:linear-gradient(90deg,'+c(a)+' 50%,'+c(b)+' 50%)"></span>';}).join('')+'</td><td>'+g.case_map.map(function(c){return '<span class="pill">'+c+'</span>';}).join('')+'</td><td>'+esc(g.weakest_link)+'<div class="id">→ '+esc(g.develop_next)+'</div></td><td class="id">'+esc(g.prior_art)+'</td></tr>';});
s2.insertAdjacentHTML("beforeend",t+'</tbody></table></div>');
s2.querySelectorAll(".dots span[data-k]").forEach(function(n){var g=rows[+n.dataset.r],k=n.dataset.k,it=g.items[k];hover(n,'<b>'+(STN[k]||ORN[k])+'</b>: '+it.v.join(' / ')+'<br>'+esc(it.quote||'no quote given (scored 0)')+(it.why?'<br><span style="color:var(--muted)">'+esc(it.why)+'</span>':''));});
</script></body></html>"""


def main() -> int:
    p = argparse.ArgumentParser(prog=LABEL, description=__doc__.splitlines()[0])
    p.add_argument("items", nargs="*"); p.add_argument("--limit", type=int); p.add_argument("--workers", type=int, default=30)
    p.add_argument("--provider", default="deepseek"); p.add_argument("--model", default="deepseek-chat")
    p.add_argument("--focus", default=""); p.add_argument("--redo", action="store_true"); p.add_argument("--out")
    p.add_argument("--publish", action="store_true", help="afterwards write the analysis onto each source note")
    p.add_argument("--dry-run", action="store_true", help="list what would run; no API")
    a = p.parse_args(); a.copy = None
    inbox = FRONT / "INBOX"
    items = deep48.resolve_items(a.items or [str(inbox)], a.limit)
    if not items: print(f"No input. Drop .md files into {inbox} or pass files.", file=sys.stderr); return 0 if a.dry_run else 2
    if a.dry_run:
        for src, _ in items: print("PLAN", src)
        return 0
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda sp: process(sp[0], a), items))
    root = Path(a.out or OUTBOX)
    print(f"\nledger: {write_ledger(root)}\n        {root / 'LEDGER.html'}")
    if a.publish: publish_on_note([sp[0] for sp, good in zip(items, results) if good and not sp[1]])
    return 0 if all(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
