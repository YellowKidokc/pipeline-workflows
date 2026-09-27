"""40_ANALYTICAL_ARMS: Fruits (flagship) + master equation x2 + axiom nodes + coherence, one run per item.

Stages (ANALYTICAL_ARMS_V1): 0 prep (local) -> 1 lexicon (local) -> 2 sentence scores (API-40.1, output
ranges in parallel, whole paper every call) -> 3 aggregate (local) -> 4 the four arms in parallel
(API-40.2..40.5) -> cross-arm checks (local) -> 5 report (.md/.html + Truth Engine-layout .xlsx).
Every stage is cached, so a rerun only pays for what changed.
"""
from __future__ import annotations

import json
import re
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(next(p for p in Path(__file__).resolve().parents if (p / "_system").is_dir()) / "_system"))
from engine.output import esc, markdown_to_html, page  # noqa: E402
from engine.paths import API_HOME, configured, external  # noqa: E402
from engine.station import ItemResult, Station  # noqa: E402
from engine.text import numbered, ranges, segment  # noqa: E402

LABEL = "40_ANALYTICAL_ARMS"
HERE = Path(__file__).resolve().parent
PLUGIN = HERE / "fruits_plugin"
CONF = json.loads((PLUGIN / "plugin.json").read_text(encoding="utf-8"))
FRUITS = CONF["fruits"]
LEX_FAMILIES = ["grounding", "contradiction", "propaganda", "jargon", "hedge", "absolute", "negation"]


# ------------------------------------------------------------------ lexicons
def _excel_terms(path: Path) -> dict:
    """Tolerant reader for the two workbooks: every sheet, any column that looks like terms.
    Fruit is taken from a fruit-named column/sheet or a 'fruit'/'category' column."""
    out = {"fruit": {f: set() for f in FRUITS}, "anti_fruit": {f: set() for f in FRUITS}, **{k: set() for k in LEX_FAMILIES}}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return out
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        name = ws.title.lower().replace("-", "_").replace(" ", "_")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h or "").strip().lower() for h in rows[0]]
        anti = "anti" in name
        family = next((k for k in LEX_FAMILIES if k in name), None)
        term_col = next((i for i, h in enumerate(header) if h in ("term", "word", "words", "lexeme", "phrase", "token")), 0)
        fruit_col = next((i for i, h in enumerate(header) if h in ("fruit", "category", "dimension", "fruit_name")), None)
        for row in rows[1:]:
            if not row or term_col >= len(row) or not row[term_col]:
                continue
            term = str(row[term_col]).strip().lower()
            if family:
                out[family].add(term)
                continue
            fruit = None
            if fruit_col is not None and fruit_col < len(row) and row[fruit_col]:
                label = str(row[fruit_col]).lower().replace("-", "_").replace(" ", "_")
                fruit = next((f for f in FRUITS if f in label), None)
            fruit = fruit or next((f for f in FRUITS if f in name), None)
            if "fruit" in name and fruit:
                out["anti_fruit" if anti else "fruit"][fruit].add(term)
    return out


def load_lexicons() -> tuple[dict, list[str]]:
    seed = json.loads((PLUGIN / CONF["lexicons"]["seed"]).read_text(encoding="utf-8"))
    lex = {"fruit": {f: set(seed["fruit"].get(f, [])) for f in FRUITS},
           "anti_fruit": {f: set(seed["anti_fruit"].get(f, [])) for f in FRUITS},
           **{k: set(seed.get(k, [])) for k in LEX_FAMILIES}}
    sources = ["seed_lexicons.json"]
    for key in CONF["lexicons"]["excel_keys"]:
        if configured(key):
            extra = _excel_terms(external(key))
            for f in FRUITS:
                lex["fruit"][f] |= extra["fruit"][f]
                lex["anti_fruit"][f] |= extra["anti_fruit"][f]
            for k in LEX_FAMILIES:
                lex[k] |= extra[k]
            sources.append(external(key).name)
    return lex, sources


def _hits(text: str, terms: set[str]) -> list[tuple[str, int]]:
    low = text.lower()
    found = []
    for term in terms:
        for m in re.finditer(r"(?<![\w-])" + re.escape(term) + r"(?![\w-])", low):
            found.append((term, m.start()))
    return found


def lexicon_pass(sentences, lex) -> list[dict]:
    rows = []
    for s in sentences:
        low = s.text.lower()
        quotes = [(m.start(), m.end()) for m in re.finditer(r"[\"“][^\"”]+[\"”]", s.text)]
        negs = [p for _, p in _hits(low, lex["negation"])]

        def flagged(term: str, pos: int) -> str:
            if any(a <= pos < b for a, b in quotes):
                return "quoted"
            if any(0 <= pos - n <= 25 for n in negs):
                return "negated"
            return ""
        row = {"id": s.id, "para": s.para, "fruit": {}, "anti": {}, "words": {}}
        for f in FRUITS:
            fh = [(t, flagged(t, p)) for t, p in _hits(low, lex["fruit"][f])]
            ah = [(t, flagged(t, p)) for t, p in _hits(low, lex["anti_fruit"][f])]
            row["fruit"][f] = sum(1 for _, fl in fh if not fl)
            row["anti"][f] = sum(1 for _, fl in ah if not fl)
            if fh or ah:
                row["words"][f] = [t + (f" ({fl})" if fl else "") for t, fl in fh] + [f"anti:{t}" + (f" ({fl})" if fl else "") for t, fl in ah]
        for k in LEX_FAMILIES:
            row[k] = len(_hits(low, lex[k]))
        rows.append(row)
    return rows


# ------------------------------------------------------------------ helpers
def rolling(values: list[float], window: int) -> list[float]:
    out = []
    for i in range(len(values)):
        chunk = values[max(0, i - window + 1): i + 1]
        out.append(round(sum(chunk) / len(chunk), 3))
    return out


def pearson(a: list[float], b: list[float]) -> float | None:
    if len(a) < 3 or len(set(a)) < 2 or len(set(b)) < 2:
        return None
    try:
        return round(statistics.correlation(a, b), 3)
    except statistics.StatisticsError:
        return None


def slot_table() -> str:
    if configured("me_pills"):
        texts = []
        for f in sorted(external("me_pills").rglob("*")):
            if f.suffix.lower() in (".md", ".yaml", ".yml") and re.search(r"ME-01-02\d", f.name + f.read_text(encoding="utf-8", errors="ignore")[:500]):
                texts.append(f.read_text(encoding="utf-8", errors="ignore")[:1500])
        if texts:
            return "\n---\n".join(texts)
    return "\n".join(line for line in (HERE / "arms" / "SLOTS.md").read_text(encoding="utf-8").splitlines() if not line.startswith("#"))


def axiom_registry() -> tuple[str, str]:
    path = external("axiom_registry") if configured("axiom_registry") else \
        API_HOME / "vendor" / "api_deep" / "AXIOM_NODES" / "AXIOMS_PART1_MODE_CLASSIFICATION.md"
    if not path.exists():
        return "(registry not found)", str(path)
    nodes, mode = [], None
    for line in path.read_text(encoding="utf-8").splitlines():
        s = line.strip()
        if s.startswith("## ") and "(" in s:
            mode = s.split("(")[0][3:].strip()
        elif s.startswith("- `") and "|" in s:
            parts = [p.strip() for p in s.split("|")]
            nodes.append(f"{parts[0].strip('- `')}: {parts[1]} ({mode})")
    return "\n".join(nodes), path.name


# ------------------------------------------------------------------ stages
def process(ctx):
    station = ctx.station
    paragraphs, sentences = segment(ctx.text)
    if not sentences:
        ctx.errors.append("no sentences found in the source")
        return None
    ctx.step(f"stage 0 prep: {len(paragraphs)} paragraphs, {len(sentences)} sentences")
    paper = numbered(paragraphs, sentences)
    ids = [s.id for s in sentences]

    lex, lex_sources = load_lexicons()
    lexrows = lexicon_pass(sentences, lex)
    ctx.step(f"stage 1 lexicon: {sum(sum(r['fruit'].values()) for r in lexrows)} fruit-word hits, "
             f"{sum(sum(r['anti'].values()) for r in lexrows)} anti-fruit hits (sources: {', '.join(lex_sources)})")
    by_id = {r["id"]: r for r in lexrows}

    # Stage 2: output ranges in parallel, whole paper every call.
    size = int(station.settings.get("sentences_per_output_range", 80))
    chunks = ranges(ids, size)
    ctx.step(f"stage 2 sentence scores: {len(chunks)} parallel call(s) of up to {size} sentences, whole paper each")
    prompt_body = (PLUGIN / CONF["sentence_prompt"]).read_text(encoding="utf-8")

    def score_range(chunk: list[str]) -> list[dict]:
        hints = "\n".join(f"{i}: " + "; ".join(f"{f}={','.join(w)}" for f, w in by_id[i]["words"].items()) for i in chunk if by_id[i]["words"])
        prompt = (f"{prompt_body}\n\nIDS: {', '.join(chunk)}\n\nLEXICON HINT (fruit words found):\n{hints or '(none)'}"
                  f"\n\nPAPER (every sentence numbered):\n{paper}")
        reply = ctx.cached(f"sentences-{chunk[0]}-{chunk[-1]}", lambda: ctx.call_json("fruits_sentences", prompt, focus=False))
        return (reply or {}).get("sentences", []) if isinstance(reply, dict) else []

    vectors = {}
    for part in ctx.parallel(score_range, chunks):
        for row in part:
            v = row.get("v") or []
            if row.get("id") in by_id and len(v) == 9:
                vectors[row["id"]] = {"v": [max(-2, min(2, int(x))) for x in v], "why": row.get("why") or {}}
    missing = [i for i in ids if i not in vectors]
    if missing:
        ctx.warnings.append(f"{len(missing)} sentence(s) got no score: {', '.join(missing[:10])}")
    ctx.step(f"stage 2 done: {len(vectors)}/{len(ids)} sentences scored")

    # Stage 3: aggregate (local).
    window = int(CONF.get("rolling_window", 5))
    curves, spikes, top = {}, [], {}
    for k, f in enumerate(FRUITS):
        values = [vectors.get(i, {"v": [0] * 9})["v"][k] for i in ids]
        roll = rolling(values, window)
        curves[f] = {"values": values, "rolling": roll, "mean": round(sum(values) / len(values), 3),
                     "share_negative": round(sum(1 for x in values if x < 0) / len(values), 3)}
        for j, i in enumerate(ids):
            if abs(values[j]) == 2:
                spikes.append({"id": i, "fruit": f, "value": values[j], "kind": "spike"})
            if j and (roll[j - 1] < 0 <= roll[j] or roll[j - 1] >= 0 > roll[j]):
                spikes.append({"id": i, "fruit": f, "value": roll[j], "kind": "turn"})
        ranked = sorted(zip(ids, values), key=lambda x: x[1])
        top[f] = {"toward": [i for i, v in reversed(ranked) if v > 0][:5], "away": [i for i, v in ranked if v < 0][:5]}
    text_of = {s.id: s.text for s in sentences}
    gap = []
    for i in ids:
        if i not in vectors:
            continue
        for k, f in enumerate(FRUITS):
            words, value = by_id[i]["fruit"][f], vectors[i]["v"][k]
            if words >= 1 and value <= -1:
                gap.append({"id": i, "fruit": f, "kind": "counterfeit", "fruit_words": words, "mechanism": value, "text": text_of[i]})
            elif words == 0 and value >= 1:
                gap.append({"id": i, "fruit": f, "kind": "hidden_fruit", "fruit_words": 0, "mechanism": value, "text": text_of[i]})
    balance = {f: pearson(curves["love"]["values"], curves[f]["values"]) for f in FRUITS[1:]}
    para_rollup = []
    for p in paragraphs:
        row = {"paragraph": p.id, "sentences": len(p.sentences)}
        for k, f in enumerate(FRUITS):
            vals = [vectors[i]["v"][k] for i in p.sentences if i in vectors]
            row[f] = round(sum(vals) / len(vals), 2) if vals else None
        para_rollup.append(row)
    ratios = {k: sum(r[k] for r in lexrows) for k in LEX_FAMILIES}
    love_mean = curves["love"]["mean"]
    character = ("Pure Fruit" if love_mean > 0.8 and not [g for g in gap if g["kind"] == "counterfeit"] else
                 "Sophisticated Propagandist" if ratios["propaganda"] > 2 and love_mean < 0 else
                 "Fear Merchant" if curves["peace"]["mean"] < -0.5 else
                 "Grounded Truth-Teller" if ratios["grounding"] >= max(3, ratios["propaganda"] * 3) and love_mean >= 0 else "Mixed")
    summary = {"curve_means": {f: curves[f]["mean"] for f in FRUITS}, "spikes": spikes[:60], "top": top,
               "counterfeit": [g for g in gap if g["kind"] == "counterfeit"][:25],
               "hidden_fruit": [g for g in gap if g["kind"] == "hidden_fruit"][:25],
               "balance_vs_love": balance, "characterization": {"label": character, "status": "provisional heuristic until the Truth Engine Characterizations sheet is wired"},
               "lexicon_totals": ratios}
    ctx.step(f"stage 3 aggregate: love mean {love_mean:+.2f}, {len(spikes)} spikes/turns, "
             f"{len(summary['counterfeit'])} counterfeit, {len(summary['hidden_fruit'])} hidden-fruit candidates")

    # Stage 4: the four arms in parallel (master equation twice).
    rubric = (PLUGIN / CONF["rubric"]).read_text(encoding="utf-8")
    registry, registry_name = axiom_registry()
    candidates = [i for i in ids if by_id[i]["contradiction"] or by_id[i]["negation"]][:60]
    arms = {
        "fruits_verdict": ("fruits_verdict", f"{(PLUGIN / CONF['verdict_prompt']).read_text(encoding='utf-8')}\n\nRUBRIC v0.3.0:\n{rubric}"
                           f"\n\nSENTENCE-LEVEL SUMMARY:\n{json.dumps(summary, ensure_ascii=False)}", None),
        "master_equation_1": ("master_equation", (HERE / "arms/MASTER_EQUATION.md").read_text(encoding="utf-8").replace("SLOTS (meanings below).", "SLOTS:\n" + slot_table()) + "\nRUN: 1", 0.7),
        "master_equation_2": ("master_equation", (HERE / "arms/MASTER_EQUATION.md").read_text(encoding="utf-8").replace("SLOTS (meanings below).", "SLOTS:\n" + slot_table()) + "\nRUN: 2", 0.7),
        "axiom_nodes": ("axiom_nodes", f"{(HERE / 'arms/AXIOM_NODES.md').read_text(encoding='utf-8')}\n\nREGISTRY ({registry_name}):\n{registry}", None),
        "coherence": ("coherence", f"{(HERE / 'arms/COHERENCE.md').read_text(encoding='utf-8')}\n\nCOHERENCE BASE:\n"
                      f"{(HERE / 'arms/COHERENCE_BASE.md').read_text(encoding='utf-8')}\n\nCANDIDATES: {', '.join(candidates) or '(none)'}", None),
    }
    ctx.step("stage 4 arms: fruits verdict, master equation x2, axiom nodes, coherence (in parallel)")

    def run_arm(name: str):
        task, body, temperature = arms[name]
        prompt = f"{body}\n\nPAPER (every sentence numbered):\n{paper}"
        return name, ctx.cached(f"arm-{name}", lambda: ctx.call_json(task, prompt, temperature=temperature))

    results = dict(ctx.parallel(run_arm, list(arms)))
    me1, me2 = results.get("master_equation_1") or {}, results.get("master_equation_2") or {}
    me_agreement = []
    fits1 = {s.get("slot"): s.get("fit") for s in me1.get("slots", [])} if isinstance(me1, dict) else {}
    fits2 = {s.get("slot"): s.get("fit") for s in me2.get("slots", [])} if isinstance(me2, dict) else {}
    for slot in sorted(set(fits1) | set(fits2)):
        me_agreement.append({"slot": slot, "run1": fits1.get(slot), "run2": fits2.get(slot),
                             "status": "agreed" if fits1.get(slot) == fits2.get(slot) else "contested"})
    s1, s2 = (me1 or {}).get("analog_strength"), (me2 or {}).get("analog_strength")
    me_summary = {"product_test": {"run1": (me1.get("product_test") or {}).get("answer") if isinstance(me1, dict) else None,
                                   "run2": (me2.get("product_test") or {}).get("answer") if isinstance(me2, dict) else None},
                  "analog_strength": min(x for x in (s1, s2) if isinstance(x, (int, float))) if any(isinstance(x, (int, float)) for x in (s1, s2)) else None,
                  "strength_spread": [s1, s2], "slots": me_agreement}
    me_summary["product_test"]["status"] = "agreed" if me_summary["product_test"]["run1"] == me_summary["product_test"]["run2"] else "contested"

    # Cross-arm checks (local).
    disagreements = []
    coh = results.get("coherence") or {}
    fv = results.get("fruits_verdict") or {}
    structural = [c for c in coh.get("contradictions", []) if c.get("type") == "structural"] if isinstance(coh, dict) else []
    truth_gate = ((fv.get("gates") or {}).get("gate.truth_evidence") or (fv.get("gates") or {}).get("truth_evidence") or {}) if isinstance(fv, dict) else {}
    truth_result = truth_gate.get("result") if isinstance(truth_gate, dict) else truth_gate
    if structural and truth_result == "pass":
        disagreements.append({"check": "coherence found a structural contradiction but the Fruits truth gate passed",
                              "detail": structural[:3]})
    zero_fruits = [f for f, d in (fv.get("dimensions") or {}).items() if isinstance(d, dict) and d.get("score") == 0] if isinstance(fv, dict) else []
    if "multiplicative" in (me_summary["product_test"]["run1"], me_summary["product_test"]["run2"]) and zero_fruits:
        disagreements.append({"check": "master equation reads multiplicative and a Fruits dimension scores 0 (zero-collapse)",
                              "detail": zero_fruits})
    loving = {i for i in ids if i in vectors and vectors[i]["v"][0] == 2}
    for node in (results.get("axiom_nodes") or {}).get("engaged", []) if isinstance(results.get("axiom_nodes"), dict) else []:
        if node.get("alignment") == "contested" and loving & set(node.get("sentence_ids", [])):
            disagreements.append({"check": "axiom node contested in a sentence scoring +2 on love (charitable disagreement?)",
                                  "detail": node})
    if me_summary["product_test"]["status"] == "contested":
        disagreements.append({"check": "the two master-equation runs disagree on the product test", "detail": me_summary["product_test"]})
    ctx.step(f"cross-arm checks: {len(disagreements)} disagreement(s); master-equation slots contested: "
             f"{sum(1 for s in me_agreement if s['status'] == 'contested')}")

    # Stage 5: report.
    data = {"paper": {"id": ctx.item.id, "title": ctx.item.title, "sentences": len(ids), "paragraphs": len(paragraphs)},
            "scale": CONF["sentence_scale"], "lexicon_sources": lex_sources, "axiom_registry": registry_name,
            "sentence_vectors": {i: vectors[i] for i in ids if i in vectors}, "curves": curves, "summary": summary,
            "paragraph_rollup": para_rollup, "arms": results, "master_equation_agreement": me_summary,
            "where_the_arms_disagree": disagreements}
    md = report_md(ctx.item.title, data, text_of)
    sheet = []
    for s in sentences:
        r = by_id[s.id]
        v = vectors.get(s.id, {"v": [None] * 9, "why": {}})
        row = {"sentence": s.id, "paragraph": s.para, "text": s.text}
        row.update({f: v["v"][k] for k, f in enumerate(FRUITS)})
        row.update({f"lex_{f}": r["fruit"][f] for f in FRUITS})
        row.update({"FRUIT": sum(r["fruit"].values()), "ANTI-FRUIT": sum(r["anti"].values()), "GROUNDING": r["grounding"],
                    "CONTRADICTION": r["contradiction"], "PROPAGANDA": r["propaganda"], "JARGON": r["jargon"],
                    "HEDGE": r["hedge"], "ABSOLUTE": r["absolute"], "trigger_words": r["words"], "why": v.get("why")})
        row["gap"] = ", ".join(f"{g['kind']}:{g['fruit']}" for g in gap if g["id"] == s.id)
        sheet.append(row)
    headline = {f"fruit_mean_{f}": curves[f]["mean"] for f in FRUITS}
    headline.update({"counterfeit_count": len(summary["counterfeit"]), "hidden_fruit_count": len(summary["hidden_fruit"]),
                     "me_analog_strength": me_summary["analog_strength"],
                     "coherence_score": coh.get("score") if isinstance(coh, dict) else None})
    return ItemResult(data, page(f"Analytical arms: {ctx.item.title}", markdown_to_html(md)),
                      {"sentences": sheet, "paragraphs": para_rollup, "gap": gap, "master_equation": me_agreement,
                       "disagreements": [{"check": d["check"], "detail": d["detail"]} for d in disagreements]}, md, headline)


def report_md(title: str, d: dict, text_of: dict) -> str:
    s, arms = d["summary"], d["arms"]
    fv = arms.get("fruits_verdict") or {}
    out = [f"# Analytical arms: {title}", "", f"*Scale {d['scale']['min']}..+{d['scale']['max']} per sentence ({d['scale']['status']}). "
           "AI proposal until David reviews it.*", "", "## Fruits of the Spirit", ""]
    if isinstance(fv, dict) and fv.get("verdict"):
        out += [f"**Verdict:** {fv['verdict']}", ""]
    out += ["| fruit | sentence mean | share negative | verdict 0-4 |", "|---|---|---|---|"]
    for f in FRUITS:
        dim = (fv.get("dimensions") or {}).get(f, {}) if isinstance(fv, dict) else {}
        out.append(f"| {f} | {d['curves'][f]['mean']:+.2f} | {d['curves'][f]['share_negative']:.0%} | {dim.get('score', '?') if isinstance(dim, dict) else dim} |")
    out += ["", "**Love curve** (rolling mean): " + " ".join("▁▂▃▄▅▆▇█"[min(7, max(0, int((x + 2) / 4 * 7.99)))] for x in d["curves"]["love"]["rolling"][:200]), ""]
    for label, key in (("Toward love", "toward"), ("Away from love", "away")):
        ids = s["top"]["love"][key][:3]
        out += [f"**{label}:**"] + [f"- {i}: {text_of[i][:200]}" for i in ids] + [""]
    out += ["### Counterfeit candidates (fruit words, anti-fruit mechanism)", ""]
    out += [f"- {g['id']} {g['fruit']} ({g['mechanism']:+d}): {g['text'][:180]}" for g in s["counterfeit"][:10]] or ["- none"]
    out += ["", "### Hidden fruit (love done, not advertised)", ""]
    out += [f"- {g['id']} {g['fruit']} (+{g['mechanism']}): {g['text'][:180]}" for g in s["hidden_fruit"][:10]] or ["- none"]
    out += ["", f"*Characterization (secondary read, {s['characterization']['status']}): {s['characterization']['label']}*", ""]
    if isinstance(fv, dict) and fv.get("repairs"):
        out += ["**Repairs:**"] + [f"- {r}" for r in fv["repairs"]] + [""]
    me = d["master_equation_agreement"]
    out += ["## Master equation", "", f"Product test: run 1 {me['product_test']['run1']}, run 2 {me['product_test']['run2']} "
            f"({me['product_test']['status']}). Analog strength {me['analog_strength']} (runs {me['strength_spread']}).", "",
            "| slot | run 1 | run 2 | status |", "|---|---|---|---|"]
    out += [f"| {x['slot']} | {x['run1']} | {x['run2']} | {x['status']} |" for x in me["slots"]]
    ax = arms.get("axiom_nodes") or {}
    out += ["", f"## Axiom nodes (registry: {d['axiom_registry']}; canonical registry still to be decided)", ""]
    out += [f"- {n.get('node')} {n.get('name', '')}: {n.get('alignment')} ({n.get('confidence')}) \"{str(n.get('quote', ''))[:120]}\""
            for n in (ax.get("engaged", []) if isinstance(ax, dict) else [])] or ["- none engaged"]
    coh = arms.get("coherence") or {}
    out += ["", "## Coherence", "", f"Score: {coh.get('score') if isinstance(coh, dict) else '?'}", ""]
    out += [f"- {c.get('type')} contradiction {c.get('a')} vs {c.get('b')}: {c.get('note', '')}" for c in (coh.get("contradictions", []) if isinstance(coh, dict) else [])] or ["- no contradictions reported"]
    out += ["", "## Where the arms disagree", ""]
    out += [f"- {x['check']}" for x in d["where_the_arms_disagree"]] or ["- no disagreements found"]
    return "\n".join(out)


if __name__ == "__main__":
    raise SystemExit(Station(LABEL, prompt_files=["PROMPT.md", "fruits_plugin/*.md", "fruits_plugin/*.json", "fruits_plugin/rubric/*",
                                                 "fruits_plugin/lexicons/*", "arms/*"]).run(process))
