#!/usr/bin/env python3
"""
Output writing -- turn a model's reply into the file format you want, and
optionally write it into a template (Excel / HTML).

Formats (config OUTPUT_FORMAT):
    md   (default) - the reply + a metadata header, as .md
    txt            - just the reply, as .txt
    html           - the reply as .html (filled into a template if provided)
    json           - the reply parsed/validated as .json
    csv            - the reply as .csv
    xlsx           - the reply's rows written into a real .xlsx
                     (into a copy of templates/*.xlsx if you provide one)

Drop a template in the folder's  templates/  folder and the model is told to
match it (its columns, for a spreadsheet; its {{OUTPUT}} placeholder, for HTML).
"""

import re
import csv
import io
import json
import shutil
import datetime
import pathlib

import providers as P


# ---------------------------------------------------------------------------
#  Helpers
# ---------------------------------------------------------------------------

def strip_fences(text: str) -> str:
    """Remove a leading ```json / ```csv / ``` fence and trailing ``` if present."""
    t = text.strip()
    m = re.match(r"^```[a-zA-Z0-9]*\s*\n(.*)\n```$", t, flags=re.DOTALL)
    if m:
        return m.group(1).strip()
    return t


def find_template(folder: pathlib.Path, fmt: str):
    """Return the first template file in templates/ matching the format, or None."""
    tdir = folder / "templates"
    if not tdir.exists():
        return None
    exts = {
        "xlsx": (".xlsx",),
        "html": (".html", ".htm"),
        "csv":  (".csv",),
        "json": (".json",),
    }.get(fmt, ())
    for p in sorted(tdir.iterdir()):
        if p.is_file() and p.suffix.lower() in exts:
            return p
    return None


def template_hint(template_path, fmt: str) -> str:
    """A short instruction describing the template so the model fits it."""
    if not template_path:
        return ""
    if fmt in ("xlsx", "csv"):
        headers = _read_headers(template_path)
        if headers:
            cols = ", ".join(headers)
            return (f"\n\nReturn the result as a JSON array of objects using EXACTLY "
                    f"these column keys: {cols}. Return ONLY the JSON array, no prose.")
    if fmt == "html":
        try:
            tpl = template_path.read_text(encoding="utf-8", errors="replace")
            return ("\n\nFill the following HTML template, replacing the marker "
                    "{{OUTPUT}} with your answer. Return ONLY the completed HTML:\n\n"
                    + tpl[:4000])
        except Exception:
            return ""
    return ""


def format_instruction(fmt: str, template_path) -> str:
    """Tell the model what shape to produce when there's no template."""
    if template_path:
        return template_hint(template_path, fmt)
    if fmt == "json":
        return "\n\nReturn ONLY valid JSON. No prose, no markdown fences."
    if fmt in ("csv", "xlsx"):
        return ("\n\nReturn ONLY the result as a JSON array of row objects "
                "(each object = one row, keys = column names). No prose.")
    if fmt == "html":
        return "\n\nReturn ONLY a complete, valid HTML document. No markdown fences."
    return ""


def _read_headers(xlsx_or_csv):
    """Read the header row from a template .xlsx or .csv."""
    p = pathlib.Path(xlsx_or_csv)
    if p.suffix.lower() == ".csv":
        with open(p, newline="", encoding="utf-8") as f:
            row = next(csv.reader(f), [])
            return [c for c in row if c]
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            return [str(c) for c in row if c is not None]
    except Exception:
        return []
    return []


def _rows_from_reply(text: str):
    """Parse the model reply into (headers, list-of-rows). Accepts JSON or CSV."""
    body = strip_fences(text)
    # Try JSON first.
    try:
        data = json.loads(body)
        if isinstance(data, dict) and "rows" in data:
            data = data["rows"]
        if isinstance(data, list) and data and isinstance(data[0], dict):
            headers = list({k: None for r in data for k in r}.keys())
            rows = [[r.get(h, "") for h in headers] for r in data]
            return headers, rows
        if isinstance(data, list) and data and isinstance(data[0], list):
            return [str(c) for c in data[0]], [list(r) for r in data[1:]]
    except (ValueError, TypeError):
        pass
    # Fall back to CSV.
    try:
        reader = list(csv.reader(io.StringIO(body)))
        if reader:
            return reader[0], reader[1:]
    except Exception:
        pass
    return None, None


# ---------------------------------------------------------------------------
#  Writers
# ---------------------------------------------------------------------------

def _header_comment(result: dict) -> str:
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    price = P.price_for(result["provider"], result["model"])
    if price:
        cost = f"${(result['input_tokens']/1000*price[0] + result['output_tokens']/1000*price[1]):.6f}"
    else:
        cost = "n/a"
    return (f"<!--\n  provider : {result['provider']}\n  model    : {result['model']}\n"
            f"  finished : {ts}\n  tokens   : {result['input_tokens']} in / "
            f"{result['output_tokens']} out\n  est cost : {cost}\n"
            f"  seconds  : {result['elapsed']:.1f}\n-->\n\n")


def _unique(dest: pathlib.Path) -> pathlib.Path:
    if not dest.exists():
        return dest
    i = 1
    while True:
        cand = dest.parent / f"{dest.stem}_{i}{dest.suffix}"
        if not cand.exists():
            return cand
        i += 1


def _write_xlsx(dest, headers, rows, template_path):
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("xlsx output needs openpyxl  ->  pip install openpyxl") from e

    if template_path and pathlib.Path(template_path).suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        tpl_headers = _read_headers(template_path) or headers
        for r in rows:
            mapped = {h: v for h, v in zip(headers, r)}
            ws.append([mapped.get(h, "") for h in tpl_headers])
        wb.save(dest)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        if headers:
            ws.append(headers)
        for r in rows:
            ws.append(list(r))
        wb.save(dest)


def write_output(outbox: pathlib.Path, stem: str, result: dict,
                 fmt: str, template_path) -> pathlib.Path:
    """Write the reply in the requested format; returns the path written."""
    fmt = (fmt or "md").lower()
    text = result["text"]

    if fmt == "txt":
        dest = _unique(outbox / f"{stem}.txt")
        dest.write_text(text, encoding="utf-8")
        return dest

    if fmt == "html":
        body = strip_fences(text)
        if template_path:
            tpl = pathlib.Path(template_path).read_text(encoding="utf-8", errors="replace")
            if "{{OUTPUT}}" in tpl:
                body = tpl.replace("{{OUTPUT}}", body)
        dest = _unique(outbox / f"{stem}.html")
        dest.write_text(body, encoding="utf-8")
        return dest

    if fmt == "json":
        body = strip_fences(text)
        dest = _unique(outbox / f"{stem}.json")
        try:
            parsed = json.loads(body)
            dest.write_text(json.dumps(parsed, indent=2), encoding="utf-8")
        except (ValueError, TypeError):
            # Not valid JSON -- keep the raw reply so nothing is lost.
            dest.write_text(body, encoding="utf-8")
        return dest

    if fmt == "csv":
        headers, rows = _rows_from_reply(text)
        dest = _unique(outbox / f"{stem}.csv")
        if headers is not None:
            with open(dest, "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(headers)
                w.writerows(rows)
        else:
            dest.write_text(strip_fences(text), encoding="utf-8")
        return dest

    if fmt == "xlsx":
        headers, rows = _rows_from_reply(text)
        if headers is None:
            # Could not structure it -- fall back to .md so nothing is lost.
            dest = _unique(outbox / f"{stem}.response.md")
            dest.write_text(_header_comment(result) +
                            "> NOTE: could not parse this reply into rows; "
                            "saved raw.\n\n" + text, encoding="utf-8")
            return dest
        dest = _unique(outbox / f"{stem}.xlsx")
        _write_xlsx(dest, headers, rows, template_path)
        return dest

    # default: markdown with metadata header
    dest = _unique(outbox / f"{stem}.response.md")
    dest.write_text(_header_comment(result) + text, encoding="utf-8")
    return dest
