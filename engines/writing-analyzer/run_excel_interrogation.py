#!/usr/bin/env python3
"""
run_excel_interrogation.py
==========================
Reads the universal_domain_mapping Excel workbook directly,
converts key sheets to text context, then runs prompt 12
(The Interrogation) through DeepSeek and OpenAI O3.

No vectorization needed — the Excel data is injected as structured text.

Saves to: universality-class-runner/excel_run_TIMESTAMP/
  excel_deepseek.md
  excel_o3.md
"""

import sys
import time
import pathlib
import datetime

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SCRIPT_DIR   = pathlib.Path(__file__).resolve().parent
CONFIG_PATH  = SCRIPT_DIR / "config.txt"
PROMPT_FILE  = SCRIPT_DIR / "prompts" / "12_interrogation.txt"
OUT_DIR      = SCRIPT_DIR / "universality-class-runner"

EXCEL_PATH   = r"\\dlowenas\h_hp\Desktop\universal_domain_mapping_with_coherence_v9.xlsx"

# Sheets to extract — ordered by importance for context
SHEETS_TO_EXTRACT = [
    ("Key Insight",          80),
    ("Core Rosetta Stone",   60),
    ("The Proof",            60),
    ("10 Laws - Universal Axioms", 80),
    ("Trinity Operating System",   60),
    ("Coherence × 10 Laws",        80),
    ("Derivative Chain",           50),
    ("Exhaustive CD Mapping",      90),   # the 70-domain table
]

DEEPSEEK_BASE_URL = "https://api.deepseek.com"

PRICING = {
    "deepseek-chat":     (0.00027, 0.00110),
    "deepseek-reasoner": (0.00055, 0.00219),
    "o3":                (0.01000, 0.04000),
    "o3-mini":           (0.00110, 0.00440),
    "gpt-4o":            (0.00250, 0.01000),
}


def parse_config(path):
    cfg = {}
    if not path.exists():
        return cfg
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" in line:
            k, _, v = line.partition("=")
            cfg[k.strip()] = v.strip()
    return cfg


def load_foundation(path_str):
    if not path_str:
        return ""
    p = pathlib.Path(path_str)
    if not p.exists():
        print(f"  [WARNING] Foundation doc not found: {p}")
        return ""
    text = p.read_text(encoding="utf-8", errors="replace")
    print(f"  Foundation doc: {p.name}  ({len(text):,} chars)")
    return (
        "[FOUNDATION DOCUMENT — CANONICAL AXIOMS — TREAT AS FIXED FRAMEWORK]\n"
        + text
        + "\n[END FOUNDATION DOCUMENT]"
    )


def extract_excel_to_text(excel_path):
    """Use COM automation to extract key sheets from the Excel workbook."""
    print(f"  Opening Excel: {excel_path}")
    try:
        import win32com.client
        xl = win32com.client.Dispatch("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Open(str(excel_path))
    except Exception as e:
        # Fallback: try openpyxl (read-only, no COM)
        print(f"  [COM unavailable: {e}] — trying openpyxl fallback...")
        return extract_excel_openpyxl(excel_path)

    sections = []
    try:
        for sheet_name, max_rows in SHEETS_TO_EXTRACT:
            try:
                ws = wb.Sheets.Item(sheet_name)
                used = ws.UsedRange
                rows = min(used.Rows.Count, max_rows)
                cols = used.Columns.Count
                lines = [f"\n=== SHEET: {sheet_name} ==="]
                for r in range(1, rows + 1):
                    cells = []
                    for c in range(1, cols + 1):
                        val = ws.Cells(r, c).Text
                        if val and str(val).strip():
                            cells.append(str(val).strip())
                    if cells:
                        lines.append(" | ".join(cells))
                sections.append("\n".join(lines))
                print(f"  [{sheet_name}] {rows} rows extracted.")
            except Exception as e:
                print(f"  [SKIP] {sheet_name}: {e}")
    finally:
        wb.Close(False)
        xl.Quit()

    return "\n\n".join(sections)


def extract_excel_openpyxl(excel_path):
    """Fallback: openpyxl read-only extraction."""
    try:
        import openpyxl
    except ImportError:
        return "[ERROR: Neither win32com nor openpyxl available. Run: pip install openpyxl]"

    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    sections = []
    for sheet_name, max_rows in SHEETS_TO_EXTRACT:
        if sheet_name not in wb.sheetnames:
            print(f"  [SKIP] Sheet not found: {sheet_name}")
            continue
        ws = wb[sheet_name]
        lines = [f"\n=== SHEET: {sheet_name} ==="]
        count = 0
        for row in ws.iter_rows(max_row=max_rows, values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
                count += 1
        sections.append("\n".join(lines))
        print(f"  [{sheet_name}] {count} rows extracted (openpyxl).")
    wb.close()
    return "\n\n".join(sections)


def build_excel_prompt(foundation, excel_text, prompt_text):
    parts = []
    if foundation:
        parts.append(foundation)
    if excel_text:
        parts.append(
            "[EXCEL DATA — universal_domain_mapping_with_coherence_v9.xlsx]\n"
            "[Key sheets extracted directly. This is the primary research data.]\n\n"
            + excel_text
            + "\n\n[END EXCEL DATA]"
        )
    parts.append(prompt_text)
    return "\n\n".join(parts)


def is_o3_model(model_name):
    return model_name.startswith("o3") or model_name.startswith("o1")


def call_deepseek(cfg, full_prompt):
    import openai
    api_key     = cfg.get("DEEPSEEK_API_KEY", "")
    model       = cfg.get("MODEL", "deepseek-chat")
    max_tokens  = int(cfg.get("MAX_TOKENS", "4096") or "0")
    temperature = float(cfg.get("TEMPERATURE", "0.3"))

    client = openai.OpenAI(api_key=api_key, base_url=DEEPSEEK_BASE_URL)
    kwargs = {
        "model": model,
        "messages": [{"role": "user", "content": [{"type": "text", "text": full_prompt}]}],
        "temperature": temperature,
    }
    if max_tokens:
        kwargs["max_tokens"] = max_tokens

    t0       = time.time()
    response = client.chat.completions.create(**kwargs)
    elapsed  = time.time() - t0
    return response.choices[0].message.content, response.usage, elapsed, model


def call_openai(cfg, full_prompt):
    import openai
    api_key    = cfg.get("OPENAI_API_KEY", "")
    model      = cfg.get("OPENAI_MODEL", "o3")
    max_tokens = int(cfg.get("OPENAI_MAX_TOKENS", "16000") or "16000")

    if not api_key or api_key.startswith("sk-PASTE"):
        print("  [OpenAI] Key not configured — skipping.")
        return None, None, 0, model

    client = openai.OpenAI(api_key=api_key)
    kwargs = {
        "model": model,
        "messages": [{"role": "user", "content": full_prompt}],
        "max_completion_tokens": max_tokens,
    }
    if not is_o3_model(model):
        kwargs["temperature"] = float(cfg.get("TEMPERATURE", "1.0"))

    t0 = time.time()
    try:
        response = client.chat.completions.create(**kwargs)
    except openai.AuthenticationError:
        print("  [OpenAI] ERROR: Invalid API key.")
        return None, None, 0, model
    except openai.APIError as e:
        print(f"  [OpenAI] ERROR: {e}")
        return None, None, 0, model

    elapsed = time.time() - t0
    return response.choices[0].message.content, response.usage, elapsed, model


def format_cost(model, in_tok, out_tok):
    if model not in PRICING:
        return f"no pricing for '{model}'"
    ir, or_ = PRICING[model]
    total = (in_tok / 1000) * ir + (out_tok / 1000) * or_
    return f"in={in_tok:,}  out={out_tok:,}  total=${total:.4f}"


def main():
    cfg = parse_config(CONFIG_PATH)

    ds_key = cfg.get("DEEPSEEK_API_KEY", "")
    if not ds_key or ds_key.startswith("sk-PASTE"):
        sys.exit("ERROR: Set DEEPSEEK_API_KEY in config.txt")

    if not PROMPT_FILE.exists():
        sys.exit(f"ERROR: Prompt not found: {PROMPT_FILE}")

    prompt_text = PROMPT_FILE.read_text(encoding="utf-8", errors="replace").strip()
    foundation  = load_foundation(cfg.get("FOUNDATION_DOC", ""))

    print()
    print("=" * 60)
    print("  EXCEL INTERROGATION RUN")
    print("  Prompt 12 × Excel Data × DeepSeek + O3")
    print("=" * 60)

    # Extract Excel
    print(f"\n  Extracting Excel data...")
    excel_text = extract_excel_to_text(EXCEL_PATH)
    print(f"  Excel context: {len(excel_text):,} chars across {len(SHEETS_TO_EXTRACT)} sheets.")

    full_prompt = build_excel_prompt(foundation, excel_text, prompt_text)
    print(f"  Total prompt:  {len(full_prompt):,} chars")

    OUT_DIR.mkdir(exist_ok=True)
    ts      = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = OUT_DIR / f"excel_run_{ts}"
    run_dir.mkdir(exist_ok=True)
    print(f"  Output folder: {run_dir}")

    ds_model = cfg.get("MODEL", "deepseek-chat")
    oa_model = cfg.get("OPENAI_MODEL", "o3")

    # DeepSeek
    print(f"\n  ── DeepSeek ({ds_model}) ──")
    try:
        reply, usage, elapsed, _ = call_deepseek(cfg, full_prompt)
        in_tok  = usage.prompt_tokens     if usage else len(full_prompt) // 4
        out_tok = usage.completion_tokens if usage else 0
        print(f"  Done {elapsed:.1f}s  |  {format_cost(ds_model, in_tok, out_tok)}")
        header = f"# Excel Interrogation × DeepSeek ({ds_model})\n\nDate: {datetime.datetime.now().isoformat()[:16]}\n\n---\n\n"
        (run_dir / "excel_deepseek.md").write_text(header + reply, encoding="utf-8")
    except Exception as e:
        print(f"  ERROR: {e}")

    # OpenAI O3
    print(f"\n  ── OpenAI ({oa_model}) ──")
    reply, usage, elapsed, _ = call_openai(cfg, full_prompt)
    if reply:
        in_tok  = usage.prompt_tokens     if usage else len(full_prompt) // 4
        out_tok = usage.completion_tokens if usage else 0
        print(f"  Done {elapsed:.1f}s  |  {format_cost(oa_model, in_tok, out_tok)}")
        header = f"# Excel Interrogation × OpenAI ({oa_model})\n\nDate: {datetime.datetime.now().isoformat()[:16]}\n\n---\n\n"
        (run_dir / "excel_o3.md").write_text(header + reply, encoding="utf-8")

    print()
    print("=" * 60)
    print(f"  Done. Results in: {run_dir}")
    print("=" * 60)
    print()


if __name__ == "__main__":
    main()
