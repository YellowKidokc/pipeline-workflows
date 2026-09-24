#!/usr/bin/env python3
"""
Offline χ + 7Q evaluator using the local 7Q engine and workbook.

Usage:
  python chi_evaluator_7q.py --input "Your claim here"
  python chi_evaluator_7q.py --input path/to/file.txt --workbook "7Q Full Method.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl


def _import_engine():
    module_root = Path(__file__).resolve().parent
    if str(module_root) not in sys.path:
        sys.path.insert(0, str(module_root))

    try:
        from .core import load_config  # type: ignore
        from .parser import extract_claims_for_paper  # type: ignore
        from .runner import run_text  # type: ignore
        from .chievaluator import CHI_CHANNELS as CHI_CHANNELS_NATIVE  # type: ignore
    except Exception:
        from core import load_config  # type: ignore
        from parser import extract_claims_for_paper  # type: ignore
        from runner import run_text  # type: ignore
        from chievaluator import CHI_CHANNELS as CHI_CHANNELS_NATIVE  # type: ignore

    # Keep one fixed canonical channel order for all outputs.
    return run_text, extract_claims_for_paper, load_config, list(CHI_CHANNELS_NATIVE)


run_text, extract_claims_for_paper, load_config, CHI_CHANNELS = _import_engine()


DEFAULT_CHI_CHANNELS = ["G", "M", "E", "S_eff", "T", "K", "R", "Q", "F", "C"]
TRUTH_WEIGHT_PROFILE = {
    "weights": {"S": 1.0, "E": 1.0, "L": 1.0, "D": 1.0, "P": 1.0, "C": 1.0},
    "labels": {
        "S": "Survivability (branch + adversarial exposure)",
        "E": "Evidence (tier + replication)",
        "L": "Logic (form + precision)",
        "D": "Dependencies (terminus + fragility)",
        "P": "Predictions (confirmed/untested ratio)",
        "C": "Coherence (ISO status + cross-domain force)",
    },
    "confidence_bins": [
        ("ESTABLISHED", 0.85, 1.0),
        ("WELL_SUPPORTED", 0.65, 0.85),
        ("TENTATIVE", 0.40, 0.65),
        ("SPECULATIVE", 0.15, 0.40),
        ("UNSUPPORTED", 0.0, 0.15),
    ],
    "name": "7Q Baseline v1 (fixed, no external tuning)",
}
TRUTH_WEIGHT_TOTAL = sum(TRUTH_WEIGHT_PROFILE["weights"].values())


def _safe_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _now() -> str:
    return datetime.utcnow().isoformat(timespec="seconds") + "Z"


def _read_input_text(raw: str) -> str:
    p = Path(raw)
    if p.exists():
        return p.read_text(encoding="utf-8", errors="ignore")
    return raw


def _normalize_claims(text: str) -> List[Dict[str, str]]:
    """Split a paper/transcript into claim chunks using the existing parser."""
    parsed = extract_claims_for_paper(text)
    if parsed:
        return [
            {
                "claim_id": item.get("claim_id", f"claim-{i}"),
                "text": item.get("statement", "").strip(),
                "heading": item.get("heading", f"claim-{i+1}"),
            }
            for i, item in enumerate(parsed)
            if _safe_str(item.get("statement")).strip()
        ]

    stripped = _safe_str(text).strip()
    if not stripped:
        return []
    return [{"claim_id": "claim-1", "text": stripped, "heading": "claim-1"}]


def _read_csv_claims(
    csv_path: Path,
    claim_column: str,
    source_column: Optional[str] = None,
    heading_column: Optional[str] = None,
    id_column: Optional[str] = None,
    max_claims: int = 0,
) -> List[Dict[str, str]]:
    if not csv_path.exists():
        raise FileNotFoundError(f"CSV not found: {csv_path}")

    claims: List[Dict[str, str]] = []
    with csv_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            return []

        headers = [h.strip() for h in reader.fieldnames if h is not None]
        if claim_column not in headers:
            raise ValueError(f"claim column '{claim_column}' not found in CSV headers: {headers}")

        for row_num, row in enumerate(reader, start=1):
            claim_text = _safe_str(row.get(claim_column, "")).strip()
            if not claim_text:
                continue

            source = ""
            if source_column and source_column in headers:
                source = _safe_str(row.get(source_column, "")).strip()

            heading = ""
            if heading_column and heading_column in headers:
                heading = _safe_str(row.get(heading_column, "")).strip()

            claim_id = ""
            if id_column and id_column in headers:
                claim_id = _safe_str(row.get(id_column, "")).strip()

            claims.append({
                "claim_id": claim_id or f"claim-{len(claims) + 1}",
                "text": claim_text,
                "heading": heading or _safe_str(row.get("heading", "")),
                "source": source or f"row:{row_num}",
            })

            if max_claims > 0 and len(claims) >= max_claims:
                break

    return claims


def _normal_ci(mean: float, std: float, count: int, z: float = 1.96) -> Tuple[float, float]:
    if count <= 0:
        return (0.0, 0.0)
    if std <= 0.0:
        m = round(mean, 6)
        return m, m
    se = std / math.sqrt(count)
    half = z * se
    lo = max(0.0, mean - half)
    hi = min(1.0, mean + half)
    return (round(lo, 6), round(hi, 6))


def _binom_ci(successes: int, total: int, z: float = 1.96) -> Tuple[float, float]:
    if total <= 0:
        return (0.0, 0.0)
    if total == 1:
        p = 1.0 if successes else 0.0
        return (p, p)
    p = successes / total
    denom = 1 + (z ** 2) / total
    centre = p + (z ** 2) / (2 * total)
    radius = z * math.sqrt((p * (1 - p) / total) + (z ** 2) / (4 * total ** 2))
    lo = max(0.0, (centre - radius) / denom)
    hi = min(1.0, (centre + radius) / denom)
    return (round(lo, 6), round(hi, 6))


def _derive_channel_stats(channel_results: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not channel_results:
        return {
            "count": 0,
            "mean_effective": 0.0,
            "mean_pos": 0.0,
            "mean_neg": 0.0,
            "mean_confidence": 0.0,
            "agreement": 0.0,
            "anti_drift_count": 0,
            "high_signal_count": 0,
        }

    effective = [float(r.get("effective_score", 0.0) or 0.0) for r in channel_results]
    pos = [float(r.get("v_pos", 0.0) or 0.0) for r in channel_results]
    neg = [float(r.get("v_neg", 0.0) or 0.0) for r in channel_results]
    conf = [float(r.get("confidence", 0.0) or 0.0) for r in channel_results]
    return {
        "count": float(len(channel_results)),
        "mean_effective": round(statistics.mean(effective), 6),
        "mean_pos": round(statistics.mean(pos), 6),
        "mean_neg": round(statistics.mean(neg), 6),
        "mean_confidence": round(statistics.mean(conf), 6),
        "agreement": round(1.0 - min(1.0, statistics.pstdev(effective)), 6),
        "anti_drift_count": sum(1 for v in effective if v < 0.35),
        "high_signal_count": sum(1 for v in effective if v > 0.7),
    }


def _derive_truth_fields(state) -> Dict[str, float]:
    branch = (_safe_str(state.branch_status).strip().lower() or "open")
    adversarial = bool(getattr(state, "adversarial_tested", False))
    if branch == "alive":
        s_val, s_reason = (1.0, "Branch is alive and adversarially tested") if adversarial else (0.6, "Branch is alive but not adversarially tested")
    elif branch == "problematic":
        s_val, s_reason = 0.3, "Branch explicitly marked problematic"
    elif branch == "dead":
        s_val, s_reason = 0.0, "Branch marked dead"
    elif branch == "untested":
        s_val, s_reason = 0.4, "Branch is untested"
    else:
        s_val, s_reason = 0.5, f"Branch '{branch}' uses fallback baseline"

    tier = _safe_str(state.evidence_tier).strip().lower() or "tier_3"
    replication = _safe_str(state.replication_status).strip().lower() or "unreplicated"
    if tier == "tier_1":
        if replication == "replicated":
            e_val, e_reason = 1.0, "Tier-1 with replication"
        elif replication == "partial":
            e_val, e_reason = 0.8, "Tier-1 with partial replication"
        else:
            e_val, e_reason = 0.9, "Tier-1 without explicit replication status"
    elif tier == "tier_2":
        if replication == "replicated":
            e_val, e_reason = 0.7, "Tier-2 with replication"
        elif replication == "partial":
            e_val, e_reason = 0.5, "Tier-2 with partial replication"
        else:
            e_val, e_reason = 0.4, "Tier-2 with unreplicated status"
    elif tier == "tier_3":
        e_val, e_reason = 0.2, "Tier-3 baseline evidence"
    else:
        e_val, e_reason = 0.0, f"Unknown tier '{tier}'"

    precision = _safe_str(state.precision).strip().lower()
    if precision in {"mathematical", "precise"}:
        l_val, l_reason = 1.0, f"Precision is {precision}"
    elif precision == "detailed":
        l_val, l_reason = 0.8, "Precision is detailed"
    else:
        l_val, l_reason = 0.3, f"Precision is '{precision or 'unknown'}'"

    claim_type = _safe_str(state.claim_type).strip().lower()
    if claim_type == "mathematical":
        l_val = min(1.0, l_val + 0.2)
        l_reason = "Claim is mathematical; +0.2 logic precision bonus"

    terminus = _safe_str(state.chain_terminus).strip().lower()
    fragility = _safe_str(state.fragility).strip().lower()
    if terminus == "axiom":
        if fragility == "survive_independently":
            d_val, d_reason = 1.0, "Axiom chain + survives independently"
        elif fragility == "degrade_gracefully":
            d_val, d_reason = 0.8, "Axiom chain + graceful degradation"
        else:
            d_val, d_reason = 0.6, f"Axiom chain + fragility '{fragility or 'unknown'}'"
    elif terminus == "brute_fact":
        d_val, d_reason = 0.5, "Brute-fact chain terminus"
    elif "circular" in fragility:
        d_val, d_reason = 0.1, "Circularity detected in dependency profile"
    elif terminus in {"open", ""}:
        d_val, d_reason = 0.0, "Chain terminus unresolved/open"
    else:
        d_val, d_reason = 0.3, f"Dependency path fallback for terminus '{terminus}'"

    total_pred = len(state.predictions) or 0
    if total_pred == 0:
        p_val = 0.0
        p_reason = "No predictions recorded"
        p_coverage = 0.0
    else:
        confirmed = len(getattr(state, "confirmed_predictions", []))
        untested = len(getattr(state, "untested_predictions", []))
        failed = len(getattr(state, "failed_predictions", []))
        p_val = (confirmed + 0.5 * untested) / total_pred
        p_val = max(0.0, min(1.0, p_val))
        p_coverage = (confirmed + untested + failed) / max(1, total_pred)
        p_reason = f"Predictions: confirmed={confirmed}, untested={untested}, failed={failed}, total={total_pred}"

    iso = _safe_str(state.iso_status).strip()
    iso_map = {"ISO_CONFIRMED": 1.0, "ISO_PARALLEL": 0.6, "ISO_ANALOGY": 0.3, "NONE": 0.0}
    c_val = iso_map.get(iso, 0.0)
    c_reason = f"ISO status '{iso or 'NONE'}'"
    if bool(getattr(state, "cross_domain_force", False)):
        c_val = min(1.0, c_val + 0.2)
        c_reason = c_reason + " with cross-domain force bonus"

    # Fixed baseline weights; no live user tuning for these.
    weighted_components = {
        "S": s_val * TRUTH_WEIGHT_PROFILE["weights"]["S"],
        "E": e_val * TRUTH_WEIGHT_PROFILE["weights"]["E"],
        "L": l_val * TRUTH_WEIGHT_PROFILE["weights"]["L"],
        "D": d_val * TRUTH_WEIGHT_PROFILE["weights"]["D"],
        "P": p_val * TRUTH_WEIGHT_PROFILE["weights"]["P"],
        "C": c_val * TRUTH_WEIGHT_PROFILE["weights"]["C"],
    }
    weighted_sum = sum(weighted_components.values())
    t_raw = weighted_sum / TRUTH_WEIGHT_TOTAL
    t_val = round(t_raw, 6)

    confidence = "UNSUPPORTED"
    for label, lo, hi in TRUTH_WEIGHT_PROFILE["confidence_bins"]:
        if lo <= t_raw < hi or (label == "ESTABLISHED" and t_raw >= lo):
            confidence = label
            break

    component_values = {
        "S": {"value": round(s_val, 6), "weight": TRUTH_WEIGHT_PROFILE["weights"]["S"], "reason": s_reason},
        "E": {"value": round(e_val, 6), "weight": TRUTH_WEIGHT_PROFILE["weights"]["E"], "reason": e_reason},
        "L": {"value": round(l_val, 6), "weight": TRUTH_WEIGHT_PROFILE["weights"]["L"], "reason": l_reason},
        "D": {"value": round(d_val, 6), "weight": TRUTH_WEIGHT_PROFILE["weights"]["D"], "reason": d_reason},
        "P": {"value": round(p_val, 6), "weight": TRUTH_WEIGHT_PROFILE["weights"]["P"], "reason": p_reason, "coverage": round(p_coverage, 6)},
        "C": {"value": round(c_val, 6), "weight": TRUTH_WEIGHT_PROFILE["weights"]["C"], "reason": c_reason},
    }
    component_disagreement = statistics.pstdev(list(v["value"] for v in component_values.values())) if len(component_values) > 1 else 0.0
    agreement = round(1.0 - min(1.0, component_disagreement), 6)

    return {
        "weights": TRUTH_WEIGHT_PROFILE,
        "S": round(s_val, 6),
        "E": round(e_val, 6),
        "L": round(l_val, 6),
        "D": round(d_val, 6),
        "P": round(p_val, 6),
        "C": round(c_val, 6),
        "T": t_val,
        "confidence": confidence,
        "weighted_sum": round(weighted_sum, 6),
        "component_breakdown": component_values,
        "component_disagreement": round(component_disagreement, 6),
        "component_agreement": agreement,
        "formula_version": TRUTH_WEIGHT_PROFILE["name"],
    }


def _truth_from_claim_output(
    base_state,
    chi_payload: Optional[Dict[str, Any]],
    claim_text: str,
    idx: int,
    source: str,
    claim_id: Optional[str] = None,
) -> Dict[str, Any]:
    truth_fields = _derive_truth_fields(base_state)

    q_scores = getattr(base_state, "q_scores", {}) or {}
    chi_enabled = bool(chi_payload)
    chi_static = 0.0
    chi_gradient = 0.0
    chi_log_gradient = 0.0
    chi_direction = ""
    chi_verdict = ""
    fruit_score = 0.0
    weakest = []
    strongest = []
    zero_channels = []
    channel_base = {}
    pressure_rows = []
    channel_summary = {}
    fruit_profile: Dict[str, Any] = {"dominant_fruits": [], "dominant_antifruits": []}
    fruit_bias_strength = 0.0
    channel_stats = _derive_channel_stats([])

    if chi_enabled:
        chi_static = float(chi_payload.get("static_chi", 0.0))
        chi_gradient = float(getattr(base_state, "chi_gradient", 0.0))
        chi_log_gradient = float(chi_payload.get("gradient_by_channel", {}).get("global", 0.0))
        chi_direction = _safe_str(chi_payload.get("gradient")).strip()
        chi_verdict = _safe_str(chi_payload.get("verdict"))
        fruit = chi_payload.get("fruit_output", {}) or {}
        fruit_score = float(fruit.get("fruit_score", 0.0))
        fruit_profile = {
            "dominant_fruits": list(fruit.get("dominant_fruits", []) or []),
            "dominant_antifruits": list(fruit.get("dominant_antifruits", []) or []),
            "notes": _safe_str(fruit.get("notes")),
        }
        fruit_bias_strength = fruit_score - 0.5
        weakest = list(chi_payload.get("weakest_channels", []) or [])
        strongest = list(chi_payload.get("strongest_channels", []) or [])
        zero_channels = list(chi_payload.get("zero_channels", []) or [])
        channel_results = list(chi_payload.get("channel_results", []) or [])
        for pr in chi_payload.get("pressure_results", []) or []:
            pressure_rows.append({
                "claim_id": f"{source}::{idx}",
                "step": _safe_str(pr.get("pressure_state")),
                "chi": float(pr.get("chi", 0.0)),
                "notes": _safe_str(pr.get("notes")),
            })
        channel_stats = _derive_channel_stats(channel_results)
        channel_base = chi_payload.get("channels_0", {}) or {}
        channel_summary = {
            "per_channel": {
                _safe_str(cr.get("channel")): {
                    "name": _safe_str(cr.get("name")),
                    "v_pos": float(cr.get("v_pos", 0.0)),
                    "v_neg": float(cr.get("v_neg", 0.0)),
                    "effective_score": float(cr.get("effective_score", 0.0)),
                    "confidence": float(cr.get("confidence", 0.0)),
                    "gradient": float(cr.get("gradient", 0.0)),
                    "failure_mode": _safe_str(cr.get("failure_mode")),
                    "repair_path": _safe_str(cr.get("repair_path")),
                    "reasoning": _safe_str(cr.get("reasoning")),
                }
                for cr in channel_results if cr.get("channel")
            }
        }

    out_id = claim_id or f"{source}::{idx}"

    return {
        "id": out_id,
        "source": source,
        "claim_index": idx,
        "claim_text": claim_text,
        "q0": round(float(q_scores.get("Q0", 0.0)), 6),
        "q1": round(float(q_scores.get("Q1", 0.0)), 6),
        "q2": round(float(q_scores.get("Q2", 0.0)), 6),
        "q3": round(float(q_scores.get("Q3", 0.0)), 6),
        "q4": round(float(q_scores.get("Q4", 0.0)), 6),
        "q5": round(float(q_scores.get("Q5", 0.0)), 6),
        "q6": round(float(q_scores.get("Q6", 0.0)), 6),
        "q7": round(float(q_scores.get("Q7", 0.0)), 6),
        "t_score": round(float(getattr(base_state, "t_score", 0.0)), 6),
        "tier": _safe_str(getattr(base_state, "tier", "")),
        "evidence_score": round(float(getattr(base_state, "e_final", 0.0)), 6),
        "S": truth_fields["S"],
        "E": truth_fields["E"],
        "L": truth_fields["L"],
        "D": truth_fields["D"],
        "P": truth_fields["P"],
        "C": truth_fields["C"],
        "truth_score": truth_fields["T"],
        "truth_confidence": truth_fields["confidence"],
        "truth_formula_version": truth_fields["formula_version"],
        "truth_weighted_sum": truth_fields["weighted_sum"],
        "truth_component_disagreement": truth_fields["component_disagreement"],
        "truth_component_agreement": truth_fields["component_agreement"],
        "truth_component_breakdown": truth_fields["component_breakdown"],
        "chi_enabled": chi_enabled,
        "chi_static": round(float(chi_static), 6),
        "chi_gradient": round(float(chi_gradient), 6),
        "chi_log_gradient": round(float(chi_log_gradient), 6),
        "chi_direction": chi_direction,
        "chi_verdict": chi_verdict,
        "fruit_score": round(fruit_score, 6),
        "fruit_bias_strength": round(fruit_bias_strength, 6),
        "fruit_dominant_fruits": ", ".join(fruit_profile.get("dominant_fruits", [])),
        "fruit_dominant_antifruits": ", ".join(fruit_profile.get("dominant_antifruits", [])),
        "fruit_notes": fruit_profile.get("notes", ""),
        "weakest_channels": ",".join(weakest),
        "strongest_channels": ",".join(strongest),
        "zero_channels": ",".join(zero_channels),
        "chi_effective_mean": channel_stats["mean_effective"],
        "chi_effective_std": round(statistics.pstdev([float(r.get("effective_score", 0.0) or 0.0) for r in list(chi_payload.get("channel_results", []) or [])]) if chi_payload and chi_payload.get("channel_results") else 0.0, 6),
        "chi_pos_mean": channel_stats["mean_pos"],
        "chi_neg_mean": channel_stats["mean_neg"],
        "chi_confidence_mean": channel_stats["mean_confidence"],
        "chi_channel_agreement": channel_stats["agreement"],
        "claims_timestamp": _now(),
        "flags": "|".join(getattr(base_state, "flags", []) or []),
        "errors": "|".join(getattr(base_state, "errors", []) or []),
        "warnings": "|".join(getattr(base_state, "warnings", []) or []),
        "channel_trace": channel_summary,
        "pressure_count": len(pressure_rows),
        "_channels_0": channel_base,
        "_pressure_rows": pressure_rows,
        "_state_json": {
            "metadata": {"source": source, "claim_index": idx},
            "chi": chi_payload or {},
            "subclaims": getattr(base_state, "subclaims", []),
            "dependencies": getattr(base_state, "dependencies", []),
            "kill_conditions": getattr(base_state, "kill_conditions", []),
        },
    }


def _numeric_stats(values: List[float]) -> Dict[str, float]:
    if not values:
        return {
            "count": 0,
            "mean": 0.0,
            "median": 0.0,
            "std": 0.0,
            "min": 0.0,
            "max": 0.0,
            "stderr": 0.0,
            "ci_95_low": 0.0,
            "ci_95_high": 0.0,
            "iqr": 0.0,
        }

    values = [float(v) for v in values]
    vals_sorted = sorted(values)
    q1 = statistics.median(vals_sorted[: max(1, len(vals_sorted) // 2)])
    q3 = statistics.median(vals_sorted[len(vals_sorted) // 2:])
    std = statistics.pstdev(values) if len(values) > 1 else 0.0
    mean = statistics.mean(values)
    ci_low, ci_high = _normal_ci(mean, std, len(values))
    return {
        "count": float(len(values)),
        "mean": round(mean, 6),
        "median": round(statistics.median(values), 6),
        "std": round(std, 6),
        "stderr": round(std / math.sqrt(len(values)), 6) if len(values) > 0 else 0.0,
        "ci_95_low": ci_low,
        "ci_95_high": ci_high,
        "min": round(min(values), 6),
        "max": round(max(values), 6),
        "iqr": round(q3 - q1, 6),
    }


def _run_level_stats(claims: List[Dict[str, Any]], chi_channel: List[str], workbook_info: Dict[str, Any]) -> Dict[str, Any]:
    t_scores = [c["t_score"] for c in claims]
    truth_scores = [c["truth_score"] for c in claims]
    chi_stat_values = [c["chi_static"] for c in claims if c["chi_enabled"]]
    chi_grad = [c["chi_gradient"] for c in claims if c["chi_enabled"]]
    fruit_scores = [c["fruit_score"] for c in claims if c.get("fruit_score") is not None]
    channel_agreement = [float(c.get("chi_channel_agreement", 0.0)) for c in claims if c.get("chi_enabled")]
    grad_by_direction = {}
    for c in claims:
        d = c.get("chi_direction") or "n/a"
        grad_by_direction[d] = grad_by_direction.get(d, 0) + 1
    verdict_counts: Dict[str, int] = {}
    tier_counts: Dict[str, int] = {}
    for c in claims:
        verdict_counts[c.get("chi_verdict", "n/a")] = verdict_counts.get(c.get("chi_verdict", "n/a"), 0) + 1
        tier_counts[c.get("tier", "n/a")] = tier_counts.get(c.get("tier", "n/a"), 0) + 1

    observed_channels = list(dict.fromkeys([ch for claim in claims for ch in (claim.get("_channels_0") or {}).keys()]))
    channels = [c for c in chi_channel if c in observed_channels]
    if not channels:
        channels = observed_channels[: len(DEFAULT_CHI_CHANNELS)] or DEFAULT_CHI_CHANNELS
    channel_stats: Dict[str, Any] = {}
    for ch in channels:
        vals = [float(c.get("_channels_0", {}).get(ch, 0.0)) for c in claims if c.get("_channels_0") is not None]
        channel_stats[ch] = _numeric_stats(vals)

    stable_claims = sum(1 for c in claims if _safe_str(c.get("chi_verdict")).lower() in {"coherent", "partially coherent", "repairable"})
    collapsed_claims = sum(1 for c in claims if _safe_str(c.get("chi_verdict")).lower() == "collapsed")
    deception_claims = sum(1 for c in claims if _safe_str(c.get("chi_verdict")).lower() == "high-signal deception")
    chi_conf_hi = _binom_ci(stable_claims, len(claims)) if claims else (0.0, 0.0)
    chi_conf_coll = _binom_ci(collapsed_claims, len(claims)) if claims else (0.0, 0.0)
    chi_conf_deceit = _binom_ci(deception_claims, len(claims)) if claims else (0.0, 0.0)

    fruit_bias = _numeric_stats(fruit_scores)
    fruit_signal_share = sum(1 for c in claims if float(c.get("fruit_score", 0.0)) >= 0.5) / len(claims) if claims else 0.0
    fruit_signal_ci = _binom_ci(sum(1 for c in claims if float(c.get("fruit_score", 0.0)) >= 0.5), len(claims))

    pressure_rows = []
    for c in claims:
        for pr in c.get("_pressure_rows", []):
            row = dict(pr)
            row["chi_verdict"] = c.get("chi_verdict", "")
            row["claim_index"] = c.get("claim_index")
            pressure_rows.append(row)

    return {
        "run_id": _now(),
        "claim_count": len(claims),
        "t_score": _numeric_stats(t_scores),
        "truth_score": _numeric_stats(truth_scores),
        "truth_weighted_sum": _numeric_stats([float(c.get("truth_weighted_sum", 0.0)) for c in claims]),
        "chi_static": _numeric_stats(chi_stat_values),
        "chi_gradient": _numeric_stats(chi_grad),
        "chi_channel_agreement": _numeric_stats(channel_agreement),
        "fruit_score": fruit_scores and _numeric_stats(fruit_scores) or {"count": 0.0, "mean": 0.0, "median": 0.0, "std": 0.0, "stderr": 0.0, "ci_95_low": 0.0, "ci_95_high": 0.0, "min": 0.0, "max": 0.0, "iqr": 0.0},
        "truth_formula": {
            "S": _numeric_stats([float(c["S"]) for c in claims]),
            "E": _numeric_stats([float(c["E"]) for c in claims]),
            "L": _numeric_stats([float(c["L"]) for c in claims]),
            "D": _numeric_stats([float(c["D"]) for c in claims]),
            "P": _numeric_stats([float(c["P"]) for c in claims]),
            "C": _numeric_stats([float(c["C"]) for c in claims]),
            "fruit_score": fruit_bias,
        },
        "chi_direction_counts": grad_by_direction,
        "chi_verdict_counts": verdict_counts,
        "tier_counts": tier_counts,
        "coherence_decision_ci": {
            "stable": {"p": round(stable_claims / max(1, len(claims)), 6), "ci_95": list(chi_conf_hi)},
            "collapsed": {"p": round(collapsed_claims / max(1, len(claims)), 6), "ci_95": list(chi_conf_coll)},
            "high_signal_deception": {"p": round(deception_claims / max(1, len(claims)), 6), "ci_95": list(chi_conf_deceit)},
            "fruit_bias_share": {"p": round(fruit_signal_share, 6), "ci_95": list(fruit_signal_ci)},
        },
        "channel_stats": channel_stats,
        "pressure_rows": pressure_rows,
        "workbook": workbook_info,
    }


def _scan_workbook(path: Path) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out: Dict[str, Any] = {"path": str(path), "sheets": []}
    for ws in wb.worksheets:
        sheet_info: Dict[str, Any] = {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "non_empty_cells": 0,
            "sample_rows": [],
        }

        row_idx = 0
        for row in ws.iter_rows(min_row=1, max_col=min(ws.max_column, 8), values_only=True):
            row_idx += 1
            if row_idx > 40:
                break
            if any(_safe_str(v).strip() for v in row):
                sheet_info["sample_rows"].append([_safe_str(v) for v in row])

        # Count non-empty cells with a soft cap to keep this fast.
        limit = ws.max_row * ws.max_column
        cap = 20000
        non_empty = 0
        for row in ws.iter_rows(min_row=1, values_only=True):
            for v in row:
                if v is not None and str(v).strip():
                    non_empty += 1
                    if non_empty >= cap:
                        break
            if non_empty >= cap:
                break
        sheet_info["non_empty_cells_capped"] = non_empty
        out["sheets"].append(sheet_info)

    return out


def _build_config(use_chi: bool = True) -> Dict[str, Any]:
    cfg = load_config()
    chi_cfg = dict(cfg.get("chi", {}) or {})
    if use_chi:
        chi_cfg.update({
            "enabled": True,
            "steps": chi_cfg.get("steps", 9),
            "zero_cutoff": chi_cfg.get("zero_cutoff", 0.05),
            "fruit": chi_cfg.get("fruit", {"beta": 0.35, "chi_threshold": 0.42}),
        })
    else:
        chi_cfg["enabled"] = False
    cfg["chi"] = chi_cfg
    return cfg


def _export_excel(path: Path, claims: List[Dict[str, Any]], run_stats: Dict[str, Any]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "claim_level"
    headers = [
        "claim_id", "source", "claim_index", "claim_text", "claim_heading", "run_mode", "t_score", "tier", "truth_score", "truth_confidence",
        "truth_formula_version", "truth_weighted_sum", "truth_component_disagreement", "truth_component_agreement",
        "q0", "q1", "q2", "q3", "q4", "q5", "q6", "q7", "evidence_score",
        "S", "E", "L", "D", "P", "C",
        "chi_enabled", "chi_static", "chi_gradient", "chi_log_gradient", "chi_direction", "chi_verdict",
        "fruit_score", "fruit_bias_strength", "fruit_dominant_fruits", "fruit_dominant_antifruits", "fruit_notes",
        "chi_effective_mean", "chi_pos_mean", "chi_neg_mean", "chi_confidence_mean", "chi_channel_agreement",
        "weakest_channels", "strongest_channels", "zero_channels", "flags", "errors", "warnings",
        "timestamp", "pressure_count",
    ]
    ws.append(headers)
    for c in claims:
        ws.append([
            c["id"],
            c["source"],
            c["claim_index"],
            c["claim_text"],
            _safe_str(c.get("claim_heading")),
            "forward",
            c["t_score"],
            c["tier"],
            c["truth_score"],
            c["truth_confidence"],
            c["truth_formula_version"],
            c["truth_weighted_sum"],
            c["truth_component_disagreement"],
            c["truth_component_agreement"],
            c["q0"], c["q1"], c["q2"], c["q3"], c["q4"], c["q5"], c["q6"], c["q7"],
            c["evidence_score"],
            c["S"], c["E"], c["L"], c["D"], c["P"], c["C"],
            c["chi_enabled"], c["chi_static"], c["chi_gradient"], c["chi_log_gradient"], c["chi_direction"], c["chi_verdict"],
            c["fruit_score"], c["fruit_bias_strength"], c["fruit_dominant_fruits"], c["fruit_dominant_antifruits"], c["fruit_notes"],
            c["chi_effective_mean"], c["chi_pos_mean"], c["chi_neg_mean"], c["chi_confidence_mean"], c["chi_channel_agreement"],
            c["weakest_channels"], c["strongest_channels"], c["zero_channels"],
            c["flags"], c["errors"], c["warnings"],
            c["claims_timestamp"], c["pressure_count"],
        ])

    ws2 = wb.create_sheet("chi_pressure")
    ws2.append(["claim_id", "claim_index", "pressure_state", "chi", "notes", "chi_verdict"])
    for c in claims:
        for pr in c.get("_pressure_rows", []):
            ws2.append([c["id"], c["claim_index"], pr["step"], pr["chi"], pr["notes"], c.get("chi_verdict", "")])

    ws3 = wb.create_sheet("run_summary")
    ws3.append(["metric", "value"])
    ws3.append(["run_id", run_stats["run_id"]])
    ws3.append(["claim_count", run_stats["claim_count"]])
    ws3.append(["t_score_mean", run_stats["t_score"]["mean"]])
    ws3.append(["truth_score_mean", run_stats["truth_score"]["mean"]])
    ws3.append(["chi_static_mean", run_stats["chi_static"]["mean"]])
    ws3.append(["chi_gradient_mean", run_stats["chi_gradient"]["mean"]])
    ws3.append(["chi_channel_agreement_mean", run_stats["chi_channel_agreement"]["mean"]])
    ws3.append(["fruit_score_mean", run_stats.get("fruit_score", {}).get("mean", 0.0)])
    ws3.append(["S_mean", run_stats["truth_formula"]["S"]["mean"]])
    ws3.append(["E_mean", run_stats["truth_formula"]["E"]["mean"]])
    ws3.append(["L_mean", run_stats["truth_formula"]["L"]["mean"]])
    ws3.append(["D_mean", run_stats["truth_formula"]["D"]["mean"]])
    ws3.append(["P_mean", run_stats["truth_formula"]["P"]["mean"]])
    ws3.append(["C_mean", run_stats["truth_formula"]["C"]["mean"]])
    ws3.append(["truth_weighted_sum_mean", run_stats["truth_weighted_sum"]["mean"]])
    ws3.append(["truth_ci_95_low", run_stats["truth_score"]["ci_95_low"]])
    ws3.append(["truth_ci_95_high", run_stats["truth_score"]["ci_95_high"]])
    ws3.append(["chi_ci_95_low", run_stats["chi_static"]["ci_95_low"]])
    ws3.append(["chi_ci_95_high", run_stats["chi_static"]["ci_95_high"]])
    ws3.append(["fruit_ci_95_low", run_stats.get("fruit_score", {}).get("ci_95_low", 0.0)])
    ws3.append(["fruit_ci_95_high", run_stats.get("fruit_score", {}).get("ci_95_high", 0.0)])
    ws3.append(["coherence_stable_p", run_stats["coherence_decision_ci"]["stable"]["p"]])
    ws3.append(["coherence_stable_ci_low", run_stats["coherence_decision_ci"]["stable"]["ci_95"][0]])
    ws3.append(["coherence_stable_ci_high", run_stats["coherence_decision_ci"]["stable"]["ci_95"][1]])
    ws3.append(["collapsed_p", run_stats["coherence_decision_ci"]["collapsed"]["p"]])
    ws3.append(["collapsed_ci_low", run_stats["coherence_decision_ci"]["collapsed"]["ci_95"][0]])
    ws3.append(["collapsed_ci_high", run_stats["coherence_decision_ci"]["collapsed"]["ci_95"][1]])
    ws3.append(["deception_p", run_stats["coherence_decision_ci"]["high_signal_deception"]["p"]])
    ws3.append(["deception_ci_low", run_stats["coherence_decision_ci"]["high_signal_deception"]["ci_95"][0]])
    ws3.append(["deception_ci_high", run_stats["coherence_decision_ci"]["high_signal_deception"]["ci_95"][1]])
    ws3.append(["fruit_bias_share", run_stats["coherence_decision_ci"]["fruit_bias_share"]["p"]])
    ws3.append(["fruit_bias_share_ci_low", run_stats["coherence_decision_ci"]["fruit_bias_share"]["ci_95"][0]])
    ws3.append(["fruit_bias_share_ci_high", run_stats["coherence_decision_ci"]["fruit_bias_share"]["ci_95"][1]])
    ws3.append(["workbook_path", run_stats["workbook"].get("path", "")])
    ws3.append(["workbook_sheet_count", len(run_stats["workbook"].get("sheets", []))])

    ws3.append([])
    ws3.append(["key", "count"])
    for k, v in run_stats["chi_direction_counts"].items():
        ws3.append([f"chi_direction::{k}", v])
    ws3.append(["tally::chi_verdict"])
    for k, v in run_stats["chi_verdict_counts"].items():
        ws3.append([k, v])
    ws3.append(["tally::tier"])
    for k, v in run_stats["tier_counts"].items():
        ws3.append([k, v])

    ws4 = wb.create_sheet("channel_stats")
    ws4.append(["channel", "count", "mean", "median", "std", "min", "max"])
    for ch, vals in run_stats["channel_stats"].items():
        ws4.append([ch, vals["count"], vals["mean"], vals["median"], vals["std"], vals["min"], vals["max"]])

    ws5 = wb.create_sheet("workbook_profile")
    ws5.append(["sheet", "max_row", "max_col", "non_empty_cells_capped"])
    for sh in run_stats["workbook"].get("sheets", []):
        ws5.append([sh["title"], sh["max_row"], sh["max_col"], sh.get("non_empty_cells_capped", 0)])

    ws6 = wb.create_sheet("truth_component_breakdown")
    ws6.append(["claim_id", "component", "weight", "value", "reason", "coverage", "run_weighted_sum"])
    for claim in claims:
        breakdown = claim.get("truth_component_breakdown", {})
        for comp_name, comp in breakdown.items():
            ws6.append([
                claim["id"],
                comp_name,
                float(comp.get("weight", 0.0)),
                float(comp.get("value", 0.0)),
                _safe_str(comp.get("reason")),
                float(comp.get("coverage", 1.0)),
                float(claim.get("truth_weighted_sum", 0.0)),
            ])

    ws7 = wb.create_sheet("chi_channel_trace")
    ws7.append(["claim_id", "channel", "name", "v_pos", "v_neg", "effective_score", "confidence", "failure_mode", "repair_path", "reasoning"])
    for claim in claims:
        trace = claim.get("channel_trace", {}).get("per_channel", {})
        for channel, payload in trace.items():
            ws7.append([
                claim["id"],
                channel,
                _safe_str(payload.get("name")),
                float(payload.get("v_pos", 0.0)),
                float(payload.get("v_neg", 0.0)),
                float(payload.get("effective_score", 0.0)),
                float(payload.get("confidence", 0.0)),
                _safe_str(payload.get("failure_mode")),
                _safe_str(payload.get("repair_path")),
                _safe_str(payload.get("reasoning")),
            ])

    wb.save(path)


def _export_json(path: Path, claims: List[Dict[str, Any]], run_stats: Dict[str, Any]) -> None:
    payload = {
        "schema": "chi_evaluator_7q.v1.1",
        "timestamp": _now(),
        "run": run_stats,
        "claims": [{k: v for k, v in c.items() if not k.startswith("_")} for c in claims],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _parse_args():
    p = argparse.ArgumentParser(description="Run the χ 7Q evaluator offline.")
    p.add_argument("--input", "-i", default=None, help="Claim text or path to text/JSON file")
    p.add_argument("--input-csv", default=None, help="Path to CSV file containing claims")
    p.add_argument("--workbook", "-w", default=str(Path(__file__).resolve().parent / "7Q Full Method.xlsx"), help="Path to 7Q Full Method.xlsx")
    p.add_argument("--output", "-o", default=None, help="Output Excel report path (default: chi_7q_run_<timestamp>.xlsx)")
    p.add_argument("--json", default=None, help="Optional JSON output path")
    p.add_argument("--no-chi", action="store_true", help="Disable χ evaluator (7Q scores only)")
    p.add_argument("--source", default="direct_input", help="Source label for claim IDs")
    p.add_argument("--max-claims", type=int, default=0, help="Limit number of claims to process (0 means all)")
    p.add_argument("--claim-column", default="claim", help="CSV claim column name (required when --input-csv is used)")
    p.add_argument("--source-column", default=None, help="CSV source column name")
    p.add_argument("--heading-column", default=None, help="CSV heading/title column name")
    p.add_argument("--id-column", default=None, help="CSV ID column name")
    return p.parse_args()


def main() -> int:
    args = _parse_args()
    if not args.input and not args.input_csv:
        print("ERROR: either --input or --input-csv is required", file=sys.stderr)
        return 1

    wb_path = Path(args.workbook)
    if not wb_path.exists():
        print(f"ERROR: workbook not found: {wb_path}", file=sys.stderr)
        return 1

    claims: List[Dict[str, str]] = []
    if args.input_csv:
        try:
            claims = _read_csv_claims(
                csv_path=Path(args.input_csv),
                claim_column=args.claim_column,
                source_column=args.source_column,
                heading_column=args.heading_column,
                id_column=args.id_column,
                max_claims=args.max_claims,
            )
        except Exception as exc:
            print(f"ERROR: failed to load csv input: {exc}", file=sys.stderr)
            return 1
    else:
        raw_input = _read_input_text(args.input)
        claims = _normalize_claims(raw_input)

        if args.max_claims > 0:
            claims = claims[: args.max_claims]

    if not claims:
        print("ERROR: no claim content found", file=sys.stderr)
        return 1

    cfg = _build_config(use_chi=not args.no_chi)

    rows = []
    for i, item in enumerate(claims, start=1):
        claim_text = _safe_str(item.get("text", "")).strip()
        if not claim_text:
            continue
        state = run_text(content=claim_text, claim_id=item.get("claim_id", f"claim-{i}"), title=_safe_str(item.get("heading")), cfg=cfg, source_path=args.input)
        if getattr(state, "run_mode", None) and not _safe_str(getattr(state, "run_mode")).strip():
            state.run_mode = "forward"
        source = item.get("source", args.source)
        row = _truth_from_claim_output(
            state,
            state.chi_result if getattr(state, "chi_result", None) else None,
            claim_text,
            i,
            source,
            claim_id=item.get("claim_id"),
        )
        row["claim_heading"] = _safe_str(item.get("heading"))
        row["claim_text_snippet"] = claim_text[:1800]
        row["source"] = _safe_str(source)
        rows.append(row)

    if not rows:
        print("ERROR: no claims emitted", file=sys.stderr)
        return 1

    workbook_info = _scan_workbook(wb_path)
    run_stats = _run_level_stats(rows, CHI_CHANNELS, workbook_info)

    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    out_path = Path(args.output or Path(".").resolve() / f"chi_7q_run_{ts}.xlsx")
    _export_excel(out_path, rows, run_stats)
    print(f"Wrote Excel: {out_path}")

    if args.json:
        json_path = Path(args.json)
        _export_json(json_path, rows, run_stats)
        print(f"Wrote JSON: {json_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
