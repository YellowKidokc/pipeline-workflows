#!/usr/bin/env python3
"""
compare_reviews.py
==================
Compare Python baseline packet metrics against model review outputs.
Produces:
- comparison.jsonl (per claim)
- comparison_summary.json
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import pathlib
from typing import Any, Dict, List, Tuple

try:
    import openpyxl  # type: ignore
except Exception:  # pragma: no cover
    openpyxl = None  # type: ignore


def read_jsonl(path: pathlib.Path) -> List[Dict[str, Any]]:
    data = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                data.append(json.loads(line))
    return data


def bucket_from_verdict(v: str) -> float:
    v = str(v or "").lower()
    mapping = {
        "coherent": 4.0,
        "partially coherent": 3.0,
        "fragile": 2.0,
        "repairable": 2.0,
        "partial": 2.0,
        "collapsed": 1.0,
        "repairable/fragile": 2.0,
    }
    return mapping.get(v, 2.0)


def label_to_bucket(label: str) -> float:
    v = (label or "").lower().strip()
    buckets = {
        "coherent": 4.0,
        "borderline": 3.0,
        "fragile": 2.0,
        "incoherent": 1.0,
        "unclear": 2.0,
    }
    return buckets.get(v, 2.0)


def safe_num(v: Any) -> float:
    try:
        if v is None:
            return float("nan")
        return float(v)
    except Exception:
        return float("nan")


def cmp_record(packet: Dict[str, Any], models: Dict[str, Any]) -> Dict[str, Any]:
    base_score = bucket_from_verdict(packet.get("chi_verdict", ""))
    base_fruit = safe_num(packet.get("fruit_score"))
    out: Dict[str, Any] = {
        "packet_id": packet.get("packet_id"),
        "source_file": packet.get("source_file"),
        "chi_verdict": packet.get("chi_verdict"),
        "chi_static": safe_num(packet.get("chi_static")),
        "fruit_score": base_fruit,
        "model_scores": {},
        "disagreement_flags": [],
    }
    max_delta = 0.0
    for model_name, payload in models.items():
        parsed = payload.get("parsed", {})
        m_label = parsed.get("label")
        m_coh = safe_num(parsed.get("coherence_score"))
        m_fruit = safe_num(parsed.get("fruit_bias_score"))
        m_conf = safe_num(parsed.get("confidence"))
        if math.isnan(m_coh):
            continue
        out["model_scores"][model_name] = {
            "coherence_score": m_coh,
            "fruit_score": m_fruit,
            "confidence": m_conf,
            "label": m_label,
        }
        m_bucket = label_to_bucket(m_label)
        delta = abs(m_bucket - base_score)
        max_delta = max(max_delta, delta)
        if delta >= 2:
            out["disagreement_flags"].append(f"{model_name}:verdict_gap={delta:.1f}")
        if not math.isnan(m_fruit) and not math.isnan(base_fruit):
            if abs(m_fruit - (2 * base_fruit - 1)) >= 0.5:
                out["disagreement_flags"].append(f"{model_name}:fruit_gap={m_fruit - (2 * base_fruit - 1):.2f}")
        if not math.isnan(m_coh) and not math.isnan(out["chi_static"]) and (m_coh - out["chi_static"] * 10.0) > 1.5:
            out["disagreement_flags"].append(f"{model_name}:high_coherence_delta")
    out["max_disagreement"] = max_delta
    return out


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Compare packet baselines with model reviews.")
    parser.add_argument("--packets", required=True, help="Feature packets JSONL")
    parser.add_argument("--reviews", required=True, help="Model reviews JSONL")
    parser.add_argument("--out", default="compare", help="Output folder")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    packets = read_jsonl(pathlib.Path(args.packets))
    reviews = {r.get("packet_id"): r for r in read_jsonl(pathlib.Path(args.reviews))}

    out_dir = pathlib.Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    records: List[Dict[str, Any]] = []
    for packet in packets:
        pid = packet.get("packet_id")
        rec = cmp_record(packet, reviews.get(pid, {}).get("models", {}))
        rec["packet_present_in_reviews"] = pid in reviews
        records.append(rec)

    cmp_path = out_dir / "comparison.jsonl"
    with cmp_path.open("w", encoding="utf-8") as handle:
        for rec in records:
            handle.write(json.dumps(rec, ensure_ascii=False) + "\n")

    # Top disagreement list
    sorted_records = sorted(records, key=lambda r: r.get("max_disagreement", 0), reverse=True)
    flags = sum(1 for r in records if r.get("disagreement_flags"))
    avg_gap = (
        sum(r.get("max_disagreement", 0) for r in records) / len(records)
        if records else 0.0
    )

    summary = {
        "total_packets": len(records),
        "matched_reviews": sum(1 for r in records if r["packet_present_in_reviews"]),
        "records_with_disagreement": flags,
        "mean_disagreement": round(avg_gap, 4),
    }
    summary_path = out_dir / "comparison_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    # CSV for easy triage
    csv_path = out_dir / "comparison_disagreements.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow([
            "packet_id",
            "source_file",
            "chi_verdict",
            "chi_static",
            "fruit_score",
            "max_disagreement",
            "flags",
        ])
        for rec in sorted_records[: min(20, len(sorted_records))]:
            writer.writerow([
                rec.get("packet_id"),
                rec.get("source_file"),
                rec.get("chi_verdict"),
                rec.get("chi_static"),
                rec.get("fruit_score"),
                rec.get("max_disagreement"),
                " | ".join(rec.get("disagreement_flags", [])),
            ])

    excel_path = out_dir / "comparison.xlsx"
    if openpyxl is not None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comparison"
        ws.append([
            "packet_id",
            "source_file",
            "chi_verdict",
            "chi_static",
            "fruit_score",
            "model",
            "label",
            "coherence_score",
            "fruit_bias_score",
            "confidence",
            "max_disagreement",
            "flags",
        ])
        model_names = {"deepseek", "openai"}
        for rec in sorted_records:
            packet_present = rec.get("packet_present_in_reviews")
            if not rec.get("model_scores"):
                ws.append([
                    rec.get("packet_id"),
                    rec.get("source_file"),
                    rec.get("chi_verdict"),
                    rec.get("chi_static"),
                    rec.get("fruit_score"),
                    "none",
                    "",
                    "",
                    "",
                    "",
                    rec.get("max_disagreement"),
                    " | ".join(rec.get("disagreement_flags", [])),
                ])
                continue
            for model_name in model_names:
                if model_name not in rec.get("model_scores", {}):
                    continue
                scores = rec["model_scores"][model_name]
                ws.append([
                    rec.get("packet_id"),
                    rec.get("source_file"),
                    rec.get("chi_verdict"),
                    rec.get("chi_static"),
                    rec.get("fruit_score"),
                    model_name,
                    scores.get("label"),
                    scores.get("coherence_score"),
                    scores.get("fruit_score"),
                    scores.get("confidence"),
                    rec.get("max_disagreement"),
                    " | ".join(rec.get("disagreement_flags", [])),
                ])

        ws_summary = wb.create_sheet("Summary")
        ws_summary.append(["metric", "value"])
        for k, v in summary.items():
            ws_summary.append([k, v])
        ws_mat = wb.create_sheet("SourceAgreement")
        ws_mat.append(["packet_id", "source_file", "has_reviews", "max_disagreement"])
        for rec in sorted_records:
            ws_mat.append([
                rec.get("packet_id"),
                rec.get("source_file"),
                rec.get("packet_present_in_reviews"),
                rec.get("max_disagreement"),
            ])
        wb.save(excel_path)
    else:
        excel_path = out_dir / "comparison.xlsx.missing_openpyxl"

    print(f"Comparison : {cmp_path}")
    print(f"Summary    : {summary_path}")
    print(f"CSV        : {csv_path}")
    print(f"Excel      : {excel_path}")
    print(f"Packets    : {len(records)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
