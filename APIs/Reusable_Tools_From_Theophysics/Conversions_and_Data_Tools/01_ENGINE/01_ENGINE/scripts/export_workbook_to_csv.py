#!/usr/bin/env python3
"""
Export every worksheet in an .xlsx workbook to separate CSV files.
"""

from __future__ import annotations

import argparse
import csv
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def safe_name(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|]+', "_", name).strip()
    s = re.sub(r"\s+", "_", s)
    return s or "Sheet"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export workbook sheets to CSV files")
    parser.add_argument("xlsx_path", help="Path to .xlsx file")
    parser.add_argument(
        "--out-dir",
        default="",
        help="Output directory (default: <workbook_stem>_CSV_EXPORT_<timestamp>)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    xlsx_path = Path(args.xlsx_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Workbook not found: {xlsx_path}")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = (
        Path(args.out_dir)
        if args.out_dir
        else xlsx_path.parent / f"{xlsx_path.stem}_CSV_EXPORT_{ts}"
    )
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(filename=str(xlsx_path), data_only=False, read_only=True)

    used_names = set()
    manifest_rows = []

    for idx, ws in enumerate(wb.worksheets, start=1):
        base = f"{idx:03d}_{safe_name(ws.title)}.csv"
        name = base
        suffix = 2
        while name.lower() in used_names:
            name = base[:-4] + f"_{suffix}.csv"
            suffix += 1
        used_names.add(name.lower())

        csv_path = out_dir / name
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            if max_row > 0 and max_col > 0:
                for row in ws.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_col,
                    values_only=True,
                ):
                    writer.writerow(["" if v is None else v for v in row])

        manifest_rows.append(
            {
                "sheet_index": idx,
                "sheet_name": ws.title,
                "csv_file": name,
                "max_row": max_row,
                "max_col": max_col,
            }
        )

    manifest_path = out_dir / "manifest.csv"
    with manifest_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["sheet_index", "sheet_name", "csv_file", "max_row", "max_col"],
        )
        writer.writeheader()
        writer.writerows(manifest_rows)

    print(f"Workbook: {xlsx_path}")
    print(f"Output: {out_dir}")
    print(f"Sheets exported: {len(manifest_rows)}")
    print(f"Manifest: {manifest_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

